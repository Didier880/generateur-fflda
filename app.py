import streamlit as st
import pandas as pd
from datetime import datetime, timedelta, time
import io
import urllib.request
import re
import collections
import streamlit.components.v1 as components
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter

# --- CONFIGURATION DE LA PAGE ---
st.set_page_config(page_title="Générateur Officiel FFLDA", page_icon="🤼", layout="wide")

# --- MENU LATÉRAL (PARAMÈTRES INTERACTIFS) ---
with st.sidebar:
    import os
    col_l1, col_l2, col_l3 = st.columns([1, 2, 1])
    with col_l2:
        if os.path.exists("logo_fflda.png"):
            st.image("logo_fflda.png", use_container_width=True)
        else:
            st.image("https://www.fflutte.com/content/uploads/2021/10/fflutte-bleu-1024x842.png", use_container_width=True)
    st.markdown("### Paramètres du tournoi")
    st.markdown("---")
    
    # --- NOM DE LA COMPÉTITION ---
    nom_competition = st.text_input("🏆 Nom de la compétition", value="Tournoi Officiel FFLDA - U7/U9/U11/U13")
    
    with st.expander("⚙️ 1. Logistique & Pesées", expanded=True):
        nb_tapis = st.number_input("Nombre de tapis", min_value=1, max_value=10, value=3)
        type_pesee = st.radio("Format des pesées", ["1 Pesée (Générale)", "2 Pesées (U7/U9 puis U11/U13)"], index=1)
        label_pesee_1 = "1ère pesée" if "1" in type_pesee else "Pesée U7/U9"
        heure_pesee_u9 = st.time_input(label_pesee_1, value=time(9, 0))
        duree_pesee = st.selectbox("Durée allouée à la pesée + échauffement (min)", [30, 45, 60, 90], index=1)
        
    with st.expander("⏱️ 2. Pause de la compétition", expanded=False):
        activer_pause = st.checkbox("Activer la pause", value=True)
        duree_pause = st.selectbox("Durée de la pause (min)", [30, 45, 60, 75, 90], index=2) if activer_pause else 0
    
    with st.expander("🤼 3. Règles Sportives & Temps", expanded=False):
        mixte_active = st.checkbox("Catégories Mixtes (U7, U9, U11 uniquement)", value=True, help="Conformément à la réglementation officielle FFLDA, la mixité n'existe plus en U13 : les catégories Féminine (LF), Libre (LL) et Gréco (LG) sont toujours strictement séparées.")
        poules_par_niveau = st.checkbox("Créer des poules par niveau (débutants/confirmés)", value=True)
        separer_clubs = st.checkbox("Éviter les lutteurs d'un même club dans la même poule (dans la mesure du possible)", value=True)
        eviter_arbitre_meme_club = st.checkbox("Éviter les matchs entre arbitres et lutteurs du même club", value=True)
        meme_tapis_poule = st.checkbox("Maintenir chaque poule / lutteur sur un même tapis", value=True, help="Chaque catégorie de même style et de même poids (ex: U13 Gréco 30 kg) est affectée intégralement à un tapis fixe unique pour tous ses combats. Les catégories de styles ou poids différents sont réparties de manière à équilibrer au mieux le nombre total de matchs par tapis.")
        tolerance_poids = st.number_input("Tolérance d'écart de poids (%) [U7, U9, U11]", min_value=10, max_value=15, value=10, step=1)
        repos_matchs = st.number_input("Matchs de repos minimum", min_value=1, max_value=10, value=3)
        duree_plateau_u7 = st.number_input("Temps de chaque plateau U7 (min)", min_value=3, max_value=30, value=10, step=1, help="L'animation U7 est sous forme de 3 plateaux d'activités avec rotation. Par exemple, 10 min par plateau = 3x10 min = 30 min consacrées aux U7 au total.")
        duree_u7 = duree_plateau_u7 * 3
        duree_u9 = st.number_input("Temps total U9 (min)", value=3)
        duree_u11 = st.number_input("Temps total U11 (min)", value=4)
        duree_u13 = st.number_input("Temps total U13 (min)", value=5)

    # JavaScript pour le comportement d'accordéon à ouverture unique (ferme les autres au clic)
    components.html("""
    <script>
    (function() {
        const parentDoc = window.parent.document;
        function initAccordion() {
            const sidebar = parentDoc.querySelector('section[data-testid="stSidebar"]');
            if (!sidebar) return;
            const expanders = sidebar.querySelectorAll('div[data-testid="stExpander"]');
            expanders.forEach(exp => {
                if (exp.dataset.accordionAttached) return;
                exp.dataset.accordionAttached = "true";
                exp.addEventListener('click', function() {
                    setTimeout(() => {
                        const details = exp.querySelector('details');
                        if (details && details.open) {
                            expanders.forEach(otherExp => {
                                if (otherExp !== exp) {
                                    const otherDetails = otherExp.querySelector('details');
                                    if (otherDetails) {
                                        otherDetails.open = false;
                                    }
                                }
                            });
                        }
                    }, 50);
                }, false);
            });
        }
        initAccordion();
        const observer = new MutationObserver(initAccordion);
        observer.observe(parentDoc.body, { childList: true, subtree: true });
    })();
    </script>
    """, height=0)

def formater_poids(val):
    if val is None or str(val).strip() in ['', 'None', 'nan']:
        return ''
    val_str = str(val).lower().replace('kg', '').replace(',', '.').strip()
    try:
        val_float = round(float(val_str), 1)
        if val_float == int(val_float):
            num_str = str(int(val_float))
        else:
            num_str = str(val_float)
        return f"{num_str} kg"
    except (ValueError, TypeError):
        s = str(val).strip()
        return f"{s} kg" if (s and not s.lower().endswith('kg')) else s

def formater_poids_court(val):
    if val is None or str(val).strip() in ['', 'None', 'nan']:
        return ''
    val_str = str(val).lower().replace('kg', '').replace(',', '.').strip()
    try:
        val_float = round(float(val_str), 1)
        if val_float == int(val_float):
            num_str = str(int(val_float))
        else:
            num_str = str(val_float)
        return f"{num_str}kg"
    except (ValueError, TypeError):
        s = str(val).strip()
        return f"{s}kg" if (s and not s.lower().endswith('kg')) else s

def abreger_nom_onglet(nom_poule):
    txt = str(nom_poule)
    txt = txt.replace("Mixte (LL/LF)", "Mxt").replace("LG (Gréco)", "LG").replace("LL (Libre)", "LL").replace("LF (Féminine)", "LF")
    txt = txt.replace(" | ", " ").replace(" (", " ").replace(")", "").replace(" - ", "-")
    txt = txt.replace("/", "-").replace("\\", "-").replace(":", "-").replace("?", "").replace("*", "")
    txt = re.sub(r'\s+', ' ', txt)
    return txt[:31].strip()

def charger_liste_arbitres(fichier_arbitres_in=None):
    """
    Charge la liste des arbitres inscrits depuis le fichier téléversé ou le fichier par défaut FFLDA - Inscription arbitres.xlsx.
    """
    arbitres = []
    filepath = None
    
    if fichier_arbitres_in is not None:
        filepath = fichier_arbitres_in
    else:
        import os
        default_name = 'FFLDA - Inscription arbitres.xlsx'
        if os.path.exists(default_name):
            filepath = default_name
            
    if filepath is None:
        return arbitres
        
    try:
        if hasattr(filepath, 'name') and filepath.name.endswith('.csv'):
            df_arb = pd.read_csv(filepath, sep=';', encoding='utf-8')
            if len(df_arb.columns) == 1:
                filepath.seek(0)
                df_arb = pd.read_csv(filepath, sep=',', encoding='utf-8')
        elif isinstance(filepath, str) and filepath.endswith('.csv'):
            df_arb = pd.read_csv(filepath, sep=';', encoding='utf-8')
            if len(df_arb.columns) == 1:
                df_arb = pd.read_csv(filepath, sep=',', encoding='utf-8')
        else:
            df_temp = pd.read_excel(filepath, nrows=5)
            header_row = 0
            for i, row in df_temp.iterrows():
                if 'Licence' in str(row.values) or 'Nom' in str(row.values) or "Inscrit Par" in str(row.values):
                    header_row = i + 1
                    break
            if hasattr(filepath, 'seek'):
                filepath.seek(0)
            df_arb = pd.read_excel(filepath, header=header_row)

        col_nom = next((c for c in df_arb.columns if str(c).strip().lower() == 'nom'), None)
        col_prenom = next((c for c in df_arb.columns if 'prénom' in str(c).strip().lower() or 'prenom' in str(c).strip().lower()), None)
        col_club = next((c for c in df_arb.columns if any(k in str(c).strip().lower() for k in ['sigle du club', 'club'])), None)
        col_comite = next((c for c in df_arb.columns if any(k in str(c).strip().lower() for k in ['comité', 'comite', 'ligue', 'région', 'region'])), None)
        col_licence = next((c for c in df_arb.columns if 'licence' in str(c).strip().lower()), None)

        for _, row in df_arb.iterrows():
            nom_val = str(row[col_nom]).strip() if (col_nom and pd.notna(row[col_nom])) else ''
            prenom_val = str(row[col_prenom]).strip() if (col_prenom and pd.notna(row[col_prenom])) else ''
            if not nom_val or nom_val.lower() in ['nan', 'none', 'photo']:
                continue
            
            club_val = str(row[col_club]).strip() if (col_club and pd.notna(row[col_club])) else 'Indépendant'
            comite_val = str(row[col_comite]).strip() if (col_comite and pd.notna(row[col_comite])) else 'Comité Non Renseigné'
            licence_val = str(row[col_licence]).strip() if (col_licence and pd.notna(row[col_licence])) else ''

            arbitres.append({
                'Nom_Complet': f"{nom_val} {prenom_val}".strip(),
                'Nom': nom_val,
                'Prenom': prenom_val,
                'Licence': licence_val,
                'Club': club_val if club_val not in ['', 'None', 'nan', '-'] else 'Indépendant',
                'Comite': comite_val if comite_val not in ['', 'None', 'nan', '-'] else 'Comité Non Renseigné'
            })
    except Exception:
        pass
        
    return arbitres

# --- CORPS PRINCIPAL ---
st.title(f"🏆 {nom_competition}")
st.markdown("**Plateforme officielle d'optimisation des tournois de jeunes et d'édition des bilans fédéraux.**")
st.markdown("---")

def generer_rondes_fflda(participants_in):
    participants = list(participants_in)
    n = len(participants)
    
    if n <= 1:
        return []
    elif n == 2:
        return [[(participants[0], participants[1])]]
    elif n == 3:
        return [[(participants[0], participants[1])],
                [(participants[2], participants[0])],
                [(participants[1], participants[2])]]
    elif n == 4:
        return [[(participants[0], participants[1]), (participants[2], participants[3])],
                [(participants[0], participants[2]), (participants[1], participants[3])],
                [(participants[0], participants[3]), (participants[1], participants[2])]]
    elif n == 5:
        return [[(participants[0], participants[1]), (participants[2], participants[3])],
                [(participants[4], participants[0]), (participants[1], participants[2])],
                [(participants[3], participants[4]), (participants[0], participants[2])],
                [(participants[1], participants[3]), (participants[2], participants[4])],
                [(participants[0], participants[3]), (participants[1], participants[4])]]
    else:
        if len(participants) % 2 != 0:
            participants.append({"Nom": "BYE", "Club": "-"})
        num_p = len(participants)
        rondes = []
        for i in range(num_p - 1):
            matchs_ronde = []
            for j in range(num_p // 2):
                p1 = participants[j]
                p2 = participants[num_p - 1 - j]
                if p1["Nom"] != "BYE" and p2["Nom"] != "BYE":
                    matchs_ronde.append((p1, p2))
            rondes.append(matchs_ronde)
            participants.insert(1, participants.pop())
        return rondes

def attribuer_categorie_poids_u13(poids_val):
    try:
        p = float(poids_val) if (poids_val is not None and str(poids_val).strip() != '') else 0.0
    except (ValueError, TypeError):
        p = 0.0
    if p <= 30.0: return "30 kg"
    elif p <= 33.0: return "33 kg"
    elif p <= 36.0: return "36 kg"
    elif p <= 39.0: return "39 kg"
    elif p <= 42.0: return "42 kg"
    elif p <= 46.0: return "46 kg"
    elif p <= 50.0: return "50 kg"
    elif p <= 55.0: return "55 kg"
    elif p <= 60.0: return "60 kg"
    else: return "+60 kg"

ORDRE_POIDS_U13 = ["30 kg", "33 kg", "36 kg", "39 kg", "42 kg", "46 kg", "50 kg", "55 kg", "60 kg", "+60 kg"]

def interleave_bracket_slots(list_a, list_b):
    res = []
    i, j = 0, 0
    while i < len(list_a) or j < len(list_b):
        if i < len(list_a):
            res.append(list_a[i])
            i += 1
        if j < len(list_b):
            res.append(list_b[j])
            j += 1
    return res

def repartir_tableau_protection_clubs(participants, nb_byes, nb_prelim):
    """
    Répartit les participants entre exempts (byes) et tour préliminaire (prelim_pts)
    selon la Règle FFLDA de Protection des clubs.
    
    1. Attribue les exemptions prioritairement de manière à ce qu'aucun club n'ait
       plus de 'nb_prelim' lutteurs dans le tour préliminaire (évite obligatoirement
       les affrontements fratricides au tour préliminaire).
    2. Répartit équitablement les exempts entre les clubs qui ont plusieurs inscrits.
    3. Garantit que dans tous les matchs préliminaires (p1, p2), p1['Club'] != p2['Club'].
    """
    clubs = collections.defaultdict(list)
    for p in participants:
        c = p.get('Club', '').strip()
        if not c or c in ['-', 'Comité Non Renseigné', 'Sans club']:
            clubs[f"_indiv_{id(p)}"].append(p)
        else:
            clubs[c].append(p)
            
    sorted_clubs = sorted(clubs.values(), key=len, reverse=True)
    
    byes_list = []
    prelim_list = []
    
    if nb_byes == 0:
        prelim_list = list(participants)
    else:
        club_byes_count = {i: 0 for i in range(len(sorted_clubs))}
        total_byes_assigned = 0
        
        # Pass 1 : Surplus > nb_prelim vers les byes
        for i, c_members in enumerate(sorted_clubs):
            surplus = len(c_members) - nb_prelim
            if surplus > 0:
                take = min(surplus, nb_byes - total_byes_assigned)
                club_byes_count[i] += take
                total_byes_assigned += take
                
        # Pass 2 : Byes attribués en priorité aux clubs multi-membres (>= 2)
        while total_byes_assigned < nb_byes:
            progress = False
            for i, c_members in enumerate(sorted_clubs):
                if total_byes_assigned >= nb_byes:
                    break
                if len(c_members) >= 2 and club_byes_count[i] < len(c_members):
                    club_byes_count[i] += 1
                    total_byes_assigned += 1
                    progress = True
            if not progress:
                # Clubs à 1 membre
                for i, c_members in enumerate(sorted_clubs):
                    if total_byes_assigned >= nb_byes:
                        break
                    if club_byes_count[i] < len(c_members):
                        club_byes_count[i] += 1
                        total_byes_assigned += 1
                        progress = True
            if not progress:
                break
                
        for i, c_members in enumerate(sorted_clubs):
            n_b = club_byes_count[i]
            byes_list.extend(c_members[:n_b])
            prelim_list.extend(c_members[n_b:])
            
    prelim_matches = []
    if nb_prelim > 0:
        club_groups = collections.defaultdict(list)
        for p in prelim_list:
            c = p.get('Club', '').strip()
            club_groups[c].append(p)
        sorted_prelim_clubs = sorted(club_groups.values(), key=len, reverse=True)
        flattened = [p for grp in sorted_prelim_clubs for p in grp]
        
        M = nb_prelim
        half1 = flattened[:M]
        half2 = flattened[M:]
        
        pairs = []
        for i in range(M):
            p1 = half1[i]
            p2 = half2[i]
            if p1.get('Club') and p1.get('Club') == p2.get('Club') and p1.get('Club') not in ['-', '']:
                for j in range(M):
                    if j != i and half2[j].get('Club') != p1.get('Club') and half2[i].get('Club') != half1[j].get('Club'):
                        half2[i], half2[j] = half2[j], half2[i]
                        p2 = half2[i]
                        break
            pairs.append((p1, p2))
        prelim_matches = pairs

    return byes_list, prelim_matches

def generer_competition_u13(age, style_grp, suffixe_niveau, cat_poids, participants, separer_clubs=True):
    n = len(participants)
    if n == 0:
        return None
    elif n == 1:
        nom = f"{age} | {style_grp}{suffixe_niveau} | {cat_poids} (1 seul inscrit)"
        return {
            'nom': nom, 
            'participants': list(participants), 
            'rondes': [], 
            'type_formule': 'seul',
            'cat_poids': cat_poids,
            'style_grp': style_grp
        }
    elif n < 6:
        # Cas 1 : Poule nordique unique (2 à 5 lutteurs)
        nom = f"{age} | {style_grp}{suffixe_niveau} | {cat_poids} (Poule unique)"
        return {
            'nom': nom, 
            'participants': list(participants), 
            'rondes': generer_rondes_fflda(participants), 
            'type_formule': 'poule',
            'cat_poids': cat_poids,
            'style_grp': style_grp
        }
    elif n == 6:
        # Cas 2 : Exactement 6 lutteurs (2 poules de 3 + Phase finale croisée)
        if separer_clubs:
            clubs_vu = {}
            poule_a, poule_b = [], []
            for p in sorted(participants, key=lambda x: x.get('Club', '')):
                c = p.get('Club', '')
                if clubs_vu.get(c, 0) % 2 == 0:
                    if len(poule_a) < 3: poule_a.append(p)
                    else: poule_b.append(p)
                else:
                    if len(poule_b) < 3: poule_b.append(p)
                    else: poule_a.append(p)
                clubs_vu[c] = clubs_vu.get(c, 0) + 1
        else:
            poule_a = [participants[0], participants[2], participants[4]]
            poule_b = [participants[1], participants[3], participants[5]]
        
        while len(poule_a) < 3 and len(poule_b) > 3:
            poule_a.append(poule_b.pop())
        while len(poule_b) < 3 and len(poule_a) > 3:
            poule_b.append(poule_a.pop())

        rondes_a = generer_rondes_fflda(poule_a)
        rondes_b = generer_rondes_fflda(poule_b)
        
        r1 = (rondes_a[0] if len(rondes_a) > 0 else []) + (rondes_b[0] if len(rondes_b) > 0 else [])
        r2 = (rondes_a[1] if len(rondes_a) > 1 else []) + (rondes_b[1] if len(rondes_b) > 1 else [])
        r3 = (rondes_a[2] if len(rondes_a) > 2 else []) + (rondes_b[2] if len(rondes_b) > 2 else [])
        
        # Tour 4 : Demi-finales croisées
        sf1 = (
            {"Nom": f"1er Poule A ({cat_poids})", "Club": "Qualifié A", "Comité": "-"},
            {"Nom": f"2ème Poule B ({cat_poids})", "Club": "Qualifié B", "Comité": "-"}
        )
        sf2 = (
            {"Nom": f"1er Poule B ({cat_poids})", "Club": "Qualifié B", "Comité": "-"},
            {"Nom": f"2ème Poule A ({cat_poids})", "Club": "Qualifié A", "Comité": "-"}
        )
        r4 = [sf1, sf2]
        
        # Tour 5 : Finales (Or/Argent et Bronze unique)
        f_or = (
            {"Nom": f"Vainqueur 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
            {"Nom": f"Vainqueur 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
        )
        f_bronze = (
            {"Nom": f"Perdant 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
            {"Nom": f"Perdant 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
        )
        r5 = [f_or, f_bronze]
        
        nom = f"{age} | {style_grp}{suffixe_niveau} | {cat_poids} (2 Poules + Finales)"
        return {
            'nom': nom, 
            'participants': list(participants), 
            'rondes': [r1, r2, r3, r4, r5], 
            'type_formule': 'poules_croisees',
            'poule_a': poule_a,
            'poule_b': poule_b,
            'cat_poids': cat_poids,
            'style_grp': style_grp
        }
    else:
        # Cas 3 : Plus de 6 lutteurs (Tableau avec repêchage des 1/4 - 2 médailles de bronze)
        rondes = []
        if n == 7:
            # 7 lutteurs : 3 quarts de finale, 1 exempt direct en demi-finale
            if separer_clubs:
                clubs = collections.defaultdict(list)
                for p in participants:
                    c = p.get('Club', '').strip()
                    if not c or c in ['-', 'Comité Non Renseigné', 'Sans club']:
                        clubs[f"_indiv_{id(p)}"].append(p)
                    else:
                        clubs[c].append(p)
                sorted_clubs = sorted(clubs.values(), key=len, reverse=True)
                # Choix de l'exempt : le 1er membre du club le plus représenté (pour protéger ses coéquipiers)
                p_exempt = sorted_clubs[0][0]
                reste = []
                first_skipped = False
                for grp in sorted_clubs:
                    for p in grp:
                        if not first_skipped and p is p_exempt:
                            first_skipped = True
                        else:
                            reste.append(p)
                # Appariement des 6 autres sans fratricide
                club_groups = collections.defaultdict(list)
                for p in reste:
                    c = p.get('Club', '').strip()
                    club_groups[c].append(p)
                sorted_reste_clubs = sorted(club_groups.values(), key=len, reverse=True)
                flattened = [p for grp in sorted_reste_clubs for p in grp]
                h1 = flattened[:3]
                h2 = flattened[3:]
                pairs = []
                for i in range(3):
                    p1 = h1[i]
                    p2 = h2[i]
                    if p1.get('Club') and p1.get('Club') == p2.get('Club') and p1.get('Club') not in ['-', '']:
                        for j in range(3):
                            if j != i and h2[j].get('Club') != p1.get('Club') and h2[i].get('Club') != h1[j].get('Club'):
                                h2[i], h2[j] = h2[j], h2[i]
                                p2 = h2[i]
                                break
                    pairs.append((p1, p2))
                # Séparation Haut / Bas : q1 et q2 sont en Haut, q3 et p_exempt sont en Bas
                # Si un match contient un coéquipier de p_exempt, il doit être en q1 ou q2 (Haut)
                c_ex = p_exempt.get('Club', '').strip()
                if c_ex and c_ex not in ['-', 'Comité Non Renseigné', 'Sans club']:
                    for i in range(3):
                        m = pairs[i]
                        has_teammate = any(p.get('Club') == c_ex for p in m)
                        if has_teammate and i == 2:
                            for target_i in [0, 1]:
                                if not any(p.get('Club') == c_ex for p in pairs[target_i]):
                                    pairs[2], pairs[target_i] = pairs[target_i], pairs[2]
                                    break
                q1, q2, q3 = pairs[0], pairs[1], pairs[2]
            else:
                q1 = (participants[0], participants[1])
                q2 = (participants[2], participants[3])
                q3 = (participants[4], participants[5])
                p_exempt = participants[6]

            rondes.append([q1, q2, q3])
            
            sf1 = (
                {"Nom": f"Vainqueur 1/4 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/4 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            sf2 = (
                {"Nom": f"Vainqueur 1/4 (3) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": p_exempt['Nom'], "Club": p_exempt.get('Club', ''), "Comité": p_exempt.get('Comité', '-')}
            )
            rep1 = (
                {"Nom": f"Perdant 1/4 (1) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/4 (2) [{cat_poids}]", "Club": "Repêché", "Comité": "-"}
            )
            rondes.append([sf1, sf2, rep1])
            
            f_or = (
                {"Nom": f"Vainqueur 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b1 = (
                {"Nom": f"Vainqueur Repêchage 1 [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b2 = (
                {"Nom": f"Perdant 1/4 (3) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            rondes.append([f_or, f_b1, f_b2])

        elif n <= 16:
            # 8 <= n <= 16 : Qualification pour 8 quarts-de-finalistes
            nb_prelim = n - 8
            nb_byes = 16 - n
            
            if separer_clubs:
                byes, prelim_matches = repartir_tableau_protection_clubs(participants, nb_byes, nb_prelim)
                prelim_winners = []
                for i in range(nb_prelim):
                    prelim_winners.append({
                        "Nom": f"Vainqueur Prél. {i+1} [{cat_poids}]",
                        "Club": "Qualifié",
                        "Comité": "-"
                    })
            else:
                byes = participants[:nb_byes]
                prelim_pts = participants[nb_byes:]
                prelim_matches = []
                prelim_winners = []
                for i in range(nb_prelim):
                    p1 = prelim_pts[i * 2]
                    p2 = prelim_pts[i * 2 + 1]
                    prelim_matches.append((p1, p2))
                    prelim_winners.append({
                        "Nom": f"Vainqueur Prél. {i+1} [{cat_poids}]",
                        "Club": "Qualifié",
                        "Comité": "-"
                    })
                
            if prelim_matches:
                rondes.append(prelim_matches)
                
            slots_qf = [None] * 8
            order_prelim_slots = [7, 3, 5, 1, 6, 2, 4, 0]
            for w_idx, slot_idx in enumerate(order_prelim_slots[:nb_prelim]):
                slots_qf[slot_idx] = prelim_winners[w_idx]
                
            if separer_clubs:
                order_upper = [s for s in [0, 2, 1, 3] if slots_qf[s] is None]
                order_lower = [s for s in [4, 6, 5, 7] if slots_qf[s] is None]
                interleaved_bye_slots = interleave_bracket_slots(order_upper, order_lower)
                for b_idx, s_idx in enumerate(interleaved_bye_slots):
                    if b_idx < len(byes):
                        slots_qf[s_idx] = byes[b_idx]
            else:
                bye_idx = 0
                for s_idx in range(8):
                    if slots_qf[s_idx] is None:
                        slots_qf[s_idx] = byes[bye_idx]
                        bye_idx += 1
                        
            qf_matches = [
                (slots_qf[0], slots_qf[1]),
                (slots_qf[2], slots_qf[3]),
                (slots_qf[4], slots_qf[5]),
                (slots_qf[6], slots_qf[7])
            ]
            rondes.append(qf_matches)
            
            sf1 = (
                {"Nom": f"Vainqueur 1/4 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/4 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            sf2 = (
                {"Nom": f"Vainqueur 1/4 (3) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/4 (4) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            rep1 = (
                {"Nom": f"Perdant 1/4 (1) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/4 (2) [{cat_poids}]", "Club": "Repêché", "Comité": "-"}
            )
            rep2 = (
                {"Nom": f"Perdant 1/4 (3) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/4 (4) [{cat_poids}]", "Club": "Repêché", "Comité": "-"}
            )
            rondes.append([sf1, sf2, rep1, rep2])
            
            f_or = (
                {"Nom": f"Vainqueur 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b1 = (
                {"Nom": f"Vainqueur Repêchage 1 [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b2 = (
                {"Nom": f"Vainqueur Repêchage 2 [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            rondes.append([f_or, f_b1, f_b2])

        elif n <= 32:
            # 17 <= n <= 32 (Tableau de 32 : 1/16, 1/8, 1/4, 1/2, Finales & Repêchages)
            nb_prelim = n - 16
            nb_byes = max(0, 32 - n)
            
            if separer_clubs:
                byes_16, prelim_matches = repartir_tableau_protection_clubs(participants, nb_byes, nb_prelim)
                prelim_winners = []
                for i in range(nb_prelim):
                    prelim_winners.append({
                        "Nom": f"Vainqueur 1/16 ({i+1}) [{cat_poids}]",
                        "Club": "Qualifié",
                        "Comité": "-"
                    })
                rondes.append(prelim_matches)
                
                slots_16 = [None] * 16
                order_prelim_slots_16 = [15, 7, 11, 3, 13, 5, 9, 1, 14, 6, 10, 2, 12, 4, 8, 0]
                for w_idx, slot_idx in enumerate(order_prelim_slots_16[:nb_prelim]):
                    slots_16[slot_idx] = prelim_winners[w_idx]
                    
                order_upper = [s for s in [0, 4, 2, 6, 1, 5, 3, 7] if slots_16[s] is None]
                order_lower = [s for s in [8, 12, 10, 14, 9, 13, 11, 15] if slots_16[s] is None]
                interleaved_bye_slots = interleave_bracket_slots(order_upper, order_lower)
                for b_idx, s_idx in enumerate(interleaved_bye_slots):
                    if b_idx < len(byes_16):
                        slots_16[s_idx] = byes_16[b_idx]
                        
                for s_idx in range(16):
                    if slots_16[s_idx] is None:
                        slots_16[s_idx] = {"Nom": f"Qualifié 1/8 ({s_idx+1})", "Club": "Qualifié", "Comité": "-"}
            else:
                byes_16 = participants[:nb_byes]
                prelim_pts = participants[nb_byes:]
                prelim_matches = []
                prelim_winners = []
                for i in range(nb_prelim):
                    p1 = prelim_pts[i * 2] if i * 2 < len(prelim_pts) else {"Nom": f"Lutteur {i*2+1}", "Club": "-"}
                    p2 = prelim_pts[i * 2 + 1] if i * 2 + 1 < len(prelim_pts) else {"Nom": f"Lutteur {i*2+2}", "Club": "-"}
                    prelim_matches.append((p1, p2))
                    prelim_winners.append({
                        "Nom": f"Vainqueur 1/16 ({i+1}) [{cat_poids}]",
                        "Club": "Qualifié",
                        "Comité": "-"
                    })
                rondes.append(prelim_matches)
                slots_16 = (byes_16 + prelim_winners)[:16]
                while len(slots_16) < 16:
                    slots_16.append({"Nom": f"Qualifié 1/8 ({len(slots_16)+1})", "Club": "Qualifié", "Comité": "-"})
            
            matches_18 = []
            winners_18 = []
            for i in range(8):
                p1 = slots_16[i * 2]
                p2 = slots_16[i * 2 + 1]
                matches_18.append((p1, p2))
                winners_18.append({
                    "Nom": f"Vainqueur 1/8 ({i+1}) [{cat_poids}]",
                    "Club": "Qualifié",
                    "Comité": "-"
                })
            rondes.append(matches_18)
            
            qf_matches = [
                (winners_18[0], winners_18[1]),
                (winners_18[2], winners_18[3]),
                (winners_18[4], winners_18[5]),
                (winners_18[6], winners_18[7])
            ]
            rondes.append(qf_matches)
            
            sf1 = (
                {"Nom": f"Vainqueur 1/4 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/4 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            sf2 = (
                {"Nom": f"Vainqueur 1/4 (3) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/4 (4) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            rep1 = (
                {"Nom": f"Perdant 1/4 (1) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/4 (2) [{cat_poids}]", "Club": "Repêché", "Comité": "-"}
            )
            rep2 = (
                {"Nom": f"Perdant 1/4 (3) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/4 (4) [{cat_poids}]", "Club": "Repêché", "Comité": "-"}
            )
            rondes.append([sf1, sf2, rep1, rep2])
            
            f_or = (
                {"Nom": f"Vainqueur 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b1 = (
                {"Nom": f"Vainqueur Repêchage 1 [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b2 = (
                {"Nom": f"Vainqueur Repêchage 2 [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            rondes.append([f_or, f_b1, f_b2])

        else:
            # 33 <= n <= 64 (Tableau de 64 : 1/32, 1/16, 1/8, 1/4, 1/2, Finales & Repêchages)
            nb_prelim = n - 32
            nb_byes = max(0, 64 - n)
            
            if separer_clubs:
                byes_32, prelim_matches = repartir_tableau_protection_clubs(participants, nb_byes, nb_prelim)
                prelim_winners = []
                for i in range(nb_prelim):
                    prelim_winners.append({
                        "Nom": f"Vainqueur 1/32 ({i+1}) [{cat_poids}]",
                        "Club": "Qualifié",
                        "Comité": "-"
                    })
                rondes.append(prelim_matches)
                
                slots_32 = [None] * 32
                order_prelim_slots_32 = [31, 15, 23, 7, 27, 11, 19, 3, 29, 13, 21, 5, 25, 9, 17, 1, 30, 14, 22, 6, 26, 10, 18, 2, 28, 12, 20, 4, 24, 8, 16, 0]
                for w_idx, slot_idx in enumerate(order_prelim_slots_32[:nb_prelim]):
                    slots_32[slot_idx] = prelim_winners[w_idx]
                    
                order_upper = [s for s in [0, 8, 4, 12, 2, 10, 6, 14, 1, 9, 5, 13, 3, 11, 7, 15] if slots_32[s] is None]
                order_lower = [s for s in [16, 24, 20, 28, 18, 26, 22, 30, 17, 25, 21, 29, 19, 27, 23, 31] if slots_32[s] is None]
                interleaved_bye_slots = interleave_bracket_slots(order_upper, order_lower)
                for b_idx, s_idx in enumerate(interleaved_bye_slots):
                    if b_idx < len(byes_32):
                        slots_32[s_idx] = byes_32[b_idx]
                        
                for s_idx in range(32):
                    if slots_32[s_idx] is None:
                        slots_32[s_idx] = {"Nom": f"Qualifié 1/16 ({s_idx+1})", "Club": "Qualifié", "Comité": "-"}
            else:
                byes_32 = participants[:nb_byes]
                prelim_pts = participants[nb_byes:]
                prelim_matches = []
                prelim_winners = []
                for i in range(nb_prelim):
                    p1 = prelim_pts[i * 2] if i * 2 < len(prelim_pts) else {"Nom": f"Lutteur {i*2+1}", "Club": "-"}
                    p2 = prelim_pts[i * 2 + 1] if i * 2 + 1 < len(prelim_pts) else {"Nom": f"Lutteur {i*2+2}", "Club": "-"}
                    prelim_matches.append((p1, p2))
                    prelim_winners.append({
                        "Nom": f"Vainqueur 1/32 ({i+1}) [{cat_poids}]",
                        "Club": "Qualifié",
                        "Comité": "-"
                    })
                rondes.append(prelim_matches)
                slots_32 = (byes_32 + prelim_winners)[:32]
                while len(slots_32) < 32:
                    slots_32.append({"Nom": f"Qualifié 1/16 ({len(slots_32)+1})", "Club": "Qualifié", "Comité": "-"})
            
            matches_116 = []
            winners_116 = []
            for i in range(16):
                p1 = slots_32[i * 2]
                p2 = slots_32[i * 2 + 1]
                matches_116.append((p1, p2))
                winners_116.append({
                    "Nom": f"Vainqueur 1/16 ({i+1}) [{cat_poids}]",
                    "Club": "Qualifié",
                    "Comité": "-"
                })
            rondes.append(matches_116)
            
            matches_18 = []
            winners_18 = []
            for i in range(8):
                p1 = winners_116[i * 2]
                p2 = winners_116[i * 2 + 1]
                matches_18.append((p1, p2))
                winners_18.append({
                    "Nom": f"Vainqueur 1/8 ({i+1}) [{cat_poids}]",
                    "Club": "Qualifié",
                    "Comité": "-"
                })
            rondes.append(matches_18)
            
            qf_matches = [
                (winners_18[0], winners_18[1]),
                (winners_18[2], winners_18[3]),
                (winners_18[4], winners_18[5]),
                (winners_18[6], winners_18[7])
            ]
            rondes.append(qf_matches)
            
            sf1 = (
                {"Nom": f"Vainqueur 1/4 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/4 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            sf2 = (
                {"Nom": f"Vainqueur 1/4 (3) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/4 (4) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            rep1 = (
                {"Nom": f"Perdant 1/4 (1) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/4 (2) [{cat_poids}]", "Club": "Repêché", "Comité": "-"}
            )
            rep2 = (
                {"Nom": f"Perdant 1/4 (3) [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/4 (4) [{cat_poids}]", "Club": "Repêché", "Comité": "-"}
            )
            rondes.append([sf1, sf2, rep1, rep2])
            
            f_or = (
                {"Nom": f"Vainqueur 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"},
                {"Nom": f"Vainqueur 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b1 = (
                {"Nom": f"Vainqueur Repêchage 1 [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (2) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            f_b2 = (
                {"Nom": f"Vainqueur Repêchage 2 [{cat_poids}]", "Club": "Repêché", "Comité": "-"},
                {"Nom": f"Perdant 1/2 (1) [{cat_poids}]", "Club": "Qualifié", "Comité": "-"}
            )
            rondes.append([f_or, f_b1, f_b2])

        nom = f"{age} | {style_grp}{suffixe_niveau} | {cat_poids} (Tableau élimination & repêchages)"
        res = {
            'nom': nom,
            'participants': list(participants),
            'rondes': rondes,
            'type_formule': 'tableau',
            'cat_poids': cat_poids,
            'style_grp': style_grp
        }
        if n == 7:
            res['p_exempt'] = p_exempt
        return res

# --- GÉNÉRATEUR VISUEL DE TABLEAU À ÉLIMINATION DIRECTE & REPÊCHAGES U13 (DE GAUCHE À DROITE) ---
def make_bracket_card(p1, p2, title='', badge=''):
    nom1 = p1.get('Nom', 'Lutteur 1')
    c1 = p1.get('Club', '')
    club1_str = f" ({c1})" if c1 and c1 not in ['-', 'Comité Non Renseigné', ''] else ''
    
    nom2 = p2.get('Nom', 'Lutteur 2')
    c2 = p2.get('Club', '')
    club2_str = f" ({c2})" if c2 and c2 not in ['-', 'Comité Non Renseigné', ''] else ''

    badge_html = f'<span style="background: #FEF3C7; color: #92400E; border: 1px solid #FCD34D; padding: 1px 5px; border-radius: 4px; font-size: 9px; font-weight: 700;">{badge}</span>' if badge else ''

    is_bye1 = any(k in str(nom1).lower() for k in ['bye', 'exempt'])
    is_bye2 = any(k in str(nom2).lower() for k in ['bye', 'exempt'])

    dot1 = '#94A3B8' if is_bye1 else '#EF4135'
    dot2 = '#94A3B8' if is_bye2 else '#0055A4'

    return f'''
    <div style="background: #FFFFFF; border: 1.5px solid #CBD5E1; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.05); width: 230px; margin: 6px 0; overflow: hidden; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; text-align: left;">
        <div style="background: #F8FAFC; border-bottom: 1px solid #E2E8F0; padding: 4px 8px; font-size: 10px; font-weight: 700; color: #475569; display: flex; justify-content: space-between; align-items: center;">
            <span>{title}</span>
            {badge_html}
        </div>
        <div style="padding: 5px 8px; display: flex; align-items: center; border-bottom: 1px solid #F1F5F9; background: {'#F8FAFC' if is_bye1 else '#FFFFFF'};">
            <span style="width: 9px; height: 9px; border-radius: 50%; background: {dot1}; display: inline-block; margin-right: 6px; flex-shrink: 0;"></span>
            <div style="flex: 1; overflow: hidden;">
                <div style="font-size: 11px; font-weight: 700; color: {'#64748B' if is_bye1 else '#1E293B'}; white-space: nowrap; overflow: hidden; text-overflow: ellipsis;">{nom1}</div>
                <div style="font-size: 9px; color: #94A3B8; white-space: nowrap; overflow: hidden; text-overflow: ellipsis;">{club1_str}</div>
            </div>
            <div style="font-size: 9px; font-weight: 700; background: #F1F5F9; border: 1px solid #CBD5E1; border-radius: 3px; padding: 1px 4px; color: #475569;">[ &nbsp; ]</div>
        </div>
        <div style="padding: 5px 8px; display: flex; align-items: center; background: {'#F8FAFC' if is_bye2 else '#FFFFFF'};">
            <span style="width: 9px; height: 9px; border-radius: 50%; background: {dot2}; display: inline-block; margin-right: 6px; flex-shrink: 0;"></span>
            <div style="flex: 1; overflow: hidden;">
                <div style="font-size: 11px; font-weight: 700; color: {'#64748B' if is_bye2 else '#1E293B'}; white-space: nowrap; overflow: hidden; text-overflow: ellipsis;">{nom2}</div>
                <div style="font-size: 9px; color: #94A3B8; white-space: nowrap; overflow: hidden; text-overflow: ellipsis;">{club2_str}</div>
            </div>
            <div style="font-size: 9px; font-weight: 700; background: #F1F5F9; border: 1px solid #CBD5E1; border-radius: 3px; padding: 1px 4px; color: #475569;">[ &nbsp; ]</div>
        </div>
    </div>
    '''

def render_svg_connectors(count, height, width=40):
    lines = []
    pairs = count // 2
    step = height / count
    for p in range(pairs):
        y_top = (p * 2 + 0.5) * step
        y_bot = (p * 2 + 1.5) * step
        y_mid = (p + 0.5) * (height / pairs)
        path = f"M 0 {y_top:.1f} H {width//2} V {y_bot:.1f} H 0 M {width//2} {y_mid:.1f} H {width}"
        lines.append(f'<path d="{path}" fill="none" stroke="#94A3B8" stroke-width="2" />')
    return f'<svg width="{width}" height="{height}" style="flex-shrink: 0; display: block;">{" ".join(lines)}</svg>'

def generer_arbre_tableau_html(p_obj):
    nom_poule = p_obj.get('nom', 'U13')
    rondes = p_obj.get('rondes', [])
    participants = p_obj.get('participants', [])
    type_formule = p_obj.get('type_formule', '')
    n = len(participants)

    if type_formule == 'poules_croisees':
        sf1, sf2 = rondes[3][0], rondes[3][1]
        f_or, f_b = rondes[4][0], rondes[4][1]
        
        c_sf1 = make_bracket_card(sf1[0], sf1[1], 'DEMI-FINALE 1', '1er A vs 2ème B')
        c_sf2 = make_bracket_card(sf2[0], sf2[1], 'DEMI-FINALE 2', '1er B vs 2ème A')
        c_for = make_bracket_card(f_or[0], f_or[1], 'FINALE OR / ARGENT', '🥇 Or / 🥈 Argent')
        c_fb = make_bracket_card(f_b[0], f_b[1], 'FINALE BRONZE (3-4)', '🥉 Bronze unique')

        h_col = 220
        svg_conn = f'<svg width="40" height="{h_col}" style="display: block; flex-shrink: 0;"><path d="M 0 55 H 20 V 165 H 0 M 20 55 H 40 M 20 165 H 40" fill="none" stroke="#94A3B8" stroke-width="2" /></svg>'

        return f'''
        <div style="background: #F8FAFC; border: 1.5px solid #E2E8F0; border-radius: 12px; padding: 18px; margin: 15px 0; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;">
            <div style="display: flex; align-items: center; justify-content: space-between; margin-bottom: 16px;">
                <span style="background: #0055A4; color: white; padding: 5px 14px; border-radius: 16px; font-size: 12px; font-weight: 800; text-transform: uppercase;">
                    🏆 Tableau Final Croisé U13 (6 Lutteurs)
                </span>
                <span style="font-size: 11px; color: #64748B; font-weight: 600;">Orientation : Gauche ➔ Droite (Demi-Finales ➔ Finales)</span>
            </div>
            
            <div style="overflow-x: auto; padding-bottom: 10px;">
                <div style="display: inline-flex; flex-direction: row; align-items: center; gap: 0;">
                    <!-- Demi-Finales -->
                    <div style="display: flex; flex-direction: column; width: 240px;">
                        <div style="background: #0055A4; color: white; padding: 6px; border-radius: 6px; font-size: 11px; font-weight: 700; text-align: center; margin-bottom: 10px;">
                            DEMI-FINALES CROISÉES
                        </div>
                        <div style="display: flex; flex-direction: column; justify-content: space-around; height: {h_col}px;">
                            {c_sf1}
                            {c_sf2}
                        </div>
                    </div>

                    {svg_conn}

                    <!-- Finales -->
                    <div style="display: flex; flex-direction: column; width: 240px;">
                        <div style="background: #0F172A; color: white; padding: 6px; border-radius: 6px; font-size: 11px; font-weight: 700; text-align: center; margin-bottom: 10px;">
                            FINALES
                        </div>
                        <div style="display: flex; flex-direction: column; justify-content: space-around; height: {h_col}px;">
                            {c_for}
                            {c_fb}
                        </div>
                    </div>

                    <!-- Podiums -->
                    <div style="margin-left: 15px; display: flex; flex-direction: column; justify-content: space-around; height: {h_col}px;">
                        <div style="background: #FEF3C7; border: 1.5px solid #F59E0B; border-radius: 8px; padding: 8px 12px; font-size: 11px;">
                            <div style="font-weight: 800; color: #B45309;">🥇 OR : Vainqueur Finale</div>
                            <div style="font-weight: 800; color: #475569; margin-top: 4px;">🥈 ARGENT : Perdant Finale</div>
                        </div>
                        <div style="background: #EFF6FF; border: 1.5px solid #3B82F6; border-radius: 8px; padding: 8px 12px; font-size: 11px;">
                            <div style="font-weight: 800; color: #1D4ED8;">🥉 BRONZE (Unique) : Vainqueur 3-4</div>
                        </div>
                    </div>
                </div>
            </div>
        </div>
        '''

    elif type_formule == 'tableau':
        f_or = rondes[-1][0]
        f_b1 = rondes[-1][1] if len(rondes[-1]) > 1 else None
        f_b2 = rondes[-1][2] if len(rondes[-1]) > 2 else None

        sf1 = rondes[-2][0]
        sf2 = rondes[-2][1]
        rep1 = rondes[-2][2] if len(rondes[-2]) > 2 else None
        rep2 = rondes[-2][3] if len(rondes[-2]) > 3 else None

        main_columns = []
        if len(rondes) == 3:
            if n == 7:
                q1, q2, q3 = rondes[0][0], rondes[0][1], rondes[0][2]
                p_ex = p_obj.get('p_exempt', participants[6])
                q_matches = [
                    (q1[0], q1[1], '1/4 DE FINALE 1'),
                    (q2[0], q2[1], '1/4 DE FINALE 2'),
                    (q3[0], q3[1], '1/4 DE FINALE 3'),
                    (p_ex, {'Nom': 'EXEMPT (BYE)', 'Club': '-'}, '1/4 DE FINALE 4 (Exempt)')
                ]
            else:
                q_matches = [(m[0], m[1], f'1/4 DE FINALE {i+1}') for i, m in enumerate(rondes[0])]
            main_columns.append(('QUARTS DE FINALE', q_matches))

        elif len(rondes) == 4:
            prelim_m = [(m[0], m[1], f'PRÉLIMINAIRE {i+1}') for i, m in enumerate(rondes[0])]
            q_matches = [(m[0], m[1], f'1/4 DE FINALE {i+1}') for i, m in enumerate(rondes[1])]
            main_columns.append(('TOUR PRÉLIMINAIRE', prelim_m))
            main_columns.append(('QUARTS DE FINALE', q_matches))

        elif len(rondes) >= 6:
            p32 = [(m[0], m[1], f'1/32 DE FINALE {i+1}') for i, m in enumerate(rondes[0])]
            p16 = [(m[0], m[1], f'1/16 DE FINALE {i+1}') for i, m in enumerate(rondes[1])]
            p18 = [(m[0], m[1], f'1/8 DE FINALE {i+1}') for i, m in enumerate(rondes[2])]
            q_matches = [(m[0], m[1], f'1/4 DE FINALE {i+1}') for i, m in enumerate(rondes[3])]
            main_columns.append(('1/32 DE FINALE', p32))
            main_columns.append(('1/16 DE FINALE', p16))
            main_columns.append(('1/8 DE FINALE', p18))
            main_columns.append(('QUARTS DE FINALE', q_matches))

        elif len(rondes) == 5:
            p16 = [(m[0], m[1], f'1/16 DE FINALE {i+1}') for i, m in enumerate(rondes[0])]
            p18 = [(m[0], m[1], f'1/8 DE FINALE {i+1}') for i, m in enumerate(rondes[1])]
            q_matches = [(m[0], m[1], f'1/4 DE FINALE {i+1}') for i, m in enumerate(rondes[2])]
            main_columns.append(('1/16 DE FINALE', p16))
            main_columns.append(('1/8 DE FINALE', p18))
            main_columns.append(('QUARTS DE FINALE', q_matches))

        main_columns.append(('DEMI-FINALES', [(sf1[0], sf1[1], 'DEMI-FINALE 1'), (sf2[0], sf2[1], 'DEMI-FINALE 2')]))
        main_columns.append(('FINALE (OR / ARGENT)', [(f_or[0], f_or[1], 'GRANDE FINALE')]))

        max_matches_in_col = max(len(col[1]) for col in main_columns)
        base_h = max(max_matches_in_col * 110, 440)

        main_cols_html = []
        for c_idx, (col_title, matches_list) in enumerate(main_columns):
            cards_html = []
            for m in matches_list:
                cards_html.append(make_bracket_card(m[0], m[1], m[2]))
            
            bg_h = '#0F172A' if 'FINALE' in col_title and 'DEMI' not in col_title else '#0055A4'

            col_div = f'''
            <div style="display: inline-flex; flex-direction: column; width: 240px; margin: 0 5px;">
                <div style="background: {bg_h}; color: white; padding: 6px; border-radius: 6px; font-size: 11px; font-weight: 700; text-align: center; margin-bottom: 10px;">
                    {col_title}
                </div>
                <div style="display: flex; flex-direction: column; justify-content: space-around; height: {base_h}px;">
                    {''.join(cards_html)}
                </div>
            </div>
            '''
            main_cols_html.append(col_div)

            if c_idx < len(main_columns) - 1:
                next_count = len(main_columns[c_idx + 1][1])
                curr_count = len(matches_list)
                if curr_count == next_count * 2:
                    svg_c = render_svg_connectors(curr_count, base_h, width=40)
                else:
                    lines = [f'<line x1="0" y1="{base_h * (i + 0.5) / curr_count:.1f}" x2="40" y2="{base_h * (i + 0.5) / curr_count:.1f}" stroke="#94A3B8" stroke-width="2" />' for i in range(curr_count)]
                    svg_c = f'<svg width="40" height="{base_h}" style="flex-shrink: 0; display: block;">{" ".join(lines)}</svg>'
                main_cols_html.append(svg_c)

        podium_main = f'''
        <div style="display: inline-flex; align-items: center; margin-left: 10px;">
            <div style="background: #FFFDF5; border: 1.5px solid #F59E0B; border-radius: 8px; padding: 10px 14px; box-shadow: 0 2px 6px rgba(245,158,11,0.15); font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;">
                <div style="font-size: 12px; font-weight: 800; color: #B45309; margin-bottom: 6px;">🥇 CHAMPION (OR)</div>
                <div style="font-size: 11px; color: #1E293B;">Vainqueur Grande Finale</div>
                <div style="border-top: 1px solid #FDE68A; margin: 6px 0;"></div>
                <div style="font-size: 12px; font-weight: 800; color: #64748B; margin-bottom: 6px;">🥈 VICE-CHAMPION (ARGENT)</div>
                <div style="font-size: 11px; color: #1E293B;">Finaliste Grande Finale</div>
            </div>
        </div>
        '''
        main_cols_html.append(podium_main)

        rep_cards = []
        if rep1: rep_cards.append(make_bracket_card(rep1[0], rep1[1], 'REPÊCHAGE 1 (1/4)', 'Perdants QF 1-2'))
        if rep2: rep_cards.append(make_bracket_card(rep2[0], rep2[1], 'REPÊCHAGE 2 (1/4)', 'Perdants QF 3-4'))
        elif n == 7: rep_cards.append('<div style="background: #F1F5F9; border: 1.5px dashed #CBD5E1; border-radius: 8px; padding: 12px; font-size: 11px; color: #64748B; text-align: center; width: 230px;">Exempt de 1er repêchage (Avance direct en Finale Bronze 2)</div>')

        bronze_cards = []
        if f_b1: bronze_cards.append(make_bracket_card(f_b1[0], f_b1[1], 'FINALE BRONZE 1', '🥉 Médaille de Bronze 1'))
        if f_b2: bronze_cards.append(make_bracket_card(f_b2[0], f_b2[1], 'FINALE BRONZE 2', '🥉 Médaille de Bronze 2'))

        h_rep = 220
        svg_rep = f'''
        <svg width="40" height="{h_rep}" style="flex-shrink: 0; display: block;">
            <line x1="0" y1="55" x2="40" y2="55" stroke="#94A3B8" stroke-width="2" />
            <line x1="0" y1="165" x2="40" y2="165" stroke="#94A3B8" stroke-width="2" />
        </svg>
        '''

        rep_html = f'''
        <div style="margin-top: 24px; padding-top: 18px; border-top: 2px dashed #CBD5E1;">
            <div style="display: flex; align-items: center; justify-content: space-between; margin-bottom: 14px;">
                <span style="background: #0284C7; color: white; padding: 5px 14px; border-radius: 16px; font-size: 12px; font-weight: 800; text-transform: uppercase;">
                    🔄 Repêchages & Attribution du Bronze (2 Troisièmes Places)
                </span>
                <span style="font-size: 11px; color: #64748B; font-weight: 600;">Perdants des 1/4 ➔ Repêchages ➔ Finales Bronze contre Perdants des 1/2</span>
            </div>

            <div style="overflow-x: auto; padding-bottom: 10px;">
                <div style="display: inline-flex; flex-direction: row; align-items: center; gap: 0;">
                    <div style="display: flex; flex-direction: column; width: 240px; margin: 0 5px;">
                        <div style="background: #0284C7; color: white; padding: 6px; border-radius: 6px; font-size: 11px; font-weight: 700; text-align: center; margin-bottom: 10px;">
                            REPÊCHAGES (1/4)
                        </div>
                        <div style="display: flex; flex-direction: column; justify-content: space-around; height: {h_rep}px;">
                            {''.join(rep_cards)}
                        </div>
                    </div>

                    {svg_rep}

                    <div style="display: flex; flex-direction: column; width: 240px; margin: 0 5px;">
                        <div style="background: #B45309; color: white; padding: 6px; border-radius: 6px; font-size: 11px; font-weight: 700; text-align: center; margin-bottom: 10px;">
                            MATCHS POUR LE BRONZE
                        </div>
                        <div style="display: flex; flex-direction: column; justify-content: space-around; height: {h_rep}px;">
                            {''.join(bronze_cards)}
                        </div>
                    </div>

                    <div style="margin-left: 15px; display: flex; flex-direction: column; justify-content: space-around; height: {h_rep}px;">
                        <div style="background: #EFF6FF; border: 1.5px solid #3B82F6; border-radius: 8px; padding: 8px 12px; font-size: 11px;">
                            <div style="font-weight: 800; color: #1D4ED8;">🥉 3ème PLACE (Bronze 1)</div>
                            <div style="color: #1E293B; margin-top: 2px;">Vainqueur Finale Bronze 1</div>
                        </div>
                        <div style="background: #EFF6FF; border: 1.5px solid #3B82F6; border-radius: 8px; padding: 8px 12px; font-size: 11px;">
                            <div style="font-weight: 800; color: #1D4ED8;">🥉 3ème PLACE (Bronze 2)</div>
                            <div style="color: #1E293B; margin-top: 2px;">Vainqueur Finale Bronze 2</div>
                        </div>
                    </div>
                </div>
            </div>
        </div>
        '''

        return f'''
        <div style="background: #F8FAFC; border: 1.5px solid #E2E8F0; border-radius: 12px; padding: 18px; margin: 15px 0; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;">
            <div style="display: flex; align-items: center; justify-content: space-between; margin-bottom: 16px;">
                <span style="background: #0055A4; color: white; padding: 5px 14px; border-radius: 16px; font-size: 12px; font-weight: 800; text-transform: uppercase;">
                    🏆 Tableau Principal à Élimination Directe — {nom_poule}
                </span>
                <span style="font-size: 11px; color: #64748B; font-weight: 600;">Orientation : Gauche ➔ Droite (Éliminatoires ➔ Quarts ➔ Demi-Finales ➔ Finale)</span>
            </div>

            <div style="overflow-x: auto; padding-bottom: 10px;">
                <div style="display: inline-flex; flex-direction: row; align-items: center; gap: 0;">
                    {''.join(main_cols_html)}
                </div>
            </div>

            {rep_html}
        </div>
        '''
    return ''

def generer_document_bracket_imprimable(nom_poule, nom_comp, bracket_html):
    return f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Tableau U13 - {nom_poule} - {nom_comp}</title>
    <style>
        @page {{
            size: landscape;
            margin: 8mm;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
            margin: 0;
            padding: 10px;
            background: white;
            color: #1E293B;
        }}
        .print-btn {{
            background-color: #0055A4;
            color: white;
            border: none;
            padding: 8px 16px;
            font-size: 13px;
            font-weight: bold;
            border-radius: 6px;
            cursor: pointer;
            margin-bottom: 12px;
        }}
        @media print {{
            .print-btn {{
                display: none !important;
            }}
            body {{
                padding: 0;
            }}
        }}
    </style>
</head>
<body>
    <button class="print-btn" onclick="window.print()">🖨️ Imprimer ce Tableau (A4 Paysage)</button>
    <div style="margin-bottom: 10px;">
        <h2 style="color: #0055A4; margin: 0 0 4px 0; font-size: 20px;">🏆 {nom_comp.upper()}</h2>
        <div style="font-size: 12px; color: #64748B; font-weight: bold;">FFLDA — TABLEAU OFFICIEL À ÉLIMINATION DIRECTE & REPÊCHAGES</div>
    </div>
    {bracket_html}
</body>
</html>"""

def generer_document_poule_imprimable(nom_poule, nom_comp, participants, rondes):
    nb_tours = len(rondes) if rondes else 0
    
    lignes_html = []
    for idx, p in enumerate(participants, 1):
        bg = "#F8FAFC" if idx % 2 == 0 else "#FFFFFF"
        nom = p.get('Nom', '')
        club = p.get('Club', '')
        comite = p.get('Comité', '')
        poids = formater_poids(p.get('Poids', ''))
        
        tours_tds = ''.join(['<td style="border: 1px solid #CBD5E1; padding: 6px; text-align: center;"></td>' for _ in range(nb_tours)])
        
        lignes_html.append(f"""
        <tr style="background: {bg};">
            <td style="border: 1px solid #CBD5E1; padding: 6px; text-align: center; font-weight: bold; color: #0055A4;"></td>
            <td style="border: 1px solid #CBD5E1; padding: 6px; text-align: center; font-weight: bold;">{idx}</td>
            <td style="border: 1px solid #CBD5E1; padding: 6px; font-weight: 600;">{nom}</td>
            <td style="border: 1px solid #CBD5E1; padding: 6px;">{club}</td>
            <td style="border: 1px solid #CBD5E1; padding: 6px;">{comite}</td>
            {tours_tds}
            <td style="border: 1px solid #CBD5E1; padding: 6px; text-align: center; font-weight: bold;"></td>
            <td style="border: 1px solid #CBD5E1; padding: 6px; text-align: center;"></td>
            <td style="border: 1px solid #CBD5E1; padding: 6px; text-align: center;">{poids}</td>
        </tr>
        """)

    headers_tours = ''.join([f'<th style="border: 1px solid #CBD5E1; padding: 8px; font-size: 11px;">Tour {t}</th>' for t in range(1, nb_tours + 1)])
    
    n_p = len(participants)
    podium_html = f"""
    <div style="display: flex; gap: 12px; margin-top: 15px; margin-bottom: 20px;">
        <div style="flex: 1; background: #FEF3C7; border: 1.5px solid #F59E0B; border-radius: 8px; padding: 10px; text-align: center;">
            <div style="font-weight: 800; font-size: 13px; color: #B45309;">🥇 1ère PLACE (OR)</div>
            <div style="font-size: 12px; font-weight: 600; color: #78350F; margin-top: 4px;">En attente</div>
        </div>
        <div style="flex: 1; background: #F1F5F9; border: 1.5px solid #94A3B8; border-radius: 8px; padding: 10px; text-align: center;">
            <div style="font-weight: 800; font-size: 13px; color: #475569;">🥈 2ème PLACE (ARGENT)</div>
            <div style="font-size: 12px; font-weight: 600; color: #1E293B; margin-top: 4px;">En attente</div>
        </div>
    """
    if n_p >= 3:
        podium_html += """
        <div style="flex: 1; background: #FFEDD5; border: 1.5px solid #F97316; border-radius: 8px; padding: 10px; text-align: center;">
            <div style="font-weight: 800; font-size: 13px; color: #9A3412;">🥉 3ème PLACE (BRONZE)</div>
            <div style="font-size: 12px; font-weight: 600; color: #7C2D12; margin-top: 4px;">En attente</div>
        </div>
        """
    podium_html += "</div>"
    
    matchs_html = []
    m_count = 1
    for tour_idx, ronde in enumerate(rondes, 1):
        tour_cards = []
        for m in ronde:
            p1, p2 = m[0], m[1]
            nom1 = p1.get('Nom', '')
            club1 = f" ({p1.get('Club', '')})" if p1.get('Club') and p1.get('Club') != '-' else ''
            nom2 = p2.get('Nom', '')
            club2 = f" ({p2.get('Club', '')})" if p2.get('Club') and p2.get('Club') != '-' else ''
            
            tour_cards.append(f"""
            <div style="background: white; border: 1px solid #CBD5E1; border-radius: 6px; margin-bottom: 8px; overflow: hidden; width: 100%;">
                <div style="background: #334155; color: white; font-size: 10px; font-weight: bold; padding: 3px 8px; text-align: center;">
                    COMBAT N°{m_count}
                </div>
                <div style="display: flex; justify-content: space-between; align-items: center; background: #FEF2F2; padding: 5px 8px; border-bottom: 1px solid #CBD5E1;">
                    <span style="font-size: 11px; font-weight: 700; color: #991B1B;">🔴 {nom1}{club1}</span>
                    <span style="border: 1px solid #CBD5E1; background: white; padding: 2px 8px; font-weight: bold; border-radius: 4px; font-size: 11px;">&nbsp;&nbsp;&nbsp;</span>
                </div>
                <div style="display: flex; justify-content: space-between; align-items: center; background: #EFF6FF; padding: 5px 8px;">
                    <span style="font-size: 11px; font-weight: 700; color: #1E40AF;">🔵 {nom2}{club2}</span>
                    <span style="border: 1px solid #CBD5E1; background: white; padding: 2px 8px; font-weight: bold; border-radius: 4px; font-size: 11px;">&nbsp;&nbsp;&nbsp;</span>
                </div>
            </div>
            """)
            m_count += 1
            
        matchs_html.append(f"""
        <div style="margin-bottom: 14px;">
            <div style="background: #0055A4; color: white; font-size: 11px; font-weight: bold; padding: 4px 10px; border-radius: 4px; margin-bottom: 6px;">
                🎯 TOUR {tour_idx}
            </div>
            <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(280px, 1fr)); gap: 10px;">
                {''.join(tour_cards)}
            </div>
        </div>
        """)

    return f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Poule Officielle - {nom_poule} - {nom_comp}</title>
    <style>
        @page {{
            size: landscape;
            margin: 8mm;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
            margin: 0;
            padding: 10px;
            background: white;
            color: #1E293B;
        }}
        .print-btn {{
            background-color: #0055A4;
            color: white;
            border: none;
            padding: 8px 16px;
            font-size: 13px;
            font-weight: bold;
            border-radius: 6px;
            cursor: pointer;
            margin-bottom: 12px;
        }}
        @media print {{
            .print-btn {{
                display: none !important;
            }}
            body {{
                padding: 0;
            }}
        }}
    </style>
</head>
<body>
    <button class="print-btn" onclick="window.print()">🖨️ Imprimer cette Poule (A4 Paysage)</button>
    <div style="margin-bottom: 12px;">
        <h2 style="color: #0055A4; margin: 0 0 4px 0; font-size: 20px;">🏆 {nom_comp.upper()}</h2>
        <div style="font-size: 12px; color: #64748B; font-weight: bold;">FFLDA — POULE OFFICIELLE : {nom_poule} (TOUS CONTRE TOUS)</div>
    </div>
    
    <table style="width: 100%; border-collapse: collapse; font-size: 11px; margin-bottom: 10px;">
        <thead>
            <tr style="background: #0055A4; color: white;">
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 45px;">CLT</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 35px;">N°</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; text-align: left;">NOM Prénom</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; text-align: left;">CLUB</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; text-align: left;">COMITÉ</th>
                {headers_tours}
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 65px;">Total Pts</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 65px;">Total Vict</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 55px;">Poids</th>
            </tr>
        </thead>
        <tbody>
            {''.join(lignes_html)}
        </tbody>
    </table>
    
    {podium_html}
    
    <div style="margin-top: 15px;">
        {''.join(matchs_html)}
    </div>
</body>
</html>"""

def generer_document_plateau_u7_imprimable(nom_poule, nom_comp, participants):
    lignes_html = []
    for idx, p in enumerate(participants, 1):
        bg = "#FFFBEB" if idx % 2 == 0 else "#FFFFFF"
        nom = p.get('Nom', '')
        club = p.get('Club', '')
        poids = formater_poids(p.get('Poids', ''))
        lignes_html.append(f"""
        <tr style="background: {bg};">
            <td style="border: 1px solid #CBD5E1; padding: 8px; text-align: center; font-weight: bold;">{idx}</td>
            <td style="border: 1px solid #CBD5E1; padding: 8px; font-weight: 600;">{nom}</td>
            <td style="border: 1px solid #CBD5E1; padding: 8px;">{club}</td>
            <td style="border: 1px solid #CBD5E1; padding: 8px; text-align: center;">{poids}</td>
            <td style="border: 1px solid #CBD5E1; padding: 8px; text-align: center; color: #166534; font-weight: bold;">[ ✓ ] Validé</td>
            <td style="border: 1px solid #CBD5E1; padding: 8px; text-align: center; color: #166534; font-weight: bold;">[ ✓ ] Validé</td>
            <td style="border: 1px solid #CBD5E1; padding: 8px; text-align: center; color: #166534; font-weight: bold;">[ ✓ ] Validé</td>
            <td style="border: 1px solid #CBD5E1; padding: 8px; text-align: center; background: #FEF3C7; color: #92400E; font-weight: 800;">🥇 Médaille d'Or</td>
        </tr>
        """)

    return f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Passeport U7 - {nom_poule} - {nom_comp}</title>
    <style>
        @page {{ size: landscape; margin: 8mm; }}
        body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; margin: 0; padding: 10px; background: white; color: #1E293B; }}
        .print-btn {{ background-color: #D97706; color: white; border: none; padding: 8px 16px; font-size: 13px; font-weight: bold; border-radius: 6px; cursor: pointer; margin-bottom: 12px; }}
        @media print {{ .print-btn {{ display: none !important; }} body {{ padding: 0; }} }}
    </style>
</head>
<body>
    <button class="print-btn" onclick="window.print()">🖨️ Imprimer cette Fiche Plateau U7 (A4 Paysage)</button>
    <div style="margin-bottom: 12px;">
        <h2 style="color: #D97706; margin: 0 0 4px 0; font-size: 20px;">🏆 {nom_comp.upper()}</h2>
        <div style="font-size: 13px; color: #64748B; font-weight: bold;">FFLDA — ANIMATION PLATEAU U7 : {nom_poule}</div>
        <div style="font-size: 11px; color: #92400E; font-style: italic;">Formule officielle FFLDA : Découverte pédagogique sous forme de 3 plateaux d'activités avec rotation. Tous les enfants sont récompensés !</div>
    </div>
    
    <table style="width: 100%; border-collapse: collapse; font-size: 12px; margin-bottom: 15px;">
        <thead>
            <tr style="background: #D97706; color: white;">
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 40px;">N°</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; text-align: left;">NOM Prénom</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; text-align: left;">CLUB</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 70px;">Poids</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px;">Plateau 1 : Motricité & Agilité</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px;">Plateau 2 : Ateliers techniques</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px;">Plateau 3 : Oppositions</th>
                <th style="border: 1px solid #CBD5E1; padding: 8px; width: 150px;">Validation / Récompense</th>
            </tr>
        </thead>
        <tbody>
            {''.join(lignes_html)}
        </tbody>
    </table>
    
    <div style="background: #FFFBEB; border: 1.5px solid #F59E0B; border-radius: 8px; padding: 12px; font-size: 11px;">
        <div style="font-weight: 800; font-size: 12px; color: #B45309; margin-bottom: 6px;">ℹ️ ORGANISATION DES 3 PLATEAUX D'ACTIVITÉ U7 :</div>
        <div style="display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 10px;">
            <div style="background: white; border: 1px solid #FCD34D; border-radius: 6px; padding: 8px;">
                <b style="color: #B45309;">🟡 Plateau 1 : Motricité & Agilité</b>
                <p style="margin: 4px 0 0 0; color: #78350F;">Parcours gymnique, franchissements, équilibre, réactivité motrice.</p>
            </div>
            <div style="background: white; border: 1px solid #FCD34D; border-radius: 6px; padding: 8px;">
                <b style="color: #166534;">🟢 Plateau 2 : Ateliers techniques</b>
                <p style="margin: 4px 0 0 0; color: #14532D;">Ateliers d'apprentissage technique, habiletés motrices et gestes de lutte adaptés.</p>
            </div>
            <div style="background: white; border: 1px solid #FCD34D; border-radius: 6px; padding: 8px;">
                <b style="color: #1E40AF;">🔵 Plateau 3 : Oppositions</b>
                <p style="margin: 4px 0 0 0; color: #1E3A8A;">Jeux de lutte et oppositions adaptées, combats éducatifs aménagés.</p>
            </div>
        </div>
        <div style="margin-top: 8px; text-align: center; font-weight: bold; color: #92400E;">
            🏅 Chaque enfant passe successivement sur les 3 plateaux. Tous reçoivent une médaille d'or FFLDA et un diplôme !
        </div>
    </div>
</body>
</html>"""

def make_winner_formula(c_r, c_b, pt_r, pt_b, placeholder_name, target_corner="🔴"):
    default_text = f"{target_corner} {placeholder_name}"
    r_val = f"IF({pt_r.coordinate}=\"\", 0, {pt_r.coordinate})"
    b_val = f"IF({pt_b.coordinate}=\"\", 0, {pt_b.coordinate})"
    
    if target_corner == "🔴":
        win_r = c_r.coordinate
        win_b = f'SUBSTITUTE({c_b.coordinate}, "🔵 ", "🔴 ")'
    else:
        win_r = f'SUBSTITUTE({c_r.coordinate}, "🔴 ", "🔵 ")'
        win_b = c_b.coordinate
        
    return (
        f'=IF({r_val}+{b_val}=0, "{default_text}", '
        f'IF({r_val}>{b_val}, {win_r}, '
        f'IF({b_val}>{r_val}, {win_b}, '
        f'"{default_text}")))'
    )

def make_loser_formula(c_r, c_b, pt_r, pt_b, placeholder_name, target_corner="🔴"):
    default_text = f"{target_corner} {placeholder_name}"
    r_val = f"IF({pt_r.coordinate}=\"\", 0, {pt_r.coordinate})"
    b_val = f"IF({pt_b.coordinate}=\"\", 0, {pt_b.coordinate})"
    
    if target_corner == "🔴":
        lose_r = f'SUBSTITUTE({c_b.coordinate}, "🔵 ", "🔴 ")'
        lose_b = c_r.coordinate
    else:
        lose_r = c_b.coordinate
        lose_b = f'SUBSTITUTE({c_r.coordinate}, "🔴 ", "🔵 ")'
        
    return (
        f'=IF({r_val}+{b_val}=0, "{default_text}", '
        f'IF({r_val}>{b_val}, {lose_r}, '
        f'IF({b_val}>{r_val}, {lose_b}, '
        f'"{default_text}")))'
    )

def draw_excel_match_card(ws, start_row, start_col, title, p1, p2, cat_poule, coords_map=None, bg_header=None, end_col=None):
    b_thin = Side(style='thin', color='CBD5E1')
    b_style = Border(left=b_thin, right=b_thin, top=b_thin, bottom=b_thin)
    font_match_h = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    font_p_bold = Font(name="Arial", size=9, bold=True)
    font_pts = Font(name="Arial", size=10, bold=True)
    
    fill_header = bg_header if bg_header is not None else PatternFill("solid", fgColor="475569")
    fill_red_card = PatternFill("solid", fgColor="FEF2F2")
    fill_blue_card = PatternFill("solid", fgColor="EFF6FF")
    fill_gray_box = PatternFill("solid", fgColor="F8FAFC")

    if end_col is not None and end_col > start_col + 1:
        col1 = start_col
        col2 = end_col
        # En-tête du match sur toute la largeur (col1..col2)
        ws.merge_cells(start_row=start_row, start_column=col1, end_row=start_row, end_column=col2)
        c_h = ws.cell(row=start_row, column=col1, value=title)
        c_h.font = font_match_h
        c_h.fill = fill_header
        c_h.alignment = Alignment(horizontal="center", vertical="center")
        
        # Combattant Rouge (fusionné de col1 à col2-1)
        ws.merge_cells(start_row=start_row+1, start_column=col1, end_row=start_row+1, end_column=col2-1)
        # Combattant Bleu (fusionné de col1 à col2-1)
        ws.merge_cells(start_row=start_row+2, start_column=col1, end_row=start_row+2, end_column=col2-1)
        
        for r_k in range(start_row, start_row + 3):
            for c_k in range(col1, col2 + 1):
                ws.cell(row=r_k, column=c_k).border = b_style
                if r_k == start_row:
                    ws.cell(row=r_k, column=c_k).fill = fill_header
                elif r_k == start_row + 1 and c_k < col2:
                    ws.cell(row=r_k, column=c_k).fill = fill_red_card
                elif r_k == start_row + 2 and c_k < col2:
                    ws.cell(row=r_k, column=c_k).fill = fill_blue_card
    else:
        col1 = start_col
        col2 = start_col + 1
        
        # En-tête du match
        ws.merge_cells(start_row=start_row, start_column=col1, end_row=start_row, end_column=col2)
        c_h = ws.cell(row=start_row, column=col1, value=title)
        c_h.font = font_match_h
        c_h.fill = fill_header
        c_h.alignment = Alignment(horizontal="center", vertical="center")
        c_h.border = b_style
        ws.cell(row=start_row, column=col2).border = b_style
    
    # Combattant Rouge
    nom1 = p1.get('Nom', '') if isinstance(p1, dict) else str(p1)
    club1 = p1.get('Club', '') if isinstance(p1, dict) else ''
    text_r = f"🔴 {nom1}" + (f" ({club1})" if club1 and club1 != '-' else "")
    c_r = ws.cell(row=start_row+1, column=col1)
    if isinstance(p1, dict) and 'formula' in p1:
        c_r.value = p1['formula']
    else:
        c_r.value = text_r
    c_r.font = font_p_bold
    c_r.fill = fill_red_card
    c_r.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c_r.border = b_style
    
    c_pt_r = ws.cell(row=start_row+1, column=col2)
    c_pt_r.font = font_pts
    c_pt_r.fill = fill_gray_box
    c_pt_r.alignment = Alignment(horizontal="center", vertical="center")
    c_pt_r.border = b_style
    
    # Combattant Bleu
    nom2 = p2.get('Nom', '') if isinstance(p2, dict) else str(p2)
    club2 = p2.get('Club', '') if isinstance(p2, dict) else ''
    text_b = f"🔵 {nom2}" + (f" ({club2})" if club2 and club2 != '-' else "")
    c_b = ws.cell(row=start_row+2, column=col1)
    if isinstance(p2, dict) and 'formula' in p2:
        c_b.value = p2['formula']
    else:
        c_b.value = text_b
    c_b.font = font_p_bold
    c_b.fill = fill_blue_card
    c_b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c_b.border = b_style
    
    c_pt_b = ws.cell(row=start_row+2, column=col2)
    c_pt_b.font = font_pts
    c_pt_b.fill = fill_gray_box
    c_pt_b.alignment = Alignment(horizontal="center", vertical="center")
    c_pt_b.border = b_style

    if coords_map:
        m_info = coords_map.get((cat_poule, nom1, nom2))
        if not m_info:
            m_info = coords_map.get((cat_poule, nom2, nom1))
        if not m_info:
            b1 = re.sub(r'\s*\[.*?\]', '', nom1).strip()
            b2 = re.sub(r'\s*\[.*?\]', '', nom2).strip()
            m_info = coords_map.get((cat_poule, b1, b2))
            if not m_info:
                m_info = coords_map.get((cat_poule, b2, b1))
        if not m_info:
            b1 = re.sub(r'\s*\[.*?\]', '', nom1).strip()
            b2 = re.sub(r'\s*\[.*?\]', '', nom2).strip()
            for (c, k1, k2), inf in coords_map.items():
                if cat_poule == c or cat_poule in c or c in cat_poule:
                    kb1 = re.sub(r'\s*\[.*?\]', '', str(k1)).strip()
                    kb2 = re.sub(r'\s*\[.*?\]', '', str(k2)).strip()
                    if (nom1 == k1 and nom2 == k2) or (b1 == kb1 and b2 == kb2) or (nom1 == k2 and nom2 == k1) or (b1 == kb2 and b2 == kb1):
                        m_info = inf
                        break
        if m_info and isinstance(m_info, dict):
            s_name = m_info.get('sheet') or m_info.get('sheet_name', '')
            b1 = re.sub(r'\s*\[.*?\]', '', nom1).strip()
            bp1 = re.sub(r'\s*\[.*?\]', '', str(m_info.get('p1', ''))).strip()
            ptr_c = m_info.get('ptr_cell')
            ptb_c = m_info.get('ptb_cell')
            if not ptr_c and 'ptr' in m_info and hasattr(m_info['ptr'], 'coordinate'):
                ptr_c = m_info['ptr'].coordinate
            if not ptb_c and 'ptb' in m_info and hasattr(m_info['ptb'], 'coordinate'):
                ptb_c = m_info['ptb'].coordinate
            if s_name and ptr_c and ptb_c:
                if nom1 == m_info.get('p1') or (b1 and b1 == bp1) or (b1 and b1 in bp1) or (bp1 and bp1 in b1):
                    c_pt_r.value = f"=IF('{s_name}'!{ptr_c}<>\"\", '{s_name}'!{ptr_c}, \"\")"
                    c_pt_b.value = f"=IF('{s_name}'!{ptb_c}<>\"\", '{s_name}'!{ptb_c}, \"\")"
                else:
                    c_pt_r.value = f"=IF('{s_name}'!{ptb_c}<>\"\", '{s_name}'!{ptb_c}, \"\")"
                    c_pt_b.value = f"=IF('{s_name}'!{ptr_c}<>\"\", '{s_name}'!{ptr_c}, \"\")"

    return c_pt_r, c_pt_b, c_r, c_b


def draw_excel_vertical_connector(ws, start_row, end_row, col):
    for r in range(start_row, end_row + 1):
        cell = ws.cell(row=r, column=col)
        cell.border = Border(right=Side(style='medium', color='94A3B8'))


def draw_excel_podium_card(ws, start_row, start_col, title, subtitle, fill_bg, font_color="000000", formula_val=None):
    b_thin = Side(style='thin', color='CBD5E1')
    b_style = Border(left=b_thin, right=b_thin, top=b_thin, bottom=b_thin)
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row+1, end_column=start_col+1)
    c = ws.cell(row=start_row, column=start_col)
    if formula_val:
        c.value = formula_val
    else:
        c.value = f"{title}\n{subtitle}"
    c.font = Font(name="Arial", size=10, bold=True, color=font_color)
    c.fill = fill_bg
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(start_row, start_row+2):
        for col in range(start_col, start_col+2):
            ws.cell(row=r, column=col).border = b_style


def link_tapis_slot(tapis_slots, cat, p_nom, ws_bracket, bracket_cell):
    if not tapis_slots or not p_nom or not bracket_cell:
        return
    slots_to_update = []
    
    def collect_slots(k):
        v = tapis_slots.get(k)
        if v:
            if isinstance(v, list):
                slots_to_update.extend(v)
            elif isinstance(v, tuple):
                slots_to_update.append(v)

    collect_slots((cat, p_nom))
    collect_slots(p_nom)
    
    base_nom = re.sub(r'\s*\[.*?\]', '', str(p_nom)).strip()
    if base_nom and base_nom != p_nom:
        collect_slots((cat, base_nom))
        collect_slots(base_nom)
        
    if not slots_to_update:
        for k, sl in tapis_slots.items():
            if isinstance(k, tuple):
                c, nom = k
                c_clean = str(c)
                nom_clean = str(nom)
                if (p_nom == nom_clean or p_nom in nom_clean or nom_clean in p_nom or (base_nom and base_nom in nom_clean)) and (cat in c_clean or c_clean in cat):
                    if isinstance(sl, list): slots_to_update.extend(sl)
                    elif isinstance(sl, tuple): slots_to_update.append(sl)
            elif isinstance(k, str):
                if p_nom == k or p_nom in k or k in p_nom or (base_nom and base_nom in k):
                    if isinstance(sl, list): slots_to_update.extend(sl)
                    elif isinstance(sl, tuple): slots_to_update.append(sl)
                    
    seen = set()
    formula_str = f'=SUBSTITUTE(SUBSTITUTE(\'{ws_bracket.title}\'!{bracket_cell.coordinate}, "🔴 ", ""), "🔵 ", "")'
    for ws_m, cell_coord in slots_to_update:
        if (ws_m.title, cell_coord) not in seen:
            seen.add((ws_m.title, cell_coord))
            ws_m[cell_coord].value = formula_str



def construire_feuille_tableau_excel(ws, p_obj, nom_poule, liste_p, coords_matchs_tapis, nom_competition, tapis_slots=None):
    font_title = Font(name="Arial", size=13, bold=True, color="0055A4")
    font_sub = Font(name="Arial", size=9, italic=True, color="64748B")
    font_hdr = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    
    fill_dark = PatternFill("solid", fgColor="0F172A")
    fill_blue = PatternFill("solid", fgColor="0055A4")
    fill_sky = PatternFill("solid", fgColor="0284C7")
    fill_amber = PatternFill("solid", fgColor="B45309")
    fill_gold = PatternFill("solid", fgColor="FEF3C7")
    fill_silver = PatternFill("solid", fgColor="F1F5F9")
    fill_bronze = PatternFill("solid", fgColor="FFEDD5")
    fill_zebra = PatternFill("solid", fgColor="F8FAFC")
    fill_gray_h = PatternFill("solid", fgColor="475569")
    
    b_thin = Side(style='thin', color='CBD5E1')
    b_style = Border(left=b_thin, right=b_thin, top=b_thin, bottom=b_thin)

    ws.views.sheetView[0].showGridLines = True
    
    ws.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()} — TABLEAU OFFICIEL U13 : {nom_poule}").font = font_title
    ws.cell(row=2, column=1, value="Formule officielle FFLDA : Élimination directe avec repêchage des 1/4 de finale (2 Médailles de Bronze) — Orientation : Gauche ➔ Droite").font = font_sub

    # Table des participants inscrits (Cols A-D)
    ws.cell(row=4, column=1, value="LISTE DES PARTICIPANTS").font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for c_i, h in enumerate(["N°", "NOM Prénom", "CLUB", "POIDS"], 1):
        c = ws.cell(row=4, column=c_i, value=h)
        c.fill, c.font, c.alignment, c.border = fill_blue, font_hdr, Alignment(horizontal="center", vertical="center"), b_style

    for idx, p in enumerate(liste_p, 1):
        r = 4 + idx
        c_num = ws.cell(row=r, column=1, value=idx)
        c_num.alignment, c_num.border = Alignment(horizontal="center", vertical="center"), b_style
        
        c_nom = ws.cell(row=r, column=2, value=p['Nom'])
        c_nom.border = b_style
        
        c_club = ws.cell(row=r, column=3, value=p.get('Club', ''))
        c_club.border = b_style
        
        c_pds = ws.cell(row=r, column=4, value=formater_poids(p.get('Poids', '')))
        c_pds.alignment, c_pds.border = Alignment(horizontal="center", vertical="center"), b_style
        
        if idx % 2 == 0:
            c_num.fill = fill_zebra
            c_nom.fill = fill_zebra
            c_club.fill = fill_zebra
            c_pds.fill = fill_zebra

    max_len_nom = max([len(str(p.get('Nom', ''))) for p in liste_p] + [12])
    max_len_club = max([len(str(p.get('Club', ''))) for p in liste_p if p.get('Club') and p.get('Club') != '-'] + [8])
    col_nom_w = max(max_len_nom + 4, 24)
    col_club_w = max(max_len_club + 4, 16)

    max_len_match = max([
        len(f"🔴 {p.get('Nom', '')}" + (f" ({p.get('Club', '')})" if p.get('Club') and p.get('Club') != '-' else ""))
        for p in liste_p
    ] + [22])
    col_match_w = max(max_len_match + 3, 24)
    col_pod_w = max(int(col_match_w * 0.55), 12)

    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = col_nom_w
    ws.column_dimensions['C'].width = col_club_w
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 3

    rondes = p_obj.get('rondes', [])
    n = len(liste_p)
    
    m_cp = re.search(r'(\+?\d+\s*kg)', nom_poule)
    cat_poids = m_cp.group(1) if m_cp else ""

    f_or = rondes[-1][0]
    f_b1 = rondes[-1][1] if len(rondes[-1]) > 1 else None
    f_b2 = rondes[-1][2] if len(rondes[-1]) > 2 else None

    sf1 = rondes[-2][0]
    sf2 = rondes[-2][1]
    rep1 = rondes[-2][2] if len(rondes[-2]) > 2 else None
    rep2 = rondes[-2][3] if len(rondes[-2]) > 3 else None

    if len(rondes) >= 3:
        # Configuration dynamique des colonnes selon le nombre de tours (1/32, 1/16, 1/8, 1/4, 1/2, Finales)
        col_cur = 6
        if len(rondes) >= 6:
            # 1/32 de finale + 1/16 + 1/8 + 1/4 + 1/2 + Finales (33 <= n <= 64)
            col_32 = col_cur
            col_c_pre = col_cur + 2
            col_cur += 3

            col_16 = col_cur
            col_c0 = col_cur + 2
            col_cur += 3

            col_18 = col_cur
            col_c1 = col_cur + 2
            col_cur += 3

            col_qf = col_cur
            col_c2 = col_cur + 2
            col_cur += 3

            col_sf = col_cur
            col_c3 = col_cur + 2
            col_cur += 3

            col_fn = col_cur
            col_pod = col_cur + 2

            stage_cols = [col_32, col_16, col_18, col_qf, col_sf]
            conn_col_qf_sf = col_c2
            conn_col_sf_fn = col_c3
        elif len(rondes) == 5:
            # 1/16 de finale + 1/8 de finale + 1/4 + 1/2 + Finales (17 <= n <= 32)
            col_16 = col_cur
            col_c0 = col_cur + 2
            col_cur += 3
            
            col_18 = col_cur
            col_c1 = col_cur + 2
            col_cur += 3
            
            col_qf = col_cur
            col_c2 = col_cur + 2
            col_cur += 3
            
            col_sf = col_cur
            col_c3 = col_cur + 2
            col_cur += 3
            
            col_fn = col_cur
            col_pod = col_cur + 2
            
            stage_cols = [col_16, col_18, col_qf, col_sf]
            conn_col_qf_sf = col_c2
            conn_col_sf_fn = col_c3
        elif len(rondes) == 4:
            # Tour préliminaire (1/8) + 1/4 + 1/2 + Finales (9 <= n <= 16)
            col_prelim = col_cur
            col_c0 = col_cur + 2
            col_cur += 3
            
            col_qf = col_cur
            col_c1 = col_cur + 2
            col_cur += 3
            
            col_sf = col_cur
            col_c2 = col_cur + 2
            col_cur += 3
            
            col_fn = col_cur
            col_pod = col_cur + 2
            
            stage_cols = [col_prelim, col_qf, col_sf]
            conn_col_qf_sf = col_c1
            conn_col_sf_fn = col_c2
        else:
            # 1/4 + 1/2 + Finales (n = 7 ou 8)
            col_qf = col_cur
            col_c1 = col_cur + 2
            col_cur += 3
            
            col_sf = col_cur
            col_c2 = col_cur + 2
            col_cur += 3
            
            col_fn = col_cur
            col_pod = col_cur + 2
            
            stage_cols = [col_qf, col_sf]
            conn_col_qf_sf = col_c1
            conn_col_sf_fn = col_c2

        for sc in stage_cols:
            ws.column_dimensions[get_column_letter(sc)].width = col_match_w
            ws.column_dimensions[get_column_letter(sc+1)].width = 6
            ws.column_dimensions[get_column_letter(sc+2)].width = 3
        ws.column_dimensions[get_column_letter(col_fn)].width = col_match_w
        ws.column_dimensions[get_column_letter(col_fn+1)].width = 6
        ws.column_dimensions[get_column_letter(col_pod)].width = col_pod_w
        ws.column_dimensions[get_column_letter(col_pod+1)].width = col_pod_w

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_pod+1)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_pod+1)

        # Bannière principale
        ws.merge_cells(start_row=4, start_column=6, end_row=4, end_column=col_pod+1)
        c_bann = ws.cell(row=4, column=6, value="🏆 TABLEAU PRINCIPAL D'ÉLIMINATION DIRECTE (OR / ARGENT)")
        c_bann.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c_bann.fill = fill_blue
        c_bann.alignment = Alignment(horizontal="center", vertical="center")

        qualif_map = {}
        max_upper_row = 23

        # 1. Traitement des tours préliminaires (1/32, 1/16 et 1/8) si n > 16 ou 9 <= n <= 16
        if len(rondes) >= 6:
            # 1/32 de finale
            ws.merge_cells(start_row=5, start_column=col_32, end_row=5, end_column=col_32+1)
            c_h32 = ws.cell(row=5, column=col_32, value="1/32 DE FINALE")
            c_h32.font, c_h32.fill, c_h32.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_dark, Alignment(horizontal="center", vertical="center")

            r32_matches = rondes[0]
            max_upper_row = max(max_upper_row, 6 + len(r32_matches) * 4)
            for i, m in enumerate(r32_matches):
                r_i = 6 + i * 4
                ptr, ptb, c_r, c_b = draw_excel_match_card(ws, r_i, col_32, f"1/32 DE FINALE {i+1}", m[0], m[1], nom_poule, coords_matchs_tapis, bg_header=fill_gray_h)
                k1 = f"Vainqueur 1/32 ({i+1}) [{cat_poids}]"
                k2 = f"Vainqueur 1/32 ({i+1})"
                w_info = {'c_r': c_r, 'c_b': c_b, 'ptr': ptr, 'ptb': ptb, 'nom': f"Vainqueur 1/32 ({i+1})"}
                qualif_map[k1] = w_info
                qualif_map[k2] = w_info

            # 1/16 de finale
            ws.merge_cells(start_row=5, start_column=col_16, end_row=5, end_column=col_16+1)
            c_h16 = ws.cell(row=5, column=col_16, value="1/16 DE FINALE")
            c_h16.font, c_h16.fill, c_h16.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_blue, Alignment(horizontal="center", vertical="center")

            r16_matches = rondes[1]
            max_upper_row = max(max_upper_row, 6 + len(r16_matches) * 4)
            for i, m in enumerate(r16_matches):
                r_i = 6 + i * 4
                p1_nom = m[0]['Nom'] if isinstance(m[0], dict) else str(m[0])
                p2_nom = m[1]['Nom'] if isinstance(m[1], dict) else str(m[1])
                
                if p1_nom in qualif_map:
                    qi = qualif_map[p1_nom]
                    p1_in = {'Nom': p1_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔴")}
                else:
                    p1_in = m[0]
                    
                if p2_nom in qualif_map:
                    qi = qualif_map[p2_nom]
                    p2_in = {'Nom': p2_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔵")}
                else:
                    p2_in = m[1]
                    
                ptr, ptb, c_r, c_b = draw_excel_match_card(ws, r_i, col_16, f"1/16 DE FINALE {i+1}", p1_in, p2_in, nom_poule, coords_matchs_tapis, bg_header=fill_gray_h)
                
                if p1_nom in qualif_map:
                    link_tapis_slot(tapis_slots, nom_poule, p1_nom, ws, c_r)
                if p2_nom in qualif_map:
                    link_tapis_slot(tapis_slots, nom_poule, p2_nom, ws, c_b)
                    
                k1 = f"Vainqueur 1/16 ({i+1}) [{cat_poids}]"
                k2 = f"Vainqueur 1/16 ({i+1})"
                w_info = {'c_r': c_r, 'c_b': c_b, 'ptr': ptr, 'ptb': ptb, 'nom': f"Vainqueur 1/16 ({i+1})"}
                qualif_map[k1] = w_info
                qualif_map[k2] = w_info

            # 1/8 de finale
            ws.merge_cells(start_row=5, start_column=col_18, end_row=5, end_column=col_18+1)
            c_h18 = ws.cell(row=5, column=col_18, value="1/8 DE FINALE")
            c_h18.font, c_h18.fill, c_h18.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_blue, Alignment(horizontal="center", vertical="center")

            r18_matches = rondes[2]
            max_upper_row = max(max_upper_row, 6 + len(r18_matches) * 4)
            for i, m in enumerate(r18_matches):
                r_i = 6 + i * 4
                p1_nom = m[0]['Nom'] if isinstance(m[0], dict) else str(m[0])
                p2_nom = m[1]['Nom'] if isinstance(m[1], dict) else str(m[1])
                
                if p1_nom in qualif_map:
                    qi = qualif_map[p1_nom]
                    p1_in = {'Nom': p1_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔴")}
                else:
                    p1_in = m[0]
                    
                if p2_nom in qualif_map:
                    qi = qualif_map[p2_nom]
                    p2_in = {'Nom': p2_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔵")}
                else:
                    p2_in = m[1]
                    
                ptr, ptb, c_r, c_b = draw_excel_match_card(ws, r_i, col_18, f"1/8 DE FINALE {i+1}", p1_in, p2_in, nom_poule, coords_matchs_tapis, bg_header=fill_sky)
                
                if p1_nom in qualif_map:
                    link_tapis_slot(tapis_slots, nom_poule, p1_nom, ws, c_r)
                if p2_nom in qualif_map:
                    link_tapis_slot(tapis_slots, nom_poule, p2_nom, ws, c_b)
                    
                k1 = f"Vainqueur 1/8 ({i+1}) [{cat_poids}]"
                k2 = f"Vainqueur 1/8 ({i+1})"
                w_info = {'c_r': c_r, 'c_b': c_b, 'ptr': ptr, 'ptb': ptb, 'nom': f"Vainqueur 1/8 ({i+1})"}
                qualif_map[k1] = w_info
                qualif_map[k2] = w_info

        elif len(rondes) == 5:
            ws.merge_cells(start_row=5, start_column=col_16, end_row=5, end_column=col_16+1)
            c_h16 = ws.cell(row=5, column=col_16, value="1/16 DE FINALE")
            c_h16.font, c_h16.fill, c_h16.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_dark, Alignment(horizontal="center", vertical="center")

            r16_matches = rondes[0]
            max_upper_row = max(max_upper_row, 6 + len(r16_matches) * 4)
            for i, m in enumerate(r16_matches):
                r_i = 6 + i * 4
                ptr, ptb, c_r, c_b = draw_excel_match_card(ws, r_i, col_16, f"1/16 DE FINALE {i+1}", m[0], m[1], nom_poule, coords_matchs_tapis, bg_header=fill_gray_h)
                k1 = f"Vainqueur 1/16 ({i+1}) [{cat_poids}]"
                k2 = f"Vainqueur 1/16 ({i+1})"
                w_info = {'c_r': c_r, 'c_b': c_b, 'ptr': ptr, 'ptb': ptb, 'nom': f"Vainqueur 1/16 ({i+1})"}
                qualif_map[k1] = w_info
                qualif_map[k2] = w_info

            # Traitement des 1/8 de finale
            ws.merge_cells(start_row=5, start_column=col_18, end_row=5, end_column=col_18+1)
            c_h18 = ws.cell(row=5, column=col_18, value="1/8 DE FINALE")
            c_h18.font, c_h18.fill, c_h18.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_blue, Alignment(horizontal="center", vertical="center")

            r18_matches = rondes[1]
            max_upper_row = max(max_upper_row, 6 + len(r18_matches) * 4)
            for i, m in enumerate(r18_matches):
                r_i = 6 + i * 4
                p1_nom = m[0]['Nom'] if isinstance(m[0], dict) else str(m[0])
                p2_nom = m[1]['Nom'] if isinstance(m[1], dict) else str(m[1])
                
                if p1_nom in qualif_map:
                    qi = qualif_map[p1_nom]
                    p1_in = {'Nom': p1_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔴")}
                else:
                    p1_in = m[0]
                    
                if p2_nom in qualif_map:
                    qi = qualif_map[p2_nom]
                    p2_in = {'Nom': p2_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔵")}
                else:
                    p2_in = m[1]
                    
                ptr, ptb, c_r, c_b = draw_excel_match_card(ws, r_i, col_18, f"1/8 DE FINALE {i+1}", p1_in, p2_in, nom_poule, coords_matchs_tapis, bg_header=fill_sky)
                
                if p1_nom in qualif_map:
                    link_tapis_slot(tapis_slots, nom_poule, p1_nom, ws, c_r)
                if p2_nom in qualif_map:
                    link_tapis_slot(tapis_slots, nom_poule, p2_nom, ws, c_b)
                    
                k1 = f"Vainqueur 1/8 ({i+1}) [{cat_poids}]"
                k2 = f"Vainqueur 1/8 ({i+1})"
                w_info = {'c_r': c_r, 'c_b': c_b, 'ptr': ptr, 'ptb': ptb, 'nom': f"Vainqueur 1/8 ({i+1})"}
                qualif_map[k1] = w_info
                qualif_map[k2] = w_info

        elif len(rondes) == 4:
            # Traitement du Tour préliminaire (1/8)
            ws.merge_cells(start_row=5, start_column=col_prelim, end_row=5, end_column=col_prelim+1)
            c_hp = ws.cell(row=5, column=col_prelim, value="TOUR PRÉLIMINAIRE (1/8)")
            c_hp.font, c_hp.fill, c_hp.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_dark, Alignment(horizontal="center", vertical="center")

            prelim_matches = rondes[0]
            max_upper_row = max(max_upper_row, 6 + len(prelim_matches) * 4)
            for i, m in enumerate(prelim_matches):
                r_i = 6 + i * 4
                ptr, ptb, c_r, c_b = draw_excel_match_card(ws, r_i, col_prelim, f"PRÉLIMINAIRE {i+1}", m[0], m[1], nom_poule, coords_matchs_tapis, bg_header=fill_gray_h)
                k1 = f"Vainqueur Prél. {i+1} [{cat_poids}]"
                k2 = f"Vainqueur Prél. {i+1}"
                w_info = {'c_r': c_r, 'c_b': c_b, 'ptr': ptr, 'ptb': ptb, 'nom': f"Vainqueur Prél. {i+1}"}
                qualif_map[k1] = w_info
                qualif_map[k2] = w_info

        # 2. Quarts de finale (Toujours rondes[-3])
        ws.merge_cells(start_row=5, start_column=col_qf, end_row=5, end_column=col_qf+1)
        c_hqf = ws.cell(row=5, column=col_qf, value="QUARTS DE FINALE")
        c_hqf.font, c_hqf.fill, c_hqf.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_blue, Alignment(horizontal="center", vertical="center")

        q_matches = rondes[-3]
        q_list = [
            q_matches[0],
            q_matches[1],
            q_matches[2],
            q_matches[3] if len(q_matches) > 3 else (p_obj.get('p_exempt', liste_p[6]), {'Nom': 'EXEMPT (BYE)', 'Club': '-'})
        ]

        qf_cards = []
        qf_rows = [6, 11, 16, 21]
        for k in range(4):
            m = q_list[k]
            p1_obj = m[0]
            p2_obj = m[1]
            p1_nom = p1_obj.get('Nom', '') if isinstance(p1_obj, dict) else str(p1_obj)
            p2_nom = p2_obj.get('Nom', '') if isinstance(p2_obj, dict) else str(p2_obj)
            
            p1_in = p1_obj
            if p1_nom in qualif_map:
                qi = qualif_map[p1_nom]
                p1_in = {'Nom': p1_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔴")}
                
            p2_in = p2_obj
            if p2_nom in qualif_map:
                qi = qualif_map[p2_nom]
                p2_in = {'Nom': p2_nom, 'formula': make_winner_formula(qi['c_r'], qi['c_b'], qi['ptr'], qi['ptb'], qi['nom'], "🔵")}
                
            q_ptr, q_ptb, q_r, q_b = draw_excel_match_card(ws, qf_rows[k], col_qf, f"1/4 DE FINALE {k+1}", p1_in, p2_in, nom_poule, coords_matchs_tapis)
            qf_cards.append((q_ptr, q_ptb, q_r, q_b))
            
            if p1_nom in qualif_map:
                link_tapis_slot(tapis_slots, nom_poule, p1_nom, ws, q_r)
            if p2_nom in qualif_map:
                link_tapis_slot(tapis_slots, nom_poule, p2_nom, ws, q_b)

        qf1_ptr, qf1_ptb, qf1_r, qf1_b = qf_cards[0]
        qf2_ptr, qf2_ptb, qf2_r, qf2_b = qf_cards[1]
        qf3_ptr, qf3_ptb, qf3_r, qf3_b = qf_cards[2]
        qf4_ptr, qf4_ptb, qf4_r, qf4_b = qf_cards[3]

        draw_excel_vertical_connector(ws, 7, 12, conn_col_qf_sf)
        ws.cell(row=10, column=conn_col_qf_sf).border = Border(right=Side(style='medium', color='94A3B8'), bottom=Side(style='medium', color='94A3B8'))

        draw_excel_vertical_connector(ws, 17, 22, conn_col_qf_sf)
        ws.cell(row=20, column=conn_col_qf_sf).border = Border(right=Side(style='medium', color='94A3B8'), bottom=Side(style='medium', color='94A3B8'))

        # 3. Demi-finales dynamiques
        ws.merge_cells(start_row=5, start_column=col_sf, end_row=5, end_column=col_sf+1)
        c_hsf = ws.cell(row=5, column=col_sf, value="DEMI-FINALES")
        c_hsf.font, c_hsf.fill, c_hsf.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_sky, Alignment(horizontal="center", vertical="center")

        sf1_p1 = {'Nom': sf1[0]['Nom'], 'formula': make_winner_formula(qf1_r, qf1_b, qf1_ptr, qf1_ptb, "Vainqueur 1/4 (1)", "🔴")}
        sf1_p2 = {'Nom': sf1[1]['Nom'], 'formula': make_winner_formula(qf2_r, qf2_b, qf2_ptr, qf2_ptb, "Vainqueur 1/4 (2)", "🔵")}
        sf1_ptr, sf1_ptb, sf1_r, sf1_b = draw_excel_match_card(ws, 8, col_sf, "DEMI-FINALE 1", sf1_p1, sf1_p2, nom_poule, coords_matchs_tapis, bg_header=fill_sky)

        sf2_p1 = {'Nom': sf2[0]['Nom'], 'formula': make_winner_formula(qf3_r, qf3_b, qf3_ptr, qf3_ptb, "Vainqueur 1/4 (3)", "🔴")}
        if n == 7:
            sf2_p2 = sf2[1]
        else:
            sf2_p2 = {'Nom': sf2[1]['Nom'], 'formula': make_winner_formula(qf4_r, qf4_b, qf4_ptr, qf4_ptb, "Vainqueur 1/4 (4)", "🔵")}
        sf2_ptr, sf2_ptb, sf2_r, sf2_b = draw_excel_match_card(ws, 18, col_sf, "DEMI-FINALE 2", sf2_p1, sf2_p2, nom_poule, coords_matchs_tapis, bg_header=fill_sky)

        draw_excel_vertical_connector(ws, 10, 19, conn_col_sf_fn)
        ws.cell(row=14, column=conn_col_sf_fn).border = Border(right=Side(style='medium', color='94A3B8'), bottom=Side(style='medium', color='94A3B8'))

        # 4. Grande Finale dynamique
        ws.merge_cells(start_row=5, start_column=col_fn, end_row=5, end_column=col_fn+1)
        c_hfn = ws.cell(row=5, column=col_fn, value="GRANDE FINALE")
        c_hfn.font, c_hfn.fill, c_hfn.alignment = Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill_dark, Alignment(horizontal="center", vertical="center")

        fn_p1 = {'Nom': f_or[0]['Nom'], 'formula': make_winner_formula(sf1_r, sf1_b, sf1_ptr, sf1_ptb, "Vainqueur 1/2 (1)", "🔴")}
        fn_p2 = {'Nom': f_or[1]['Nom'], 'formula': make_winner_formula(sf2_r, sf2_b, sf2_ptr, sf2_ptb, "Vainqueur 1/2 (2)", "🔵")}
        fn_ptr, fn_ptb, fn_r, fn_b = draw_excel_match_card(ws, 13, col_fn, "GRANDE FINALE (OR)", fn_p1, fn_p2, nom_poule, coords_matchs_tapis, bg_header=fill_dark)

        # 5. Podiums Or & Argent
        r_fn = f"IF({fn_ptr.coordinate}=\"\",0,{fn_ptr.coordinate})"
        b_fn = f"IF({fn_ptb.coordinate}=\"\",0,{fn_ptb.coordinate})"
        form_gold = f'=IF({r_fn}+{b_fn}=0, "🥇 CHAMPION (OR)" & CHAR(10) & "Vainqueur Grande Finale", "🥇 CHAMPION (OR)" & CHAR(10) & SUBSTITUTE(SUBSTITUTE(IF({r_fn}>{b_fn}, {fn_r.coordinate}, IF({b_fn}>{r_fn}, {fn_b.coordinate}, "En attente")), "🔴 ", ""), "🔵 ", ""))'
        form_silver = f'=IF({r_fn}+{b_fn}=0, "🥈 VICE-CHAMPION (ARGENT)" & CHAR(10) & "Perdant Grande Finale", "🥈 VICE-CHAMPION (ARGENT)" & CHAR(10) & SUBSTITUTE(SUBSTITUTE(IF({r_fn}>{b_fn}, {fn_b.coordinate}, IF({b_fn}>{r_fn}, {fn_r.coordinate}, "En attente")), "🔴 ", ""), "🔵 ", ""))'
        draw_excel_podium_card(ws, 12, col_pod, "🥇 CHAMPION (OR)", "Vainqueur Grande Finale", fill_gold, font_color="B45309", formula_val=form_gold)
        draw_excel_podium_card(ws, 15, col_pod, "🥈 VICE-CHAMPION (ARGENT)", "Perdant Grande Finale", fill_silver, font_color="475569", formula_val=form_silver)

        # 6. Repêchages & Bronze
        row_rep = max(26, max_upper_row + 3)
        ws.merge_cells(start_row=row_rep, start_column=col_qf, end_row=row_rep, end_column=col_pod+1)
        c_rep_h = ws.cell(row=row_rep, column=col_qf, value="🔄 TABLEAU DE REPÊCHAGE & MATCHS POUR LE BRONZE (2 Troisièmes Places)")
        c_rep_h.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c_rep_h.fill = fill_amber
        c_rep_h.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells(start_row=row_rep+1, start_column=col_qf, end_row=row_rep+1, end_column=col_pod+1)
        c_rep_sub = ws.cell(row=row_rep+1, column=col_qf, value="Perdants des 1/4 ➔ Repêchages ➔ Finales Bronze contre les perdants des Demi-Finales")
        c_rep_sub.font = font_sub
        c_rep_sub.alignment = Alignment(horizontal="left", vertical="center")

        # Repêchage 1 (Perdant QF 1 vs Perdant QF 2)
        rep1_p1 = {'Nom': rep1[0]['Nom'], 'formula': make_loser_formula(qf1_r, qf1_b, qf1_ptr, qf1_ptb, "Perdant 1/4 (1)", "🔴")}
        rep1_p2 = {'Nom': rep1[1]['Nom'], 'formula': make_loser_formula(qf2_r, qf2_b, qf2_ptr, qf2_ptb, "Perdant 1/4 (2)", "🔵")}
        rep1_ptr, rep1_ptb, rep1_r, rep1_b = draw_excel_match_card(ws, row_rep+3, col_qf, "REPÊCHAGE 1/4 (1)", rep1_p1, rep1_p2, nom_poule, coords_matchs_tapis, bg_header=fill_sky)

        if n >= 8:
            rep2_p1 = {'Nom': rep2[0]['Nom'], 'formula': make_loser_formula(qf3_r, qf3_b, qf3_ptr, qf3_ptb, "Perdant 1/4 (3)", "🔴")}
            rep2_p2 = {'Nom': rep2[1]['Nom'], 'formula': make_loser_formula(qf4_r, qf4_b, qf4_ptr, qf4_ptb, "Perdant 1/4 (4)", "🔵")}
            rep2_ptr, rep2_ptb, rep2_r, rep2_b = draw_excel_match_card(ws, row_rep+8, col_qf, "REPÊCHAGE 1/4 (2)", rep2_p1, rep2_p2, nom_poule, coords_matchs_tapis, bg_header=fill_sky)
        elif n == 7:
            ws.merge_cells(start_row=row_rep+8, start_column=col_qf, end_row=row_rep+10, end_column=col_qf+1)
            c_ex = ws.cell(row=row_rep+8, column=col_qf, value="Exempt de repêchage 1\n(Avance direct en Finale Bronze 2)")
            c_ex.font = Font(name="Arial", size=9, italic=True, color="64748B")
            c_ex.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for r_k in range(row_rep+8, row_rep+11):
                for c_k in range(col_qf, col_qf+2):
                    ws.cell(row=r_k, column=c_k).border = b_style

        ws.cell(row=row_rep+4, column=conn_col_qf_sf).border = Border(bottom=Side(style='medium', color='94A3B8'))
        ws.cell(row=row_rep+9, column=conn_col_qf_sf).border = Border(bottom=Side(style='medium', color='94A3B8'))

        # Finale Bronze 1 (Vainqueur Repêchage 1 vs Perdant Demi-Finale 2)
        fb1_p1 = {'Nom': f_b1[0]['Nom'], 'formula': make_winner_formula(rep1_r, rep1_b, rep1_ptr, rep1_ptb, "Vainqueur Repêchage 1", "🔴")}
        fb1_p2 = {'Nom': f_b1[1]['Nom'], 'formula': make_loser_formula(sf2_r, sf2_b, sf2_ptr, sf2_ptb, "Perdant 1/2 (2)", "🔵")}
        b1_ptr, b1_ptb, b1_r, b1_b = draw_excel_match_card(ws, row_rep+3, col_sf, "FINALE BRONZE 1", fb1_p1, fb1_p2, nom_poule, coords_matchs_tapis, bg_header=fill_amber)

        r_b1 = f"IF({b1_ptr.coordinate}=\"\",0,{b1_ptr.coordinate})"
        b_b1 = f"IF({b1_ptb.coordinate}=\"\",0,{b1_ptb.coordinate})"
        form_b1 = f'=IF({r_b1}+{b_b1}=0, "🥉 3ème PLACE (Bronze 1)" & CHAR(10) & "Vainqueur Finale Bronze 1", "🥉 3ème PLACE (Bronze 1)" & CHAR(10) & SUBSTITUTE(SUBSTITUTE(IF({r_b1}>{b_b1}, {b1_r.coordinate}, IF({b_b1}>{r_b1}, {b1_b.coordinate}, "En attente")), "🔴 ", ""), "🔵 ", ""))'
        draw_excel_podium_card(ws, row_rep+3, col_pod, "🥉 3ème PLACE (Bronze 1)", "Vainqueur Finale Bronze 1", fill_bronze, font_color="9A3412", formula_val=form_b1)

        # Finale Bronze 2 (Vainqueur Repêchage 2 ou Perdant QF 3 vs Perdant Demi-Finale 1)
        if n == 7:
            fb2_p1 = {'Nom': f_b2[0]['Nom'], 'formula': make_loser_formula(qf3_r, qf3_b, qf3_ptr, qf3_ptb, "Perdant 1/4 (3)", "🔴")}
        else:
            fb2_p1 = {'Nom': f_b2[0]['Nom'], 'formula': make_winner_formula(rep2_r, rep2_b, rep2_ptr, rep2_ptb, "Vainqueur Repêchage 2", "🔴")}
        fb2_p2 = {'Nom': f_b2[1]['Nom'], 'formula': make_loser_formula(sf1_r, sf1_b, sf1_ptr, sf1_ptb, "Perdant 1/2 (1)", "🔵")}
        b2_ptr, b2_ptb, b2_r, b2_b = draw_excel_match_card(ws, row_rep+8, col_sf, "FINALE BRONZE 2", fb2_p1, fb2_p2, nom_poule, coords_matchs_tapis, bg_header=fill_amber)

        r_b2 = f"IF({b2_ptr.coordinate}=\"\",0,{b2_ptr.coordinate})"
        b_b2 = f"IF({b2_ptb.coordinate}=\"\",0,{b2_ptb.coordinate})"
        form_b2 = f'=IF({r_b2}+{b_b2}=0, "🥉 3ème PLACE (Bronze 2)" & CHAR(10) & "Vainqueur Finale Bronze 2", "🥉 3ème PLACE (Bronze 2)" & CHAR(10) & SUBSTITUTE(SUBSTITUTE(IF({r_b2}>{b_b2}, {b2_r.coordinate}, IF({b_b2}>{r_b2}, {b2_b.coordinate}, "En attente")), "🔴 ", ""), "🔵 ", ""))'
        draw_excel_podium_card(ws, row_rep+8, col_pod, "🥉 3ème PLACE (Bronze 2)", "Vainqueur Finale Bronze 2", fill_bronze, font_color="9A3412", formula_val=form_b2)

        # Liaison dynamique vers les cartes de match sur les Grilles Tapis
        link_tapis_slot(tapis_slots, nom_poule, sf1[0]['Nom'], ws, sf1_r)
        link_tapis_slot(tapis_slots, nom_poule, sf1[1]['Nom'], ws, sf1_b)
        link_tapis_slot(tapis_slots, nom_poule, sf2[0]['Nom'], ws, sf2_r)
        if n >= 8 and isinstance(sf2[1], dict):
            link_tapis_slot(tapis_slots, nom_poule, sf2[1]['Nom'], ws, sf2_b)

        if rep1:
            link_tapis_slot(tapis_slots, nom_poule, rep1[0]['Nom'], ws, rep1_r)
            link_tapis_slot(tapis_slots, nom_poule, rep1[1]['Nom'], ws, rep1_b)
        if n >= 8 and rep2:
            link_tapis_slot(tapis_slots, nom_poule, rep2[0]['Nom'], ws, rep2_r)
            link_tapis_slot(tapis_slots, nom_poule, rep2[1]['Nom'], ws, rep2_b)

        link_tapis_slot(tapis_slots, nom_poule, f_or[0]['Nom'], ws, fn_r)
        link_tapis_slot(tapis_slots, nom_poule, f_or[1]['Nom'], ws, fn_b)

        if f_b1:
            link_tapis_slot(tapis_slots, nom_poule, f_b1[0]['Nom'], ws, b1_r)
            link_tapis_slot(tapis_slots, nom_poule, f_b1[1]['Nom'], ws, b1_b)
        if f_b2:
            link_tapis_slot(tapis_slots, nom_poule, f_b2[0]['Nom'], ws, b2_r)
            link_tapis_slot(tapis_slots, nom_poule, f_b2[1]['Nom'], ws, b2_b)

    else:
        col_cur = 6
        for r_idx, ronde in enumerate(rondes):
            col_match = col_cur
            col_conn = col_cur + 2
            
            ws.column_dimensions[get_column_letter(col_match)].width = col_match_w
            ws.column_dimensions[get_column_letter(col_match+1)].width = 6
            ws.column_dimensions[get_column_letter(col_conn)].width = 3

            r_title = f"RONDE {r_idx+1}"
            if r_idx == 0: r_title = "TOUR PRÉLIMINAIRE"
            elif r_idx == len(rondes)-3: r_title = "QUARTS DE FINALE"
            elif r_idx == len(rondes)-2: r_title = "DEMI-FINALES & REP."
            elif r_idx == len(rondes)-1: r_title = "FINALES"

            ws.merge_cells(start_row=4, start_column=col_match, end_row=4, end_column=col_match+1)
            c_rt = ws.cell(row=4, column=col_match, value=r_title)
            c_rt.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c_rt.fill = fill_blue if r_idx < len(rondes)-1 else fill_dark
            c_rt.alignment = Alignment(horizontal="center", vertical="center")

            row_cur = 6
            for m_i, m in enumerate(ronde):
                m_label = f"Match {m_i+1}"
                draw_excel_match_card(ws, row_cur, col_match, m_label, m[0], m[1], nom_poule, coords_matchs_tapis)
                row_cur += 4
            
            col_cur += 3

        col_pod = col_cur
        ws.column_dimensions[get_column_letter(col_pod)].width = col_pod_w
        ws.column_dimensions[get_column_letter(col_pod+1)].width = col_pod_w
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_pod+1)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_pod+1)
        draw_excel_podium_card(ws, 6, col_pod, "🥇 CHAMPION (OR)", "Vainqueur Finale", fill_gold, font_color="B45309")
        draw_excel_podium_card(ws, 9, col_pod, "🥈 VICE-CHAMPION", "Finaliste", fill_silver, font_color="475569")
        draw_excel_podium_card(ws, 12, col_pod, "🥉 3ème PLACE (1)", "Bronze 1", fill_bronze, font_color="9A3412")
        draw_excel_podium_card(ws, 15, col_pod, "🥉 3ème PLACE (2)", "Bronze 2", fill_bronze, font_color="9A3412")

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def construire_feuille_poules_croisees_excel(ws, p_obj, nom_poule, liste_p, coords_matchs_tapis, nom_competition, tapis_slots=None):
    font_title = Font(name="Arial", size=13, bold=True, color="0055A4")
    font_sub = Font(name="Arial", size=9, italic=True, color="64748B")
    font_hdr = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_match_h = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    font_pts = Font(name="Arial", size=10, bold=True)
    
    fill_dark = PatternFill("solid", fgColor="0F172A")
    fill_blue = PatternFill("solid", fgColor="0055A4")
    fill_sky = PatternFill("solid", fgColor="0284C7")
    fill_amber = PatternFill("solid", fgColor="B45309")
    fill_gray_h = PatternFill("solid", fgColor="475569")
    fill_gold = PatternFill("solid", fgColor="FEF3C7")
    fill_silver = PatternFill("solid", fgColor="F1F5F9")
    fill_bronze = PatternFill("solid", fgColor="FFEDD5")
    
    b_thin = Side(style='thin', color='CBD5E1')
    b_style = Border(left=b_thin, right=b_thin, top=b_thin, bottom=b_thin)

    ws.views.sheetView[0].showGridLines = True
    
    ws.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()} — POULES CROISÉES U13 (6 LUTTEURS) : {nom_poule}").font = font_title
    ws.cell(row=2, column=1, value="Formule officielle FFLDA : Phase 1 (2 Poules de 3 Nordiques) ➔ Phase 2 (Demi-Finales Croisées & Finales Or/Argent et Bronze unique)").font = font_sub

    poule_a = p_obj.get('poule_a', liste_p[:3])
    poule_b = p_obj.get('poule_b', liste_p[3:])
    rondes = p_obj.get('rondes', [])

    lignes_lutteurs = {}

    def render_sub_poule_table(ws, start_row, title, participants, bg_color):
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=8)
        c_title = ws.cell(row=start_row, column=1, value=title)
        c_title.font, c_title.fill, c_title.alignment = font_hdr, bg_color, Alignment(horizontal="center", vertical="center")
        
        headers = ["CLT", "N°", "NOM Prénom", "CLUB", "Tour 1", "Tour 2", "Tour 3", "Total Pts"]
        for c_i, h in enumerate(headers, 1):
            c = ws.cell(row=start_row+1, column=c_i, value=h)
            c.font, c.fill, c.alignment, c.border = font_match_h, fill_gray_h, Alignment(horizontal="center", vertical="center"), b_style
            
        row_cur = start_row + 2
        sub_poule_match_cols = {(1, 2): 'E', (2, 1): 'E', (2, 3): 'F', (3, 2): 'F', (1, 3): 'G', (3, 1): 'G'}
        for idx, p in enumerate(participants, 1):
            lignes_lutteurs[p['Nom']] = row_cur
            
            comparisons = []
            for idx_j in range(1, len(participants) + 1):
                if idx_j == idx:
                    continue
                r_j = start_row + 1 + idx_j
                col_t_lettre = sub_poule_match_cols.get((idx, idx_j))
                if col_t_lettre:
                    comp = (
                        f"IF(H{r_j}>H{row_cur}, 1, "
                        f"IF(H{r_j}<H{row_cur}, 0, "
                        f"IF({col_t_lettre}{r_j}>{col_t_lettre}{row_cur}, 1, "
                        f"IF({col_t_lettre}{r_j}<{col_t_lettre}{row_cur}, 0, "
                        f"IF({r_j}<{row_cur}, 1, 0)))))"
                    )
                else:
                    comp = f"IF(H{r_j}>H{row_cur}, 1, IF(H{r_j}<H{row_cur}, 0, IF({r_j}<{row_cur}, 1, 0)))"
                comparisons.append(comp)

            somme_comp = " + ".join(comparisons) if comparisons else "0"
            plage_tot = f"H${start_row+2}:H${start_row+1+len(participants)}"
            c_clt = ws.cell(row=row_cur, column=1, value=f'=IF(SUM({plage_tot})=0, "", 1 + {somme_comp})')
            c_clt.alignment, c_clt.border = Alignment(horizontal="center", vertical="center"), b_style
            c_clt.font = Font(name="Arial", size=10, bold=True, color="0055A4")
            
            c_num = ws.cell(row=row_cur, column=2, value=idx)
            c_num.alignment, c_num.border = Alignment(horizontal="center", vertical="center"), b_style
            
            c_nom = ws.cell(row=row_cur, column=3, value=p['Nom'])
            c_nom.border = b_style
            
            c_club = ws.cell(row=row_cur, column=4, value=p.get('Club', ''))
            c_club.border = b_style
            
            for t_i in range(3):
                c_t = ws.cell(row=row_cur, column=5+t_i)
                c_t.alignment, c_t.border = Alignment(horizontal="center", vertical="center"), b_style
            
            c_tot = ws.cell(row=row_cur, column=8, value=f"=SUM(E{row_cur}:G{row_cur})")
            c_tot.font, c_tot.alignment, c_tot.border = font_pts, Alignment(horizontal="center", vertical="center"), b_style
            
            row_cur += 1

    render_sub_poule_table(ws, 4, "🥋 PHASE 1 : POULE A (3 Lutteurs)", poule_a, fill_blue)
    render_sub_poule_table(ws, 10, "🥋 PHASE 1 : POULE B (3 Lutteurs)", poule_b, fill_dark)

    max_len_nom = max([len(str(p.get('Nom', ''))) for p in liste_p] + [12])
    max_len_club = max([len(str(p.get('Club', ''))) for p in liste_p if p.get('Club') and p.get('Club') != '-'] + [8])
    col_nom_w = max(max_len_nom + 4, 22)
    col_club_w = max(max_len_club + 4, 16)

    max_len_match = max([
        len(f"🔴 {p.get('Nom', '')}" + (f" ({p.get('Club', '')})" if p.get('Club') and p.get('Club') != '-' else ""))
        for p in liste_p
    ] + [22])
    col_match_w = max(max_len_match + 3, 23)
    col_pod_w = max(int(col_match_w * 0.55), 12)

    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 5
    ws.column_dimensions['C'].width = col_nom_w
    ws.column_dimensions['D'].width = col_club_w
    ws.column_dimensions['E'].width = 8
    ws.column_dimensions['F'].width = 8
    ws.column_dimensions['G'].width = 8
    ws.column_dimensions['H'].width = 10
    ws.column_dimensions['I'].width = 3

    # Matchs des Tours 1, 2, 3 en bas à gauche (Rows 16+)
    r_matches = 16
    ws.merge_cells(start_row=r_matches, start_column=1, end_row=r_matches, end_column=8)
    c_m_hdr = ws.cell(row=r_matches, column=1, value="🤼 RENCONTRES DES POULES A & B (TOURS 1 À 3)")
    c_m_hdr.font, c_m_hdr.fill, c_m_hdr.alignment = font_hdr, fill_gray_h, Alignment(horizontal="center", vertical="center")
    
    r_matches += 1
    for tour_idx in range(3):
        if tour_idx < len(rondes):
            ronde = rondes[tour_idx]
            for m_idx, m in enumerate(ronde):
                grp_tag = "Poule A" if m_idx == 0 else "Poule B"
                title_m = f"T{tour_idx+1} ({grp_tag})"
                pt_r, pt_b, _, _ = draw_excel_match_card(ws, r_matches, 1, title_m, m[0], m[1], nom_poule, coords_matchs_tapis, end_col=8)
                
                p1_nom = m[0]['Nom']
                p2_nom = m[1]['Nom']
                if p1_nom in lignes_lutteurs:
                    ws.cell(row=lignes_lutteurs[p1_nom], column=5+tour_idx, value=f"={pt_r.coordinate}")
                if p2_nom in lignes_lutteurs:
                    ws.cell(row=lignes_lutteurs[p2_nom], column=5+tour_idx, value=f"={pt_b.coordinate}")
                
                r_matches += 4

    # Phase 2 : Demi-Finales Croisées & Finales (Cols J+)
    col_sf = 10
    col_c = 12
    col_fn = 13
    col_pod = 15

    ws.column_dimensions['J'].width = col_match_w
    ws.column_dimensions['K'].width = 6
    ws.column_dimensions['L'].width = 3
    ws.column_dimensions['M'].width = col_match_w
    ws.column_dimensions['N'].width = 6
    ws.column_dimensions['O'].width = col_pod_w
    ws.column_dimensions['P'].width = col_pod_w

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_pod+1)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_pod+1)

    ws.merge_cells(start_row=4, start_column=col_sf, end_row=4, end_column=col_pod+1)
    c_fin_h = ws.cell(row=4, column=col_sf, value="🏆 PHASE 2 : PHASE FINALE CROISÉE (Gauche ➔ Droite)")
    c_fin_h.font, c_fin_h.fill, c_fin_h.alignment = font_hdr, fill_sky, Alignment(horizontal="center", vertical="center")

    sf1 = rondes[3][0]
    sf2 = rondes[3][1]
    f_or = rondes[4][0]
    f_b = rondes[4][1]

    # Qualifications dynamiques depuis les rangs des poules A et B
    # Poule A est en rangs 6 à 8 (Col A = Rang, Col C = Nom)
    # Poule B est en rangs 12 à 14 (Col A = Rang, Col C = Nom)
    form_1er_a = '=IF(OR(SUM(H$6:H$8)=0, ISNA(MATCH(1, A$6:A$8, 0))), "🔴 1er Poule A", "🔴 " & INDEX(C$6:C$8, MATCH(1, A$6:A$8, 0)) & IF(INDEX(D$6:D$8, MATCH(1, A$6:A$8, 0))<>"", " (" & INDEX(D$6:D$8, MATCH(1, A$6:A$8, 0)) & ")", ""))'
    form_2e_b  = '=IF(OR(SUM(H$12:H$14)=0, ISNA(MATCH(2, A$12:A$14, 0))), "🔵 2ème Poule B", "🔵 " & INDEX(C$12:C$14, MATCH(2, A$12:A$14, 0)) & IF(INDEX(D$12:D$14, MATCH(2, A$12:A$14, 0))<>"", " (" & INDEX(D$12:D$14, MATCH(2, A$12:A$14, 0)) & ")", ""))'
    form_1er_b = '=IF(OR(SUM(H$12:H$14)=0, ISNA(MATCH(1, A$12:A$14, 0))), "🔴 1er Poule B", "🔴 " & INDEX(C$12:C$14, MATCH(1, A$12:A$14, 0)) & IF(INDEX(D$12:D$14, MATCH(1, A$12:A$14, 0))<>"", " (" & INDEX(D$12:D$14, MATCH(1, A$12:A$14, 0)) & ")", ""))'
    form_2e_a  = '=IF(OR(SUM(H$6:H$8)=0, ISNA(MATCH(2, A$6:A$8, 0))), "🔵 2ème Poule A", "🔵 " & INDEX(C$6:C$8, MATCH(2, A$6:A$8, 0)) & IF(INDEX(D$6:D$8, MATCH(2, A$6:A$8, 0))<>"", " (" & INDEX(D$6:D$8, MATCH(2, A$6:A$8, 0)) & ")", ""))'

    p_sf1_1 = {'Nom': sf1[0]['Nom'], 'formula': form_1er_a}
    p_sf1_2 = {'Nom': sf1[1]['Nom'], 'formula': form_2e_b}
    sf1_ptr, sf1_ptb, sf1_r, sf1_b = draw_excel_match_card(ws, 6, col_sf, "DEMI-FINALE 1 (1er A vs 2ème B)", p_sf1_1, p_sf1_2, nom_poule, coords_matchs_tapis, bg_header=fill_blue)

    p_sf2_1 = {'Nom': sf2[0]['Nom'], 'formula': form_1er_b}
    p_sf2_2 = {'Nom': sf2[1]['Nom'], 'formula': form_2e_a}
    sf2_ptr, sf2_ptb, sf2_r, sf2_b = draw_excel_match_card(ws, 12, col_sf, "DEMI-FINALE 2 (1er B vs 2ème A)", p_sf2_1, p_sf2_2, nom_poule, coords_matchs_tapis, bg_header=fill_blue)

    draw_excel_vertical_connector(ws, 7, 13, col_c)
    ws.cell(row=8, column=col_c).border = Border(right=Side(style='medium', color='94A3B8'), bottom=Side(style='medium', color='94A3B8'))
    ws.cell(row=14, column=col_c).border = Border(right=Side(style='medium', color='94A3B8'), bottom=Side(style='medium', color='94A3B8'))

    # Finales dynamiques
    f_or_p1 = {'Nom': f_or[0]['Nom'], 'formula': make_winner_formula(sf1_r, sf1_b, sf1_ptr, sf1_ptb, "Vainqueur SF1", "🔴")}
    f_or_p2 = {'Nom': f_or[1]['Nom'], 'formula': make_winner_formula(sf2_r, sf2_b, sf2_ptr, sf2_ptb, "Vainqueur SF2", "🔵")}
    for_ptr, for_ptb, for_r, for_b = draw_excel_match_card(ws, 7, col_fn, "FINALE 1-2 (OR / ARGENT)", f_or_p1, f_or_p2, nom_poule, coords_matchs_tapis, bg_header=fill_dark)

    f_b_p1 = {'Nom': f_b[0]['Nom'], 'formula': make_loser_formula(sf1_r, sf1_b, sf1_ptr, sf1_ptb, "Perdant SF1", "🔴")}
    f_b_p2 = {'Nom': f_b[1]['Nom'], 'formula': make_loser_formula(sf2_r, sf2_b, sf2_ptr, sf2_ptb, "Perdant SF2", "🔵")}
    fb_ptr, fb_ptb, fb_r, fb_b = draw_excel_match_card(ws, 13, col_fn, "FINALE 3-4 (BRONZE UNIQUE)", f_b_p1, f_b_p2, nom_poule, coords_matchs_tapis, bg_header=fill_amber)

    # Podiums dynamiques
    r_for = f"IF({for_ptr.coordinate}=\"\",0,{for_ptr.coordinate})"
    b_for = f"IF({for_ptb.coordinate}=\"\",0,{for_ptb.coordinate})"
    r_fb = f"IF({fb_ptr.coordinate}=\"\",0,{fb_ptr.coordinate})"
    b_fb = f"IF({fb_ptb.coordinate}=\"\",0,{fb_ptb.coordinate})"
    form_or = f'=IF({r_for}+{b_for}=0, "🥇 OR" & CHAR(10) & "Vainqueur Finale 1-2", "🥇 OR" & CHAR(10) & SUBSTITUTE(SUBSTITUTE(IF({r_for}>{b_for}, {for_r.coordinate}, IF({b_for}>{r_for}, {for_b.coordinate}, "En attente")), "🔴 ", ""), "🔵 ", ""))'
    form_arg = f'=IF({r_for}+{b_for}=0, "🥈 ARGENT" & CHAR(10) & "Perdant Finale 1-2", "🥈 ARGENT" & CHAR(10) & SUBSTITUTE(SUBSTITUTE(IF({r_for}>{b_for}, {for_b.coordinate}, IF({b_for}>{r_for}, {for_r.coordinate}, "En attente")), "🔴 ", ""), "🔵 ", ""))'
    form_brz = f'=IF({r_fb}+{b_fb}=0, "🥉 BRONZE (Unique)" & CHAR(10) & "Vainqueur Finale 3-4", "🥉 BRONZE" & CHAR(10) & SUBSTITUTE(SUBSTITUTE(IF({r_fb}>{b_fb}, {fb_r.coordinate}, IF({b_fb}>{r_fb}, {fb_b.coordinate}, "En attente")), "🔴 ", ""), "🔵 ", ""))'

    draw_excel_podium_card(ws, 6, col_pod, "🥇 OR", "Vainqueur Finale 1-2", fill_gold, font_color="B45309", formula_val=form_or)
    draw_excel_podium_card(ws, 9, col_pod, "🥈 ARGENT", "Perdant Finale 1-2", fill_silver, font_color="475569", formula_val=form_arg)
    draw_excel_podium_card(ws, 13, col_pod, "🥉 BRONZE (Unique)", "Vainqueur Finale 3-4", fill_bronze, font_color="9A3412", formula_val=form_brz)

    # Liaison dynamique vers les cartes de match sur les Grilles Tapis
    link_tapis_slot(tapis_slots, nom_poule, sf1[0]['Nom'], ws, sf1_r)
    link_tapis_slot(tapis_slots, nom_poule, sf1[1]['Nom'], ws, sf1_b)
    link_tapis_slot(tapis_slots, nom_poule, sf2[0]['Nom'], ws, sf2_r)
    link_tapis_slot(tapis_slots, nom_poule, sf2[1]['Nom'], ws, sf2_b)

    link_tapis_slot(tapis_slots, nom_poule, f_or[0]['Nom'], ws, for_r)
    link_tapis_slot(tapis_slots, nom_poule, f_or[1]['Nom'], ws, for_b)

    link_tapis_slot(tapis_slots, nom_poule, f_b[0]['Nom'], ws, fb_r)
    link_tapis_slot(tapis_slots, nom_poule, f_b[1]['Nom'], ws, fb_b)

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def construire_feuille_plateau_u7_excel(ws, p_obj, nom_poule, liste_p, coords_matchs_tapis, nom_competition):
    font_title = Font(name="Arial", size=13, bold=True, color="B45309")
    font_sub = Font(name="Arial", size=9, italic=True, color="64748B")
    font_hdr = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_pts = Font(name="Arial", size=10, bold=True)
    
    fill_amber = PatternFill("solid", fgColor="D97706")  # Ambre U7
    fill_pastel = PatternFill("solid", fgColor="FEF3C7") # Pastel U7
    fill_zebra = PatternFill("solid", fgColor="FFFBEB")
    
    b_thin = Side(style='thin', color='CBD5E1')
    b_style = Border(left=b_thin, right=b_thin, top=b_thin, bottom=b_thin)
    
    ws.views.sheetView[0].showGridLines = True
    
    ws.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()} — ANIMATION PLATEAU U7 : {nom_poule}").font = font_title
    ws.cell(row=2, column=1, value="Formule officielle FFLDA : Découverte pédagogique sous forme de 3 plateaux d'activités avec rotation (Tous les enfants sont récompensés)").font = font_sub
    
    headers = [
        "N°", "NOM Prénom", "CLUB", "POIDS", 
        "Plateau 1 : Motricité & Agilité", "Plateau 2 : Ateliers techniques", "Plateau 3 : Oppositions", 
        "Validation / Récompense"
    ]
    
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    
    for c_i, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=c_i, value=h)
        c.fill, c.font, c.alignment, c.border = fill_amber, font_hdr, Alignment(horizontal="center", vertical="center", wrap_text=True), b_style
    ws.row_dimensions[4].height = 28
    
    for idx, p in enumerate(liste_p, 1):
        r = 4 + idx
        ws.row_dimensions[r].height = 22
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra if is_even else PatternFill(fill_type=None)
        
        c_num = ws.cell(row=r, column=1, value=idx)
        c_num.alignment, c_num.border, c_num.fill = Alignment(horizontal="center", vertical="center"), b_style, row_fill
        
        c_nom = ws.cell(row=r, column=2, value=p.get('Nom', ''))
        c_nom.border, c_nom.fill = b_style, row_fill
        c_nom.alignment = Alignment(vertical="center", indent=1)
        
        c_club = ws.cell(row=r, column=3, value=p.get('Club', ''))
        c_club.border, c_club.fill = b_style, row_fill
        c_club.alignment = Alignment(vertical="center", indent=1)
        
        c_pds = ws.cell(row=r, column=4, value=formater_poids(p.get('Poids', '')))
        c_pds.alignment, c_pds.border, c_pds.fill = Alignment(horizontal="center", vertical="center"), b_style, row_fill
        
        for p_col in [5, 6, 7]:
            c_plat = ws.cell(row=r, column=p_col, value="[ ✓ ] Validé")
            c_plat.alignment, c_plat.border, c_plat.fill = Alignment(horizontal="center", vertical="center"), b_style, row_fill
            c_plat.font = Font(name="Arial", size=9, bold=True, color="166534")
            
        c_rec = ws.cell(row=r, column=8, value="🥇 Médaille d'Or / Diplôme")
        c_rec.alignment, c_rec.border, c_rec.fill = Alignment(horizontal="center", vertical="center"), b_style, fill_pastel
        c_rec.font = Font(name="Arial", size=9, bold=True, color="92400E")
        
    max_len_nom = max([len(str(p.get('Nom', ''))) for p in liste_p] + [12])
    max_len_club = max([len(str(p.get('Club', ''))) for p in liste_p if p.get('Club') and p.get('Club') != '-'] + [10])
    
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = max(max_len_nom + 4, 22)
    ws.column_dimensions['C'].width = max(max_len_club + 4, 16)
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 24
    ws.column_dimensions['F'].width = 24
    ws.column_dimensions['G'].width = 24
    ws.column_dimensions['H'].width = 24
    
    r_info = 4 + len(liste_p) + 2
    ws.merge_cells(start_row=r_info, start_column=1, end_row=r_info, end_column=8)
    c_info_h = ws.cell(row=r_info, column=1, value="ℹ️ DÉROULEMENT DES 3 PLATEAUX D'ACTIVITÉ U7")
    c_info_h.fill, c_info_h.font, c_info_h.alignment = fill_amber, font_hdr, Alignment(horizontal="center", vertical="center")
    
    explications = [
        ("🟡 Plateau 1 : Motricité & Agilité", "Parcours gymnique, franchissements, agilité, équilibre, réactivité motrice."),
        ("🟢 Plateau 2 : Ateliers techniques", "Ateliers d'apprentissage technique, habiletés motrices et gestes de lutte adaptés."),
        ("🔵 Plateau 3 : Oppositions", "Jeux de lutte et oppositions adaptées, combats éducatifs aménagés."),
        ("🏅 Esprit FFLDA", "Chaque enfant passe successivement sur les 3 plateaux. Tous reçoivent un diplôme et une médaille !")
    ]
    for idx_e, (tit, desc) in enumerate(explications, start=1):
        r_e = r_info + idx_e
        ws.cell(row=r_e, column=1, value=tit).font = Font(name="Arial", size=9, bold=True, color="92400E")
        ws.merge_cells(start_row=r_e, start_column=2, end_row=r_e, end_column=8)
        ws.cell(row=r_e, column=2, value=desc).font = Font(name="Arial", size=9, italic=True)
        
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def construire_feuille_poule_nordique_excel(ws, nom_poule, liste_p, rondes, coords_matchs_tapis, nom_competition):
    font_title = Font(name="Arial", size=13, bold=True, color="0055A4")
    font_sub = Font(name="Arial", size=9, italic=True, color="64748B")
    font_hdr = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_pts = Font(name="Arial", size=10, bold=True)
    
    fill_blue = PatternFill("solid", fgColor="0055A4")
    fill_gold = PatternFill("solid", fgColor="FEF3C7")
    fill_silver = PatternFill("solid", fgColor="F1F5F9")
    fill_bronze = PatternFill("solid", fgColor="FFEDD5")
    fill_zebra = PatternFill("solid", fgColor="F8FAFC")
    fill_gray_h = PatternFill("solid", fgColor="475569")
    
    b_thin = Side(style='thin', color='CBD5E1')
    b_style = Border(left=b_thin, right=b_thin, top=b_thin, bottom=b_thin)

    ws.views.sheetView[0].showGridLines = True
    
    # Titre officiel FFLDA & sous-titre
    ws.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()} — POULE OFFICIELLE : {nom_poule}").font = font_title
    ws.cell(row=2, column=1, value="Formule officielle FFLDA : Tournoi nordique (Tous contre tous) — Point de classement : 2 pt = victoire, 1 pt = nul, 0 pt = défaite").font = font_sub

    nb_tours = len(rondes) if rondes else 0
    headers = ["CLT", "N°", "NOM Prénom", "CLUB", "COMITÉ"]
    for t in range(1, nb_tours + 1):
        headers.append(f"Tour {t}")
    headers.extend(["Total Pts", "Total Vict", "Poids"])
    
    col_fin_table = len(headers)
    col_pts_idx = 6 + nb_tours
    col_vict_idx = 7 + nb_tours
    col_poids_idx = 8 + nb_tours
    
    row_cursor = 4
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=row_cursor, column=col_idx, value=h)
        c.font, c.alignment, c.border = font_hdr, Alignment(horizontal="center", vertical="center"), b_style
        c.fill = fill_blue
        
    lignes_lutteurs = {}
    ligne_debut_poule = row_cursor + 1
    
    col_pts_lettre = get_column_letter(col_pts_idx)
    col_debut_tours_lettre = get_column_letter(6)
    col_fin_tours_lettre = get_column_letter(5 + nb_tours) if nb_tours > 0 else col_pts_lettre
    plage_totaux = f"${col_pts_lettre}${ligne_debut_poule}:${col_pts_lettre}${ligne_debut_poule + len(liste_p) - 1}"
    plage_clt = f"$A${ligne_debut_poule}:$A${ligne_debut_poule + len(liste_p) - 1}"
    plage_nom = f"$C${ligne_debut_poule}:$C${ligne_debut_poule + len(liste_p) - 1}"
    plage_club = f"$D${ligne_debut_poule}:$D${ligne_debut_poule + len(liste_p) - 1}"

    # Correspondance des combats directs entre lutteurs pour le départage FFLDA
    match_col_map = {}
    for tour_idx, ronde in enumerate(rondes, 1):
        col_t_lettre = get_column_letter(5 + tour_idx)
        for match in ronde:
            p1_nom = match[0].get('Nom', '') if isinstance(match[0], dict) else str(match[0])
            p2_nom = match[1].get('Nom', '') if isinstance(match[1], dict) else str(match[1])
            match_col_map[(p1_nom, p2_nom)] = col_t_lettre
            match_col_map[(p2_nom, p1_nom)] = col_t_lettre

    for idx, p in enumerate(liste_p, 1):
        r = ligne_debut_poule + idx - 1
        lignes_lutteurs[p['Nom']] = r
        
        # Formule CLT officielle FFLDA : classement par Total Pts, avec départage au combat direct en cas d'égalité
        comparisons = []
        for idx_j, p_j in enumerate(liste_p, 1):
            if idx_j == idx:
                continue
            r_j = ligne_debut_poule + idx_j - 1
            col_t_lettre = match_col_map.get((p['Nom'], p_j['Nom']))
            
            if col_t_lettre:
                # Si j a plus de points que r, j est devant (+1).
                # Si j a moins de points que r, j est derrière (+0).
                # En cas d'égalité de points totaux (col_pts_lettre), départage par la victoire directe (col_t_lettre) :
                # Si j a battu r au combat direct, j est devant (+1). Si r a battu j, r est devant (+0).
                # En cas d'égalité absolue (ex: match nul direct 1-1 ou 0-0), départage résiduel par la ligne.
                comp = (
                    f"IF({col_pts_lettre}{r_j}>{col_pts_lettre}{r}, 1, "
                    f"IF({col_pts_lettre}{r_j}<{col_pts_lettre}{r}, 0, "
                    f"IF({col_t_lettre}{r_j}>{col_t_lettre}{r}, 1, "
                    f"IF({col_t_lettre}{r_j}<{col_t_lettre}{r}, 0, "
                    f"IF({r_j}<{r}, 1, 0)))))"
                )
            else:
                comp = f"IF({col_pts_lettre}{r_j}>{col_pts_lettre}{r}, 1, IF({col_pts_lettre}{r_j}<{col_pts_lettre}{r}, 0, IF({r_j}<{r}, 1, 0)))"
            comparisons.append(comp)

        if comparisons:
            somme_comp = " + ".join(comparisons)
            form_clt = f'=IF(SUM({plage_totaux})=0, "", 1 + {somme_comp})'
        else:
            form_clt = f'=IF(SUM({plage_totaux})=0, "", 1)'

        c_clt = ws.cell(row=r, column=1, value=form_clt)
        c_clt.alignment, c_clt.border = Alignment(horizontal="center", vertical="center"), b_style
        c_clt.font = Font(name="Arial", size=10, bold=True, color="0055A4")
        
        c_num = ws.cell(row=r, column=2, value=idx)
        c_num.alignment, c_num.border = Alignment(horizontal="center", vertical="center"), b_style
        
        c_nom = ws.cell(row=r, column=3, value=p['Nom'])
        c_nom.border = b_style
        
        c_club = ws.cell(row=r, column=4, value=p.get('Club', ''))
        c_club.border = b_style
        
        c_comite = ws.cell(row=r, column=5, value=p.get('Comité', ''))
        c_comite.border = b_style
        
        for t in range(nb_tours):
            c_tour = ws.cell(row=r, column=6 + t)
            c_tour.alignment, c_tour.border = Alignment(horizontal="center", vertical="center"), b_style
            
        if nb_tours > 0:
            c_tot_pts = ws.cell(row=r, column=col_pts_idx, value=f"=SUM({col_debut_tours_lettre}{r}:{col_fin_tours_lettre}{r})")
            c_tot_vict = ws.cell(row=r, column=col_vict_idx, value=f"=COUNTIF({col_debut_tours_lettre}{r}:{col_fin_tours_lettre}{r}, 2)")
        else:
            c_tot_pts = ws.cell(row=r, column=col_pts_idx, value=0)
            c_tot_vict = ws.cell(row=r, column=col_vict_idx, value="")
            
        c_tot_pts.font, c_tot_pts.alignment, c_tot_pts.border = font_pts, Alignment(horizontal="center", vertical="center"), b_style
        c_tot_vict.alignment, c_tot_vict.border = Alignment(horizontal="center", vertical="center"), b_style
        
        c_poids = ws.cell(row=r, column=col_poids_idx, value=formater_poids(p.get('Poids', '')))
        c_poids.alignment, c_poids.border = Alignment(horizontal="center", vertical="center"), b_style
        
        if idx % 2 == 0:
            for col_k in range(1, col_fin_table + 1):
                ws.cell(row=r, column=col_k).fill = fill_zebra

    # Ajustement dynamique des largeurs de colonnes
    max_len_nom = max([len(str(p.get('Nom', ''))) for p in liste_p] + [12])
    max_len_club = max([len(str(p.get('Club', ''))) for p in liste_p if p.get('Club') and p.get('Club') != '-'] + [10])
    max_len_comite = max([len(str(p.get('Comité', ''))) for p in liste_p if p.get('Comité') and p.get('Comité') != '-'] + [10])

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 5
    ws.column_dimensions['C'].width = max(max_len_nom + 4, 24)
    ws.column_dimensions['D'].width = max(max_len_club + 4, 16)
    ws.column_dimensions['E'].width = max(max_len_comite + 4, 16)
    for t in range(nb_tours):
        col_t_lettre = get_column_letter(6 + t)
        ws.column_dimensions[col_t_lettre].width = 9
    ws.column_dimensions[col_pts_lettre].width = 11
    ws.column_dimensions[get_column_letter(col_vict_idx)].width = 11
    ws.column_dimensions[get_column_letter(col_poids_idx)].width = 10

    # Cartes Podium Dynamiques (reliées au classement)
    col_pod = col_poids_idx + 2
    col_pod_lettre1 = get_column_letter(col_pod)
    col_pod_lettre2 = get_column_letter(col_pod + 1)
    ws.column_dimensions[get_column_letter(col_pod - 1)].width = 3
    ws.column_dimensions[col_pod_lettre1].width = 16
    ws.column_dimensions[col_pod_lettre2].width = 16
    
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_pod+1)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_pod+1)
    
    form_gold = (
        f'=IF(SUM({plage_totaux})=0, "🥇 CHAMPION (OR)" & CHAR(10) & "En attente", '
        f'"🥇 CHAMPION (OR)" & CHAR(10) & INDEX({plage_nom}, MATCH(1, {plage_clt}, 0)) & '
        f'IF(INDEX({plage_club}, MATCH(1, {plage_clt}, 0))<>"", " (" & INDEX({plage_club}, MATCH(1, {plage_clt}, 0)) & ")", ""))'
    )
    form_silver = (
        f'=IF(SUM({plage_totaux})=0, "🥈 VICE-CHAMPION (ARGENT)" & CHAR(10) & "En attente", '
        f'"🥈 VICE-CHAMPION (ARGENT)" & CHAR(10) & INDEX({plage_nom}, MATCH(2, {plage_clt}, 0)) & '
        f'IF(INDEX({plage_club}, MATCH(2, {plage_clt}, 0))<>"", " (" & INDEX({plage_club}, MATCH(2, {plage_clt}, 0)) & ")", ""))'
    )
    
    draw_excel_podium_card(ws, 4, col_pod, "🥇 CHAMPION (OR)", "En attente", fill_gold, font_color="B45309", formula_val=form_gold)
    draw_excel_podium_card(ws, 7, col_pod, "🥈 VICE-CHAMPION (ARGENT)", "En attente", fill_silver, font_color="475569", formula_val=form_silver)
    
    if len(liste_p) >= 3:
        form_bronze = (
            f'=IF(SUM({plage_totaux})=0, "🥉 3ème PLACE (BRONZE)" & CHAR(10) & "En attente", '
            f'"🥉 3ème PLACE (BRONZE)" & CHAR(10) & INDEX({plage_nom}, MATCH(3, {plage_clt}, 0)) & '
            f'IF(INDEX({plage_club}, MATCH(3, {plage_clt}, 0))<>"", " (" & INDEX({plage_club}, MATCH(3, {plage_clt}, 0)) & ")", ""))'
        )
        draw_excel_podium_card(ws, 10, col_pod, "🥉 3ème PLACE (BRONZE)", "En attente", fill_bronze, font_color="9A3412", formula_val=form_bronze)

    # Fiches de combat modernes en bas de feuille
    row_cursor = max(ligne_debut_poule + len(liste_p) + 2, 13)
    match_cpt = 1
    
    for tour_idx, ronde in enumerate(rondes, 1):
        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=col_fin_table)
        c_tour = ws.cell(row=row_cursor, column=1, value=f"🎯 TOUR {tour_idx}")
        c_tour.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c_tour.fill = fill_blue
        c_tour.alignment = Alignment(horizontal="center", vertical="center")
        for c_k in range(1, col_fin_table + 1):
            ws.cell(row=row_cursor, column=c_k).border = b_style
        row_cursor += 2
        
        for match in ronde:
            p1, p2 = match[0], match[1]
            m_title = f"COMBAT N°{match_cpt} (TOUR {tour_idx})"
            
            c_pt_r, c_pt_b, c_r, c_b = draw_excel_match_card(
                ws=ws,
                start_row=row_cursor,
                start_col=1,
                title=m_title,
                p1=p1,
                p2=p2,
                cat_poule=nom_poule,
                coords_map=coords_matchs_tapis,
                bg_header=fill_gray_h,
                end_col=col_fin_table
            )
            
            if p1['Nom'] in lignes_lutteurs:
                lig1 = lignes_lutteurs[p1['Nom']]
                ws.cell(row=lig1, column=6 + (tour_idx - 1), value=f"={c_pt_r.coordinate}").alignment = Alignment(horizontal="center", vertical="center")
            if p2['Nom'] in lignes_lutteurs:
                lig2 = lignes_lutteurs[p2['Nom']]
                ws.cell(row=lig2, column=6 + (tour_idx - 1), value=f"={c_pt_b.coordinate}").alignment = Alignment(horizontal="center", vertical="center")
                
            match_cpt += 1
            row_cursor += 4
        row_cursor += 1

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def fusionner_poules_isolees(poules, multiplicateur_poids, max_size):
    """
    Évite d'avoir un lutteur seul dans une poule de 1 tout en respectant 
    STRICTEMENT la tolérance d'écart de poids et la taille maximale de poule.
    """
    if len(poules) <= 1:
        return poules
    
    ameliore = True
    iterations = 0
    max_iterations = 20
    
    while ameliore and iterations < max_iterations:
        ameliore = False
        iterations += 1
        
        i = 0
        while i < len(poules):
            if len(poules[i]['participants']) == 1:
                p_iso = poules[i]['participants'][0]
                fusionne = False
                
                # Option 1: Essayer de fusionner directement avec la poule précédente (i-1)
                if i > 0:
                    cand = sorted(poules[i-1]['participants'] + [p_iso], key=lambda x: float(x.get('Poids_Num') or 0.0))
                    if len(cand) <= max_size and poule_poids_valide(cand, multiplicateur_poids):
                        poules[i-1]['participants'] = cand
                        poules[i-1]['rondes'] = generer_rondes_fflda(cand)
                        p_min = cand[0].get('Poids_Num') or 0.0
                        p_max = cand[-1].get('Poids_Num') or 0.0
                        prefix = poules[i-1]['nom'].split(' (')[0]
                        poules[i-1]['nom'] = f"{prefix} ({formater_poids_court(p_min)} - {formater_poids_court(p_max)})"
                        poules.pop(i)
                        fusionne = True
                        ameliore = True
                        break
                
                # Option 2: Essayer de fusionner directement avec la poule suivante (i+1)
                if not fusionne and i < len(poules) - 1:
                    cand = sorted([p_iso] + poules[i+1]['participants'], key=lambda x: float(x.get('Poids_Num') or 0.0))
                    if len(cand) <= max_size and poule_poids_valide(cand, multiplicateur_poids):
                        poules[i+1]['participants'] = cand
                        poules[i+1]['rondes'] = generer_rondes_fflda(cand)
                        p_min = cand[0].get('Poids_Num') or 0.0
                        p_max = cand[-1].get('Poids_Num') or 0.0
                        prefix = poules[i+1]['nom'].split(' (')[0]
                        poules[i+1]['nom'] = f"{prefix} ({formater_poids_court(p_min)} - {formater_poids_court(p_max)})"
                        poules.pop(i)
                        fusionne = True
                        ameliore = True
                        break
                
                # Option 3: Rééquilibrer avec la poule précédente (i-1)
                if not fusionne and i > 0 and len(poules[i-1]['participants']) >= 3:
                    prev_parts = poules[i-1]['participants']
                    for k in range(1, len(prev_parts) - 1):
                        new_prev = sorted(prev_parts[:-k], key=lambda x: float(x.get('Poids_Num') or 0.0))
                        new_iso = sorted(prev_parts[-k:] + [p_iso], key=lambda x: float(x.get('Poids_Num') or 0.0))
                        
                        if (len(new_prev) >= 2 and len(new_iso) >= 2 and len(new_iso) <= max_size and
                            poule_poids_valide(new_prev, multiplicateur_poids) and
                            poule_poids_valide(new_iso, multiplicateur_poids)):
                            
                            poules[i-1]['participants'] = new_prev
                            poules[i-1]['rondes'] = generer_rondes_fflda(new_prev)
                            p_min1 = new_prev[0].get('Poids_Num') or 0.0
                            p_max1 = new_prev[-1].get('Poids_Num') or 0.0
                            prefix1 = poules[i-1]['nom'].split(' (')[0]
                            poules[i-1]['nom'] = f"{prefix1} ({formater_poids_court(p_min1)} - {formater_poids_court(p_max1)})"
                            
                            poules[i]['participants'] = new_iso
                            poules[i]['rondes'] = generer_rondes_fflda(new_iso)
                            p_min2 = new_iso[0].get('Poids_Num') or 0.0
                            p_max2 = new_iso[-1].get('Poids_Num') or 0.0
                            prefix2 = poules[i]['nom'].split(' (')[0]
                            poules[i]['nom'] = f"{prefix2} ({formater_poids_court(p_min2)} - {formater_poids_court(p_max2)})"
                            
                            fusionne = True
                            ameliore = True
                            break
                    if fusionne:
                        break
                
                # Option 4: Rééquilibrer avec la poule suivante (i+1)
                if not fusionne and i < len(poules) - 1 and len(poules[i+1]['participants']) >= 3:
                    next_parts = poules[i+1]['participants']
                    for k in range(1, len(next_parts) - 1):
                        new_iso = sorted([p_iso] + next_parts[:k], key=lambda x: float(x.get('Poids_Num') or 0.0))
                        new_next = sorted(next_parts[k:], key=lambda x: float(x.get('Poids_Num') or 0.0))
                        
                        if (len(new_iso) >= 2 and len(new_next) >= 2 and len(new_iso) <= max_size and
                            poule_poids_valide(new_iso, multiplicateur_poids) and
                            poule_poids_valide(new_next, multiplicateur_poids)):
                            
                            poules[i]['participants'] = new_iso
                            poules[i]['rondes'] = generer_rondes_fflda(new_iso)
                            p_min1 = new_iso[0].get('Poids_Num') or 0.0
                            p_max1 = new_iso[-1].get('Poids_Num') or 0.0
                            prefix1 = poules[i]['nom'].split(' (')[0]
                            poules[i]['nom'] = f"{prefix1} ({formater_poids_court(p_min1)} - {formater_poids_court(p_max1)})"
                            
                            poules[i+1]['participants'] = new_next
                            poules[i+1]['rondes'] = generer_rondes_fflda(new_next)
                            p_min2 = new_next[0].get('Poids_Num') or 0.0
                            p_max2 = new_next[-1].get('Poids_Num') or 0.0
                            prefix2 = poules[i+1]['nom'].split(' (')[0]
                            poules[i+1]['nom'] = f"{prefix2} ({formater_poids_court(p_min2)} - {formater_poids_court(p_max2)})"
                            
                            fusionne = True
                            ameliore = True
                            break
                    if fusionne:
                        break
            i += 1
            
    return poules

def compter_collisions_club(participants_poule):
    clubs = [str(p.get('Club', '')).strip().lower() for p in participants_poule if str(p.get('Club', '')).strip() not in ['', '-', 'indépendant', 'independant', 'none', 'nan']]
    if not clubs:
        return 0
    from collections import Counter
    counts = Counter(clubs)
    return sum(c - 1 for c in counts.values() if c > 1)

def poule_poids_valide(participants_poule, multiplicateur_poids):
    if not participants_poule:
        return True
    poids_list = []
    for p in participants_poule:
        try:
            val = p.get('Poids_Num')
            if val is not None and str(val).strip() != '':
                f_val = float(val)
                if f_val > 0:
                    poids_list.append(f_val)
        except (ValueError, TypeError):
            continue
    if not poids_list:
        return True
    p_min = min(poids_list)
    p_max = max(poids_list)
    try:
        mult = float(multiplicateur_poids)
    except (ValueError, TypeError):
        mult = 1.15
    return p_max <= round(p_min * mult, 4)

def optimiser_poules_clubs(poules_groupe, multiplicateur_poids):
    """
    Permute les lutteurs entre poules d'un même groupe (âge/sexe/niveau) pour réduire 
    au maximum les affrontements entre lutteurs d'un même club, tout en respectant 
    strictement la tolérance d'écart de poids.
    """
    if len(poules_groupe) <= 1:
        return poules_groupe
    
    ameliore = True
    iterations = 0
    max_iterations = 50
    
    while ameliore and iterations < max_iterations:
        ameliore = False
        iterations += 1
        
        for i in range(len(poules_groupe)):
            for j in range(i + 1, len(poules_groupe)):
                p1 = poules_groupe[i]['participants']
                p2 = poules_groupe[j]['participants']
                
                cost_before = compter_collisions_club(p1) + compter_collisions_club(p2)
                if cost_before == 0:
                    continue
                
                best_swap = None
                best_cost = cost_before
                
                for idx1, w1 in enumerate(p1):
                    for idx2, w2 in enumerate(p2):
                        p1_test = p1[:idx1] + [w2] + p1[idx1+1:]
                        p2_test = p2[:idx2] + [w1] + p2[idx2+1:]
                        
                        if poule_poids_valide(p1_test, multiplicateur_poids) and poule_poids_valide(p2_test, multiplicateur_poids):
                            cost_after = compter_collisions_club(p1_test) + compter_collisions_club(p2_test)
                            if cost_after < best_cost:
                                best_cost = cost_after
                                best_swap = (idx1, idx2, p1_test, p2_test)
                
                if best_swap:
                    idx1, idx2, p1_test, p2_test = best_swap
                    poules_groupe[i]['participants'] = sorted(p1_test, key=lambda x: float(x.get('Poids_Num') or 0.0))
                    poules_groupe[j]['participants'] = sorted(p2_test, key=lambda x: float(x.get('Poids_Num') or 0.0))
                    
                    for idx_p in [i, j]:
                        parts = poules_groupe[idx_p]['participants']
                        prefix = poules_groupe[idx_p]['nom'].split(' (')[0]
                        p_min = parts[0].get('Poids_Num') or 0.0
                        p_max = parts[-1].get('Poids_Num') or 0.0
                        poules_groupe[idx_p]['nom'] = f"{prefix} ({formater_poids_court(p_min)} - {formater_poids_court(p_max)})"
                        poules_groupe[idx_p]['rondes'] = generer_rondes_fflda(parts)
                    
                    ameliore = True
                    break
            if ameliore:
                break
                
    return poules_groupe

# --- CONFIGURATION HARMONISÉE DES COULEURS PAR CATÉGORIE D'ÂGE (GRILLES DE PASSAGE & TAPIS) ---
COULEURS_AGE_GRILLE = {
    'U7': {
        'bg_excel_pastel': PatternFill("solid", fgColor="FEF3C7"),  # Ambre / Jaune doré doux
        'bg_excel_header': PatternFill("solid", fgColor="D97706"),  # Ambre soutenu
        'hex_pastel': '#FEF3C7',
        'hex_header': '#D97706',
        'hex_text': '#92400E',
        'badge': '🟡 U7',
        'nom': 'U7 (Jaune ambré)'
    },
    'U9': {
        'bg_excel_pastel': PatternFill("solid", fgColor="DCFCE7"),  # Vert menthe doux
        'bg_excel_header': PatternFill("solid", fgColor="059669"),  # Vert émeraude soutenu
        'hex_pastel': '#DCFCE7',
        'hex_header': '#059669',
        'hex_text': '#166534',
        'badge': '🟢 U9',
        'nom': 'U9 (Vert menthe)'
    },
    'U11': {
        'bg_excel_pastel': PatternFill("solid", fgColor="E0F2FE"),  # Bleu ciel doux
        'bg_excel_header': PatternFill("solid", fgColor="0284C7"),  # Bleu ciel soutenu
        'hex_pastel': '#E0F2FE',
        'hex_header': '#0284C7',
        'hex_text': '#075985',
        'badge': '🔵 U11',
        'nom': 'U11 (Bleu ciel)'
    },
    'U13': {
        'bg_excel_pastel': PatternFill("solid", fgColor="F3E8FF"),  # Violet / Lavande doux
        'bg_excel_header': PatternFill("solid", fgColor="7C3AED"),  # Violet / Indigo soutenu
        'hex_pastel': '#F3E8FF',
        'hex_header': '#7C3AED',
        'hex_text': '#6B21A8',
        'badge': '🟣 U13',
        'nom': 'U13 (Violet)'
    },
    'U15': {
        'bg_excel_pastel': PatternFill("solid", fgColor="FCE7F3"),  # Rose doux
        'bg_excel_header': PatternFill("solid", fgColor="DB2777"),  # Rose soutenu
        'hex_pastel': '#FCE7F3',
        'hex_header': '#DB2777',
        'hex_text': '#9D174D',
        'badge': '🌸 U15',
        'nom': 'U15 (Rose)'
    },
    'U17': {
        'bg_excel_pastel': PatternFill("solid", fgColor="FFEDD5"),  # Orange pêche doux
        'bg_excel_header': PatternFill("solid", fgColor="EA580C"),  # Orange soutenu
        'hex_pastel': '#FFEDD5',
        'hex_header': '#EA580C',
        'hex_text': '#9A3412',
        'badge': '🟠 U17',
        'nom': 'U17 (Orange)'
    },
    'U20': {
        'bg_excel_pastel': PatternFill("solid", fgColor="E2E8F0"),  # Ardoise argenté doux
        'bg_excel_header': PatternFill("solid", fgColor="475569"),  # Ardoise soutenu
        'hex_pastel': '#E2E8F0',
        'hex_header': '#475569',
        'hex_text': '#1E293B',
        'badge': '⚪ U20',
        'nom': 'U20 (Ardoise)'
    },
    'SENIOR': {
        'bg_excel_pastel': PatternFill("solid", fgColor="CCFBF1"),  # Turquoise doux
        'bg_excel_header': PatternFill("solid", fgColor="0D9488"),  # Sarcelle soutenu
        'hex_pastel': '#CCFBF1',
        'hex_header': '#0D9488',
        'hex_text': '#115E59',
        'badge': '🔘 Senior',
        'nom': 'Senior (Turquoise)'
    },
    'AUTRE': {
        'bg_excel_pastel': PatternFill("solid", fgColor="F1F5F9"),  # Gris clair neutre
        'bg_excel_header': PatternFill("solid", fgColor="0055A4"),  # Bleu officiel FFLDA
        'hex_pastel': '#F1F5F9',
        'hex_header': '#0055A4',
        'hex_text': '#1E293B',
        'badge': '🥋 Autre',
        'nom': 'Général (Bleu)'
    }
}

def extraire_age_de_texte(texte):
    """
    Extrait la catégorie d'âge (U7, U9, U11, U13, etc.) depuis le libellé d'un match ou d'une poule.
    """
    if not texte:
        return 'AUTRE'
    txt = str(texte).upper()
    for age in ['U7', 'U9', 'U11', 'U13', 'U15', 'U17', 'U20', 'SENIOR']:
        if re.search(rf'\[\s*{age}\b|\b{age}\b', txt):
            return age
    return 'AUTRE'

def generer_html_grille_coloree(grille_lignes, nb_tapis):
    """
    Génère un tableau HTML stylé avec les couleurs de fond spécifiques à chaque catégorie d'âge.
    Parfait pour l'affichage interactif Streamlit et l'export HTML d'impression.
    """
    html = []
    html.append('''
    <div style="display:flex; flex-wrap:wrap; gap:10px; align-items:center; margin-bottom:14px; padding:10px 14px; background:#F8FAFC; border:1px solid #E2E8F0; border-radius:8px;">
        <span style="font-weight:bold; color:#1E293B; font-size:13px;">🎨 Légende des Catégories d'Âge :</span>
        <span style="background:#FEF3C7; color:#92400E; padding:3px 10px; border-radius:12px; font-weight:bold; font-size:12px; border:1px solid #FCD34D;">🟡 U7</span>
        <span style="background:#DCFCE7; color:#166534; padding:3px 10px; border-radius:12px; font-weight:bold; font-size:12px; border:1px solid #86EFAC;">🟢 U9</span>
        <span style="background:#E0F2FE; color:#075985; padding:3px 10px; border-radius:12px; font-weight:bold; font-size:12px; border:1px solid #7DD3FC;">🔵 U11</span>
        <span style="background:#F3E8FF; color:#6B21A8; padding:3px 10px; border-radius:12px; font-weight:bold; font-size:12px; border:1px solid #D8B4FE;">🟣 U13</span>
        <span style="background:#EF4135; color:#FFFFFF; padding:3px 10px; border-radius:12px; font-weight:bold; font-size:12px;">⏸️ Pause</span>
        <span style="background:#EFEFEF; color:#666666; padding:3px 10px; border-radius:12px; font-style:italic; font-size:12px;">⏳ Repos / Pesée</span>
    </div>
    ''')
    
    html.append('<div style="overflow-x:auto;"><table style="width:100%; border-collapse:collapse; font-family:Arial, sans-serif;">')
    html.append('<thead><tr>')
    for t in range(nb_tapis):
        html.append(f'<th style="background:#0055A4; color:#ffffff; padding:10px 8px; border:1px solid #CBD5E1; text-align:center; font-size:14px; font-weight:bold;">🥋 Tapis {t + 1}</th>')
    html.append('</tr></thead><tbody>')
    
    for row in grille_lignes:
        html.append('<tr>')
        for t in range(nb_tapis):
            col_name = f"Tapis {t + 1}"
            val = row.get(col_name, "")
            if not val:
                html.append('<td style="background:#FFFFFF; border:1px solid #E2E8F0; padding:8px;"></td>')
                continue
            val_s = str(val)
            if "PAUSE" in val_s:
                bg = "#EF4135"
                color = "#FFFFFF"
                border_c = "#DC2626"
                weight = "bold"
            elif any(k in val_s for k in ["Attente", "Pesée", "échauffement", "Repos"]):
                bg = "#F1F5F9"
                color = "#64748B"
                border_c = "#CBD5E1"
                weight = "normal"
            else:
                age = extraire_age_de_texte(val_s)
                cfg = COULEURS_AGE_GRILLE.get(age, COULEURS_AGE_GRILLE['AUTRE'])
                bg = cfg['hex_pastel']
                color = "#0F172A"
                border_c = cfg['hex_header']
                weight = "normal"
            
            cell_content = val_s.replace("\n", "<br>")
            html.append(f'<td style="background:{bg}; color:{color}; border:1px solid #CBD5E1; border-top:3px solid {border_c}; padding:8px 6px; text-align:center; font-size:12px; font-weight:{weight}; vertical-align:middle; line-height:1.4;">{cell_content}</td>')
        html.append('</tr>')
        
    html.append('</tbody></table></div>')
    return "\n".join(html)

# --- GÉNÉRATEUR DE DOCUMENTS HTML AUTONOMES POUR IMPRESSION PAYSAGE A4 ---
def generer_document_html_imprimable(titre, nom_comp, sections):
    """
    Génère un document HTML 100% autonome prêt pour l'impression A4 Paysage.
    Chaque section / onglet occupe sa propre page (page-break-after: always) et les tables ne sont jamais coupées.
    """
    html_sections = []
    for section_title, content in sections:
        if isinstance(content, pd.DataFrame):
            if "Grille Globale" in section_title or "Grille de Passage" in section_title:
                table_html = generer_html_grille_coloree(content.to_dict('records'), len(content.columns))
            else:
                table_html = content.to_html(index=False, classes="print-table")
        else:
            table_html = str(content)
        
        html_sections.append(f"""
        <div class="block-table">
            <h3 class="block-title">{section_title}</h3>
            {table_html}
        </div>
        """)
    
    sections_str = "\n".join(html_sections)
    
    html_doc = f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>{titre} - {nom_comp}</title>
    <style>
        @page {{
            size: landscape;
            margin: 10mm;
        }}
        body {{
            font-family: Arial, Helvetica, sans-serif;
            margin: 0;
            padding: 15px;
            background: #ffffff;
            color: #111;
        }}
        .header-print {{
            text-align: center;
            border-bottom: 3px solid #0055A4;
            padding-bottom: 12px;
            margin-bottom: 25px;
        }}
        .header-print h1 {{
            color: #0055A4;
            margin: 0 0 6px 0;
            font-size: 24px;
            text-transform: uppercase;
        }}
        .header-print p {{
            margin: 0;
            color: #555;
            font-size: 13px;
        }}
        .block-table {{
            page-break-after: always !important;
            break-after: page !important;
            page-break-inside: avoid !important;
            break-inside: avoid-page !important;
            margin-bottom: 30px;
            width: 100%;
            clear: both;
        }}
        .block-table:last-child {{
            page-break-after: auto !important;
            break-after: auto !important;
        }}
        .block-title {{
            background-color: #0055A4;
            color: #ffffff;
            padding: 8px 14px;
            font-size: 15px;
            font-weight: bold;
            border-radius: 4px 4px 0 0;
            margin: 0 0 5px 0;
        }}
        table, .print-table {{
            width: 100% !important;
            border-collapse: collapse;
            margin-top: 0;
            margin-bottom: 10px;
            page-break-inside: avoid !important;
            break-inside: avoid-page !important;
        }}
        th {{
            background-color: #EF4135;
            color: #ffffff;
            padding: 8px 12px;
            font-size: 13px;
            border: 1px solid #000;
            text-align: center;
        }}
        td {{
            padding: 7px 12px;
            font-size: 12px;
            border: 1px solid #ccc;
            text-align: center;
        }}
        tr:nth-child(even) {{
            background-color: #f8f9fa;
        }}
        .no-print-bar {{
            text-align: center;
            padding: 12px;
            background-color: #f0f4f8;
            border: 1px solid #d0d7de;
            border-radius: 8px;
            margin-bottom: 20px;
        }}
        .btn-imprimer {{
            background-color: #0055A4;
            color: white;
            font-size: 15px;
            font-weight: bold;
            padding: 10px 24px;
            border: none;
            border-radius: 6px;
            cursor: pointer;
            box-shadow: 0 3px 6px rgba(0,0,0,0.15);
        }}
        .btn-imprimer:hover {{
            background-color: #003f7d;
        }}
        @media print {{
            .no-print-bar {{
                display: none !important;
            }}
        }}
    </style>
</head>
<body onload="window.print()">
    <div class="no-print-bar">
        <button class="btn-imprimer" onclick="window.print()">🖨️ Imprimer le Document (Format Paysage A4)</button>
    </div>
    <div class="header-print">
        <h1>🏆 {nom_comp}</h1>
        <p><strong>{titre}</strong> — Document Officiel FFLDA — Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}</p>
    </div>
    {sections_str}
</body>
</html>"""
    return html_doc

# --- GÉNÉRATEUR DE DOCUMENTS PDF VECTORIELS DEPUIS LE CLASSEUR EXCEL OFFICIEL (A4 PORTRAIT) ---
def generer_pdf_depuis_classeur_excel(workbook_or_sheets, nom_competition="Tournoi FFLDA", wb=None):
    """
    Génère un fichier PDF vectoriel A4 Portrait à partir des feuilles du classeur Excel officiel FFLDA (openpyxl).
    Chaque feuille Excel (Poule nordique, Tableau éliminatoire, Poules croisées, Plateaux U7,
    Grille de Tapis, Planning) est fidèlement convertie en page(s) PDF en mode Portrait avec :
    - La disposition exacte des cellules et fusions (SPAN)
    - Les largeurs de colonnes proportionnelles adaptées au mode Portrait A4
    - Les couleurs de fond officielles FFLDA
    - Les bordures de cellules
    - Les styles de police (gras, tailles, alignements)
    - La résolution intelligente des formules et liens inter-onglets
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab import rl_config
    rl_config.allowTableBoundsErrors = 1

    wb_ref = wb
    if wb_ref is None:
        if hasattr(workbook_or_sheets, 'worksheets'):
            sheets = list(workbook_or_sheets.worksheets)
            wb_ref = workbook_or_sheets
        elif hasattr(workbook_or_sheets, 'sheets'):
            sheets = list(workbook_or_sheets.sheets.values()) if isinstance(workbook_or_sheets.sheets, dict) else list(workbook_or_sheets.sheets)
            wb_ref = getattr(workbook_or_sheets, 'book', None)
        elif isinstance(workbook_or_sheets, (list, tuple)):
            sheets = list(workbook_or_sheets)
            if sheets and hasattr(sheets[0], 'parent'):
                wb_ref = sheets[0].parent
        else:
            sheets = [workbook_or_sheets]
            if hasattr(workbook_or_sheets, 'parent'):
                wb_ref = workbook_or_sheets.parent
    else:
        if hasattr(workbook_or_sheets, 'worksheets'):
            sheets = list(workbook_or_sheets.worksheets)
        elif hasattr(workbook_or_sheets, 'sheets'):
            sheets = list(workbook_or_sheets.sheets.values()) if isinstance(workbook_or_sheets.sheets, dict) else list(workbook_or_sheets.sheets)
        elif isinstance(workbook_or_sheets, (list, tuple)):
            sheets = list(workbook_or_sheets)
        else:
            sheets = [workbook_or_sheets]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, 
        pagesize=A4, # Format A4 Portrait
        rightMargin=12, 
        leftMargin=12, 
        topMargin=12, 
        bottomMargin=12
    )
    styles = getSampleStyleSheet()
    page_width = A4[0] - 24   # 571.27 pt
    page_height = A4[1] - 24  # 817.89 pt

    def to_col_letter(col_idx):
        res = ""
        while col_idx > 0:
            col_idx, rem = divmod(col_idx - 1, 26)
            res = chr(65 + rem) + res
        return res

    def get_hex(openpyxl_color, default=None):
        if not openpyxl_color:
            return default
        rgb = getattr(openpyxl_color, 'rgb', None)
        if not rgb:
            return default
        rgb_s = str(rgb).strip()
        if rgb_s in ['00000000', '0', 'None', '', 'none']:
            return default
        if re.match(r'^[0-9a-fA-F]{6}$', rgb_s):
            return '#' + rgb_s
        elif re.match(r'^[0-9a-fA-F]{8}$', rgb_s):
            return '#' + rgb_s[2:]
        return default

    def safe_hex_color(hex_str, default='#CBD5E1'):
        try:
            if hex_str and re.match(r'^#[0-9a-fA-F]{6}$', str(hex_str)):
                return colors.HexColor(hex_str)
        except Exception:
            pass
        return colors.HexColor(default)

    def resolve_formula_cell(formula_str, current_ws, wb_reference, depth=0):
        if depth > 4:
            return ""
        s = str(formula_str).strip()
        if not s.startswith('='):
            return s
            
        # Détection podiums et classements
        if '🥇' in s or 'CHAMPION' in s or 'OR' in s:
            quoted = re.findall(r'"([^"]*)"', s)
            p_parts = [q for q in quoted if any(k in q for k in ['🥇', 'OR', 'CHAMPION', 'En attente', 'Vainqueur'])]
            if p_parts:
                return "\n".join(p_parts[:2])
        if '🥈' in s or 'ARGENT' in s or 'VICE' in s:
            quoted = re.findall(r'"([^"]*)"', s)
            p_parts = [q for q in quoted if any(k in q for k in ['🥈', 'ARGENT', 'VICE', 'En attente', 'Perdant'])]
            if p_parts:
                return "\n".join(p_parts[:2])
        if '🥉' in s or 'BRONZE' in s or 'Bronze' in s:
            quoted = re.findall(r'"([^"]*)"', s)
            p_parts = [q for q in quoted if any(k in q for k in ['🥉', 'BRONZE', 'Bronze', 'En attente', 'Vainqueur'])]
            if p_parts:
                return "\n".join(p_parts[:2])

        quoted = re.findall(r'"([^"]*)"', s)
        priority_keywords = ['🔴', '🔵', 'Vainqueur', 'Perdant', 'Qualifié', 'Repêché', 'TOUR', 'COMBAT', 'Plateau']
        for q in quoted:
            if any(k in q for k in priority_keywords):
                return q
                
        try:
            m_ref = re.search(r"(?:'([^']+)'|([A-Za-z0-9_]+))!([A-Z]+[0-9]+)", s)
            if m_ref and wb_reference:
                target_sheet_name = m_ref.group(1) or m_ref.group(2)
                coord = m_ref.group(3)
                if hasattr(wb_reference, 'sheetnames') and target_sheet_name in wb_reference.sheetnames:
                    cell_obj = wb_reference[target_sheet_name][coord]
                    t_val = getattr(cell_obj, 'value', None)
                    res = resolve_formula_cell(t_val, wb_reference[target_sheet_name], wb_reference, depth + 1)
                    if res:
                        return res.replace('🔴 ', '').replace('🔵 ', '').strip()
        except Exception:
            pass
                    
        try:
            m_local = re.match(r"^=([A-Z]+[0-9]+)$", s)
            if m_local and current_ws:
                coord = m_local.group(1)
                cell_obj = current_ws[coord]
                t_val = getattr(cell_obj, 'value', None)
                return resolve_formula_cell(t_val, current_ws, wb_reference, depth + 1)
        except Exception:
            pass
            
        if quoted:
            for q in quoted:
                if q.strip() and q not in [" ", "🔴 ", "🔵 "]:
                    return q
        return ""

    def clean_val(val, current_ws):
        if val is None:
            return ""
        s = str(val).strip()
        if s.startswith('='):
            return resolve_formula_cell(s, current_ws, wb_ref)
        return s

    def cell_has_content(c_obj):
        if c_obj.value not in [None, '']:
            return True
        f = getattr(c_obj, 'fill', None)
        if getattr(f, 'fill_type', None) in ['solid']:
            c_hex = get_hex(getattr(f, 'fgColor', None))
            if c_hex and c_hex.upper() not in ['#FFFFFF', '#00000000']:
                return True
        b = getattr(c_obj, 'border', None)
        if b and (getattr(b.left, 'style', None) or getattr(b.top, 'style', None) or getattr(b.right, 'style', None) or getattr(b.bottom, 'style', None)):
            return True
        return False

    story = []
    sheet_has_pages = False

    for sheet_idx, ws in enumerate(sheets):
        if not hasattr(ws, 'cell'):
            continue
        max_r = ws.max_row or 1
        max_c = ws.max_column or 1
        
        min_allowed_r = 1
        min_allowed_c = 1
        for mr in list(ws.merged_cells.ranges):
            if mr.max_row > min_allowed_r:
                min_allowed_r = min(mr.max_row, max_r)
            if mr.max_col > min_allowed_c:
                min_allowed_c = min(mr.max_col, max_c)
                
        while max_r > min_allowed_r and all(not cell_has_content(ws.cell(row=max_r, column=c)) for c in range(1, max_c + 1)):
            max_r -= 1
        while max_c > min_allowed_c and all(not cell_has_content(ws.cell(row=r, column=max_c)) for r in range(1, max_r + 1)):
            max_c -= 1
            
        if max_r == 1 and max_c == 1 and not cell_has_content(ws.cell(1, 1)):
            continue

        col_widths = []
        for c in range(1, max_c + 1):
            col_letter = to_col_letter(c)
            w = ws.column_dimensions[col_letter].width if col_letter in ws.column_dimensions else None
            try:
                w_val = float(w) if (w is not None and str(w).strip() != '') else 10.0
                col_widths.append(w_val if w_val > 0 else 10.0)
            except (ValueError, TypeError):
                col_widths.append(10.0)
            
        total_w = sum(col_widths)
        scale = page_width / total_w if total_w > 0 else 1.0
        scaled_widths = [max(cw * scale, 5.0) for cw in col_widths]
        sum_sw = sum(scaled_widths)
        if sum_sw > page_width:
            scaled_widths = [sw * (page_width / sum_sw) for sw in scaled_widths]

        data = []
        t_styles = [
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 1.5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 1.5),
            ('TOPPADDING', (0, 0), (-1, -1), 1.0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 1.0),
        ]

        base_scale = 0.50 if max_c > 14 else (0.60 if max_c > 9 else 0.70)

        for r in range(1, max_r + 1):
            row_cells = []
            for c in range(1, max_c + 1):
                cell = ws.cell(row=r, column=c)
                raw = cell.value
                txt = clean_val(raw, ws)
                
                font = cell.font
                is_bold = bool(font.bold) if font else False
                orig_f_size = 10
                if font and getattr(font, 'size', None) is not None:
                    try:
                        orig_f_size = float(font.size)
                    except (ValueError, TypeError):
                        orig_f_size = 10
                
                if r == 1:
                    f_size = max(9.5, min(int(orig_f_size * 0.9), 13))
                elif r == 2:
                    f_size = max(7.0, min(int(orig_f_size * 0.85), 9))
                else:
                    f_size = max(4.5, min(int(orig_f_size * base_scale), 11))
                
                f_color = get_hex(getattr(font, 'color', None), default='#1E293B')
                if not f_color or f_color in ['#00000000', '#000000']:
                    f_color = '#1E293B'
                    
                align_obj = cell.alignment
                h_align = getattr(align_obj, 'horizontal', 'center') or 'center'
                align_code = 1 # Center
                if h_align == 'left': align_code = 0
                elif h_align == 'right': align_code = 2
                
                fill_obj = getattr(cell, 'fill', None)
                if getattr(fill_obj, 'fill_type', None) in ['solid', 'lightGrid', 'darkGrid']:
                    bg_hex = get_hex(getattr(fill_obj, 'fgColor', None))
                    if bg_hex and bg_hex.upper() not in ['#FFFFFF', '#00000000']:
                        t_styles.append(('BACKGROUND', (c - 1, r - 1), (c - 1, r - 1), safe_hex_color(bg_hex)))
                        if bg_hex.upper() in ['#0055A4', '#EF4135', '#E53935', '#000000', '#334155', '#475569', '#1E88E5']:
                            f_color = '#FFFFFF'
                
                if (not fill_obj or not getattr(fill_obj, 'fill_type', None)) and f_color == '#FFFFFF':
                    f_color = '#1E293B'
                    
                b = cell.border
                if b and (getattr(b.left, 'style', None) or getattr(b.top, 'style', None) or getattr(b.right, 'style', None) or getattr(b.bottom, 'style', None)):
                    b_col = get_hex(getattr(b.top, 'color', None), default='#CBD5E1')
                    t_styles.append(('BOX', (c - 1, r - 1), (c - 1, r - 1), 0.5, safe_hex_color(b_col)))
                    
                p_style = ParagraphStyle(
                    f'S_{sheet_idx}_{r}_{c}',
                    parent=styles['Normal'],
                    fontName='Helvetica-Bold' if is_bold else 'Helvetica',
                    fontSize=f_size,
                    leading=f_size + 1.2,
                    textColor=safe_hex_color(f_color, default='#1E293B'),
                    alignment=align_code
                )
                
                txt_clean = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', txt)
                txt_safe = txt_clean.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('\n', '<br/>')
                try:
                    p = Paragraph(txt_safe, p_style) if txt_safe else Paragraph("&nbsp;", p_style)
                except Exception:
                    p = Paragraph("&nbsp;", p_style)
                row_cells.append(p)
            data.append(row_cells)

        spanned_cells = set()
        for m_range in list(ws.merged_cells.ranges):
            min_col, min_row, max_col_r, max_row_r = m_range.min_col, m_range.min_row, m_range.max_col, m_range.max_row
            if min_row <= max_r and min_col <= max_c:
                end_c = min(max_col_r, max_c)
                end_r = min(max_row_r, max_r)
                if (end_c > min_col or end_r > min_row) and min_col >= 1 and min_row >= 1:
                    c1 = min_col - 1
                    r1 = min_row - 1
                    c2 = min(end_c - 1, len(data[0]) - 1)
                    r2 = min(end_r - 1, len(data) - 1)
                    if c2 >= c1 and r2 >= r1 and (c2 > c1 or r2 > r1):
                        span_coords = [(col_k, row_k) for col_k in range(c1, c2 + 1) for row_k in range(r1, r2 + 1)]
                        if not any(coord in spanned_cells for coord in span_coords):
                            t_styles.append(('SPAN', (c1, r1), (c2, r2)))
                            for coord in span_coords:
                                spanned_cells.add(coord)

        # Auto-fusion horizontale pour lignes d'en-tête / bannières non fusionnées
        for r_chk in range(1, min(max_r + 1, 5)):
            r_idx = r_chk - 1
            if (0, r_idx) not in spanned_cells:
                v1 = ws.cell(row=r_chk, column=1).value
                if v1 and str(v1).strip():
                    last_empty_c = 1
                    for c_chk in range(2, max_c + 1):
                        if (c_chk - 1, r_idx) in spanned_cells or cell_has_content(ws.cell(row=r_chk, column=c_chk)):
                            break
                        last_empty_c = c_chk
                    if last_empty_c > 1:
                        c2_span = last_empty_c - 1
                        span_coords = [(col_k, r_idx) for col_k in range(0, c2_span + 1)]
                        if not any(coord in spanned_cells for coord in span_coords):
                            t_styles.append(('SPAN', (0, r_idx), (c2_span, r_idx)))
                            for coord in span_coords:
                                spanned_cells.add(coord)
                
        safe_styles = []
        for cmd in t_styles:
            try:
                op = cmd[0]
                if op in ('BACKGROUND', 'BOX', 'SPAN'):
                    sc1, sr1 = cmd[1]
                    sc2, sr2 = cmd[2]
                    if 0 <= sc1 < max_c and 0 <= sc2 < max_c and 0 <= sr1 < max_r and 0 <= sr2 < max_r:
                        if sc2 >= sc1 and sr2 >= sr1:
                            safe_styles.append(cmd)
                else:
                    safe_styles.append(cmd)
            except Exception:
                pass

        rep_rows = 2 if ('tapis' in ws.title.lower() or 'passage' in ws.title.lower()) and max_r > 4 else 0
        try:
            table = Table(data, colWidths=scaled_widths, repeatRows=rep_rows, splitByRow=1)
            table.setStyle(TableStyle(safe_styles))
        except Exception:
            table = Table(data, colWidths=scaled_widths, splitByRow=1)
            table.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ]))

        if sheet_has_pages:
            story.append(PageBreak())
        story.append(table)
        sheet_has_pages = True

    try:
        doc.build(story)
        return buffer.getvalue()
    except Exception:
        # Fallback d'urgence
        try:
            buf_fb = io.BytesIO()
            doc_fb = SimpleDocTemplate(buf_fb, pagesize=A4, rightMargin=12, leftMargin=12, topMargin=12, bottomMargin=12)
            fb_story = []
            for item in story:
                if isinstance(item, Table):
                    fb_t = Table(item._cellvalues, colWidths=item._colWidths, splitByRow=1)
                    fb_t.setStyle(TableStyle([
                        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                    ]))
                    fb_story.append(fb_t)
                else:
                    fb_story.append(item)
            doc_fb.build(fb_story)
            return buf_fb.getvalue()
        except Exception:
            return None

# --- GÉNÉRATEUR DE DOCUMENTS PDF VECTORIELS (REPORTLAB - A4 PORTRAIT - 1 PAGE PAR ONGLET) ---
def generer_pdf_tournoi_complet(titre, nom_comp, sections):
    """
    Génère un fichier PDF vectoriel (A4 Portrait) prêt à imprimer et télécharger.
    Chaque onglet / section commence sur une nouvelle page (PageBreak) et aucun tableau n'est coupé.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, 
        pagesize=A4, 
        rightMargin=20, 
        leftMargin=20, 
        topMargin=20, 
        bottomMargin=20
    )
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'PDFTitle', 
        parent=styles['Heading1'], 
        fontName='Helvetica-Bold', 
        fontSize=16, 
        textColor=colors.HexColor('#0055A4'), 
        alignment=1,
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'PDFSubtitle', 
        parent=styles['Normal'], 
        fontName='Helvetica', 
        fontSize=10, 
        textColor=colors.HexColor('#555555'), 
        alignment=1,
        spaceAfter=12
    )

    sec_banner_style = ParagraphStyle(
        'SecBanner', 
        parent=styles['Heading2'], 
        fontName='Helvetica-Bold', 
        fontSize=12, 
        textColor=colors.white, 
        backColor=colors.HexColor('#0055A4'), 
        borderPadding=6,
        spaceAfter=10,
        alignment=0
    )
    
    cell_head_style = ParagraphStyle(
        'CellHead', 
        parent=styles['Normal'], 
        fontName='Helvetica-Bold', 
        fontSize=9, 
        textColor=colors.white, 
        alignment=1
    )
    
    cell_body_style = ParagraphStyle(
        'CellBody', 
        parent=styles['Normal'], 
        fontName='Helvetica', 
        fontSize=8, 
        textColor=colors.black, 
        alignment=1
    )

    story = []
    page_width = A4[0] - 40  # 555.27 pt

    for idx, (sec_title, content) in enumerate(sections):
        if idx > 0:
            story.append(PageBreak())
        
        story.append(Paragraph(f"🏆 {nom_comp.upper()}", title_style))
        story.append(Paragraph(f"<b>{titre}</b> — Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", subtitle_style))
        story.append(Paragraph(f"<b>{sec_title}</b>", sec_banner_style))
        story.append(Spacer(1, 8))
        
        if isinstance(content, pd.DataFrame):
            df = content
            if df.empty:
                continue
            
            headers = [Paragraph(str(col), cell_head_style) for col in df.columns]
            data = [headers]
            
            for _, row in df.iterrows():
                r_cells = []
                for val in row:
                    txt = str(val).replace('\n', '<br/>') if pd.notna(val) else ''
                    r_cells.append(Paragraph(txt, cell_body_style))
                data.append(r_cells)
            
            nb_cols = len(df.columns)
            col_w = page_width / nb_cols if nb_cols > 0 else page_width
            col_widths = [col_w] * nb_cols
            
            table_obj = Table(data, colWidths=col_widths, repeatRows=1)
            t_styles = [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0055A4')),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]
            if "Grille Globale" in sec_title or "Grille de Passage" in sec_title:
                for r_idx, (_, row_data) in enumerate(df.iterrows(), 1):
                    for c_idx, val in enumerate(row_data):
                        if pd.notna(val) and val:
                            val_str = str(val)
                            if "PAUSE" in val_str:
                                t_styles.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), colors.HexColor('#EF4135')))
                            elif any(k in val_str for k in ["Attente", "Pesée", "échauffement", "Repos"]):
                                t_styles.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), colors.HexColor('#EFEFEF')))
                            else:
                                age_k = extraire_age_de_texte(val_str)
                                hex_col = COULEURS_AGE_GRILLE.get(age_k, COULEURS_AGE_GRILLE['AUTRE'])['hex_pastel']
                                t_styles.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), colors.HexColor(hex_col)))
            else:
                t_styles.append(('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]))
            table_obj.setStyle(TableStyle(t_styles))
            story.append(KeepTogether(table_obj))
        else:
            story.append(Paragraph(str(content), cell_body_style))

    doc.build(story)
    return buffer.getvalue()

# --- STYLES D'IMPRESSION DIRECTE (CSS @media print) ---
st.markdown("""
    <style>
    @media print {
        @page {
            size: landscape;
            margin: 8mm;
        }

        html, body, .stApp, .main, .block-container, div[data-testid="stMain"], div[data-testid="stBlock"] {
            width: 100% !important;
            max-width: 100% !important;
            margin: 0 !important;
            padding: 0 !important;
            background: white !important;
            color: black !important;
            overflow: visible !important;
        }

        table, .stTable, div[data-testid="stTable"], .stDataFrame, div[data-testid="stDataFrame"] {
            width: 100% !important;
            max-width: 100% !important;
            table-layout: auto !important;
            page-break-inside: avoid !important;
            break-inside: avoid-page !important;
            break-inside: avoid !important;
        }

        tr, tbody, thead {
            page-break-inside: avoid !important;
            break-inside: avoid-page !important;
            break-inside: avoid !important;
        }

        section[data-testid="stSidebar"], 
        header, 
        footer, 
        button, 
        .stButton,
        iframe,
        div[data-baseweb="tab-list"] {
            display: none !important;
        }
    }
    </style>
""", unsafe_allow_html=True)

def bouton_imprimer(html_data=None, filename="Fiche_Impression_Paysage.html", label="🖨️ Imprimer / Télécharger Fiche Paysage (HTML)", key=None):
    if html_data:
        st.download_button(
            label=label,
            data=html_data,
            file_name=filename,
            mime="text/html",
            key=key
        )
    else:
        print_code = f"""
            <button onclick="window.parent.print()" style="
                background-color: #0055A4; 
                color: white; 
                border: none; 
                padding: 10px 18px; 
                font-size: 14px; 
                font-weight: bold; 
                border-radius: 6px; 
                cursor: pointer;
                box-shadow: 0px 2px 5px rgba(0,0,0,0.2);
            ">
                {label}
            </button>
        """
        components.html(print_code, height=50)

# Sélection du mode de travail
col_mode1, col_mode2 = st.columns([2, 3])
with col_mode1:
    mode_app = st.selectbox("📌 Sélectionnez le mode de travail", ["1. Générer un Tournoi (Planning & Poules)", "2. Importer les scores & Éditer le Bilan (Excel)"])

if mode_app.startswith("2"):
    st.markdown("### 📂 Module de Fin de Tournoi & Bilans Fédéraux")
    st.info(f"Importez votre fichier Excel complété pour la compétition **{nom_competition}** afin de générer automatiquement les classements officiels individuels et par club.")
    
    fichier_resultats = st.file_uploader("Sélectionner le fichier Excel complété (.xlsx)", type=["xlsx"])
    
    if fichier_resultats is not None:
        try:
            wb_res = openpyxl.load_workbook(fichier_resultats, data_only=True)
            st.success("✨ Fichier de résultats analysé avec succès !")
            
            tous_les_resultats = []
            onglets_poules = [f for f in wb_res.sheetnames if not any(x in f for x in ["Résumé", "Grille", "Classement Général"])]
            onglets_poules.sort(key=lambda x: (0 if "U7" in x.upper() else (1 if "U9" in x.upper() else (2 if "U11" in x.upper() else 3)), x))
            
            for nom_feuille in onglets_poules:
                ws = wb_res[nom_feuille]
                
                # --- REPÉRAGE DYNAMIQUE DES COLONNES (Ligne 4 - En-têtes) ---
                col_total_pts = None
                col_poids = None
                col_comite = None
                for c_idx in range(1, ws.max_column + 1):
                    val_head = str(ws.cell(row=4, column=c_idx).value or "").strip()
                    val_lower = val_head.lower()
                    if val_head == "Total Pts":
                        col_total_pts = c_idx
                    elif val_head == "Poids":
                        col_poids = c_idx
                    elif any(k in val_lower for k in ["comité", "comite", "ligue", "région", "region", "c.r."]):
                        col_comite = c_idx
                
                r = 5
                lutteurs_poule = []
                while ws.cell(row=r, column=3).value is not None:
                    nom = ws.cell(row=r, column=3).value
                    club = ws.cell(row=r, column=4).value
                    comite_val = ws.cell(row=r, column=col_comite).value if col_comite else None
                    comite = str(comite_val).strip() if (comite_val and str(comite_val).strip() not in ["", "None", "nan", "-"]) else "Comité Non Renseigné"
                    
                    total_pts = ws.cell(row=r, column=col_total_pts).value if col_total_pts else 0
                    poids_raw = ws.cell(row=r, column=col_poids).value if col_poids else 0
                    
                    try:
                        pts_val = int(round(float(total_pts))) if total_pts is not None else 0
                    except (ValueError, TypeError):
                        pts_val = 0

                    def formater_poids(val):
                        if val is None or str(val).strip() in ['', 'None', 'nan']:
                            return ''
                        try:
                            val_float = float(str(val).replace(',', '.'))
                            if val_float == int(val_float):
                                return int(val_float)
                            else:
                                return round(val_float, 1)
                        except (ValueError, TypeError):
                            return val

                    poids_val = formater_poids(poids_raw)

                    clt_excel_raw = ws.cell(row=r, column=1).value
                    try:
                        clt_excel = int(float(str(clt_excel_raw).strip())) if clt_excel_raw is not None and str(clt_excel_raw).strip() not in ['', 'NR', 'None'] else None
                    except (ValueError, TypeError):
                        clt_excel = None

                    lutteurs_poule.append({
                        "Poule": nom_feuille,
                        "Nom": nom,
                        "Club": club if club else "Indépendant",
                        "Comité": comite,
                        "Poids": poids_val,
                        "Points": pts_val,
                        "Clt_Excel": clt_excel
                    })
                    r += 1
                
                df_poule = pd.DataFrame(lutteurs_poule)
                if not df_poule.empty:
                    if df_poule["Points"].sum() == 0:
                        df_poule["Clt"] = "NR"
                    elif all(p.get("Clt_Excel") is not None for p in lutteurs_poule):
                        # Classement officiel calculé par la feuille Excel avec départage direct FFLDA
                        df_poule["Clt"] = [p["Clt_Excel"] for p in lutteurs_poule]
                        df_poule = df_poule.sort_values(by="Clt").reset_index(drop=True)
                    else:
                        df_poule = df_poule.sort_values(by="Points", ascending=False).reset_index(drop=True)
                        rangs = []
                        current_rang = 1
                        for idx, row in df_poule.iterrows():
                            if idx > 0 and row["Points"] == df_poule.iloc[idx-1]["Points"]:
                                rangs.append(rangs[-1])
                            else:
                                rangs.append(current_rang)
                            current_rang += 1
                        df_poule["Clt"] = rangs
                        
                    tous_les_resultats.extend(df_poule.to_dict('records'))

            df_bilan = pd.DataFrame(tous_les_resultats)
            
            # --- CALCUL DU CLASSEMENT DES CLUBS ET DES COMITÉS RÉGIONAUX ---
            bareme_points = {1: 4, 2: 3, 3: 2, 4: 1}
            points_clubs = {}
            points_comites = {}

            for _, row in df_bilan.iterrows():
                club = str(row.get("Club", "")).strip()
                comite = str(row.get("Comité", "Comité Non Renseigné")).strip()
                if not comite or comite in ["None", "nan", "-"]:
                    comite = "Comité Non Renseigné"
                
                if club and club not in ["", "-", "None", "nan"] and club not in points_clubs:
                    points_clubs[club] = {"Club": club, "Points Club": 0, "1ers": 0, "2èmes": 0, "3èmes": 0, "4èmes": 0}
                if comite and comite not in points_comites:
                    points_comites[comite] = {"Comité Régional": comite, "Points Comité": 0, "1ers": 0, "2èmes": 0, "3èmes": 0, "4èmes": 0}
                
                clt = str(row.get("Clt", "NR")).strip()
                if clt == "NR" or not clt.isdigit():
                    continue
                
                clt_num = int(clt)
                pts_attribués = bareme_points.get(clt_num, 0)
                
                # Ranking Clubs
                if club in points_clubs:
                    points_clubs[club]["Points Club"] += pts_attribués
                    if clt_num == 1: points_clubs[club]["1ers"] += 1
                    elif clt_num == 2: points_clubs[club]["2èmes"] += 1
                    elif clt_num == 3: points_clubs[club]["3èmes"] += 1
                    elif clt_num == 4: points_clubs[club]["4èmes"] += 1

                # Ranking Comités Régionaux
                if comite in points_comites:
                    points_comites[comite]["Points Comité"] += pts_attribués
                    if clt_num == 1: points_comites[comite]["1ers"] += 1
                    elif clt_num == 2: points_comites[comite]["2èmes"] += 1
                    elif clt_num == 3: points_comites[comite]["3èmes"] += 1
                    elif clt_num == 4: points_comites[comite]["4èmes"] += 1

            if points_clubs:
                df_clubs = pd.DataFrame(list(points_clubs.values())).sort_values(
                    by=["Points Club", "1ers", "2èmes", "3èmes", "4èmes"], 
                    ascending=False
                ).reset_index(drop=True)
                df_clubs.index = range(1, len(df_clubs) + 1)
                df_clubs.insert(0, "Clt Club", df_clubs.index)
            else:
                df_clubs = pd.DataFrame(columns=["Clt Club", "Club", "Points Club", "1ers", "2èmes", "3èmes", "4èmes"])

            if points_comites:
                df_comites = pd.DataFrame(list(points_comites.values())).sort_values(
                    by=["Points Comité", "1ers", "2èmes", "3èmes", "4èmes"], 
                    ascending=False
                ).reset_index(drop=True)
                df_comites.index = range(1, len(df_comites) + 1)
                df_comites.insert(0, "Clt Comité", df_comites.index)
            else:
                df_comites = pd.DataFrame(columns=["Clt Comité", "Comité Régional", "Points Comité", "1ers", "2èmes", "3èmes", "4èmes"])

            # --- GÉNÉRATION DU DOCUMENT HTML PAYSAGE POUR IMPRESSION DES BILANS ---
            tab_bilan_1, tab_bilan_2, tab_bilan_3 = st.tabs([
                "🏆 Classements Individuels (U7 / U9 / U11 / U13)", 
                "🛡️ Classement Général des Clubs",
                "🏛️ Classement des Comités Régionaux"
            ])
            
            with tab_bilan_1:
                st.subheader("Classements Individuels Officiels")
                for poule in df_bilan['Poule'].unique():
                    st.markdown(f"#### 🤼 {poule}")
                    sous_df = df_bilan[df_bilan['Poule'] == poule][['Clt', 'Nom', 'Club', 'Poids', 'Points']].copy()
                    sous_df['Points'] = sous_df['Points'].astype(int)
                    st.table(sous_df)

            with tab_bilan_2:
                st.subheader("🛡️ Podium des Clubs Engagés")
                st.markdown("*Barème officiel : 1er = 4 pts | 2ème = 3 pts | 3ème = 2 pts | 4ème = 1 pt*")
                st.table(df_clubs)

            with tab_bilan_3:
                st.subheader("🏛️ Classement Officiel des Comités Régionaux")
                st.markdown("*Barème officiel : 1er = 4 pts | 2ème = 3 pts | 3ème = 2 pts | 4ème = 1 pt*")
                st.table(df_comites)

            output_bilan = io.BytesIO()
            with pd.ExcelWriter(output_bilan, engine='openpyxl') as writer:
                df_clubs.to_excel(writer, sheet_name="Classement Clubs", index=False, startrow=5)
                df_comites.to_excel(writer, sheet_name="Classement Comités", index=False, startrow=5)
                ws_indiv = writer.book.create_sheet("Classements Individuels")
                
                bleu_fflda = PatternFill("solid", fgColor="0055A4")
                rouge_fflda = PatternFill("solid", fgColor="EF4135")
                gris_zebrage = PatternFill("solid", fgColor="F2F5F8")
                fond_blanc = PatternFill("solid", fgColor="FFFFFF")
                or_fill = PatternFill("solid", fgColor="FFF2CC")
                argent_fill = PatternFill("solid", fgColor="EFEFEF")
                bronze_fill = PatternFill("solid", fgColor="F8CBAD")
                
                font_titre = Font(name="Arial", size=15, bold=True, color="0055A4")
                font_section = Font(name="Arial", size=12, bold=True, color="FFFFFF")
                font_entete = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                font_data = Font(name="Arial", size=11, color="000000")
                font_data_bold = Font(name="Arial", size=11, bold=True, color="0055A4")
                b_fin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), 
                               top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
                
                # Feuille Classement Clubs
                ws_clubs = writer.sheets["Classement Clubs"]
                ws_clubs.views.sheetView[0].showGridLines = True
                ws_clubs.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()}").font = font_titre
                ws_clubs.cell(row=2, column=1, value="🛡️ CLASSEMENT OFFICIEL DES CLUBS - FFLDA").font = Font(name="Arial", size=12, bold=True, color="666666")
                ws_clubs.cell(row=3, column=1, value=f"Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}").font = Font(name="Arial", size=9, italic=True, color="888888")
                
                for col_idx in range(1, len(df_clubs.columns) + 1):
                    cell = ws_clubs.cell(row=5, column=col_idx)
                    cell.fill, cell.font, cell.alignment = bleu_fflda, font_entete, Alignment(horizontal="center", vertical="center")
                    ws_clubs.row_dimensions[5].height = 25
                
                for row_idx in range(6, ws_clubs.max_row + 1):
                    ws_clubs.row_dimensions[row_idx].height = 22
                    is_even = (row_idx % 2 == 0)
                    for col_idx in range(1, len(df_clubs.columns) + 1):
                        cell = ws_clubs.cell(row=row_idx, column=col_idx)
                        cell.border, cell.font = b_fin, font_data
                        cell.fill = gris_zebrage if is_even else fond_blanc
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                ws_clubs.column_dimensions['A'].width = 12
                ws_clubs.column_dimensions['B'].width = 30
                ws_clubs.column_dimensions['C'].width = 15
                ws_clubs.column_dimensions['D'].width = 12
                ws_clubs.column_dimensions['E'].width = 12
                ws_clubs.column_dimensions['F'].width = 12
                ws_clubs.column_dimensions['G'].width = 12

                # Feuille Classement Comités Régionaux
                ws_comites = writer.sheets["Classement Comités"]
                ws_comites.views.sheetView[0].showGridLines = True
                ws_comites.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()}").font = font_titre
                ws_comites.cell(row=2, column=1, value="🏛️ CLASSEMENT OFFICIEL DES COMITÉS RÉGIONAUX - FFLDA").font = Font(name="Arial", size=12, bold=True, color="666666")
                ws_comites.cell(row=3, column=1, value=f"Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}").font = Font(name="Arial", size=9, italic=True, color="888888")
                
                for col_idx in range(1, len(df_comites.columns) + 1):
                    cell = ws_comites.cell(row=5, column=col_idx)
                    cell.fill, cell.font, cell.alignment = bleu_fflda, font_entete, Alignment(horizontal="center", vertical="center")
                    ws_comites.row_dimensions[5].height = 25
                
                for row_idx in range(6, ws_comites.max_row + 1):
                    ws_comites.row_dimensions[row_idx].height = 22
                    is_even = (row_idx % 2 == 0)
                    for col_idx in range(1, len(df_comites.columns) + 1):
                        cell = ws_comites.cell(row=row_idx, column=col_idx)
                        cell.border, cell.font = b_fin, font_data
                        cell.fill = gris_zebrage if is_even else fond_blanc
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                ws_comites.column_dimensions['A'].width = 12
                ws_comites.column_dimensions['B'].width = 30
                ws_comites.column_dimensions['C'].width = 15
                ws_comites.column_dimensions['D'].width = 12
                ws_comites.column_dimensions['E'].width = 12
                ws_comites.column_dimensions['F'].width = 12
                ws_comites.column_dimensions['G'].width = 12

                ws_indiv.views.sheetView[0].showGridLines = True
                ws_indiv.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()}").font = font_titre
                ws_indiv.cell(row=2, column=1, value="🏆 CLASSEMENTS INDIVIDUELS OFFICIELS - FFLDA").font = Font(name="Arial", size=12, bold=True, color="666666")
                ws_indiv.cell(row=3, column=1, value=f"Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}").font = Font(name="Arial", size=9, italic=True, color="888888")
                
                row_cursor = 5
                headers_indiv = ["Clt", "NOM Prénom", "CLUB", "POIDS (kg)", "POINTS"]
                
                for poule in df_bilan['Poule'].unique():
                    groupe = df_bilan[df_bilan['Poule'] == poule]
                    ws_indiv.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=5)
                    cell_cat = ws_indiv.cell(row=row_cursor, column=1, value=f"  CATÉGORIE / POULE : {poule}")
                    cell_cat.fill, cell_cat.font, cell_cat.alignment = bleu_fflda, font_section, Alignment(horizontal="left", vertical="center")
                    ws_indiv.row_dimensions[row_cursor].height = 28
                    row_cursor += 1
                    
                    for col_idx, h in enumerate(headers_indiv, 1):
                        cell = ws_indiv.cell(row=row_cursor, column=col_idx, value=h)
                        cell.fill, cell.font, cell.alignment = rouge_fflda, font_entete, Alignment(horizontal="center", vertical="center")
                        ws_indiv.row_dimensions[row_cursor].height = 22
                    row_cursor += 1
                    
                    for _, row in groupe.iterrows():
                        current_row = row_cursor
                        ws_indiv.row_dimensions[current_row].height = 20
                        
                        c1 = ws_indiv.cell(row=current_row, column=1, value=row['Clt'])
                        c2 = ws_indiv.cell(row=current_row, column=2, value=row['Nom'])
                        c3 = ws_indiv.cell(row=current_row, column=3, value=row['Club'])
                        c4 = ws_indiv.cell(row=current_row, column=4, value=row['Poids'])
                        c5 = ws_indiv.cell(row=current_row, column=5, value=row['Points'])
                        
                        c1.font = font_data_bold
                        c2.font = font_data
                        c3.font = font_data
                        c4.font = font_data
                        c5.font = font_data_bold
                        
                        for c in [c1, c2, c3, c4, c5]:
                            c.border = b_fin
                            c.alignment = Alignment(horizontal="center", vertical="center")
                        c2.alignment = Alignment(horizontal="left", vertical="center")
                        
                        if row['Clt'] == 1: c1.fill = or_fill
                        elif row['Clt'] == 2: c1.fill = argent_fill
                        elif row['Clt'] == 3: c1.fill = bronze_fill
                        
                        row_cursor += 1
                    row_cursor += 2
                
                ws_indiv.column_dimensions['A'].width = 10
                ws_indiv.column_dimensions['B'].width = 30
                ws_indiv.column_dimensions['C'].width = 25
                ws_indiv.column_dimensions['D'].width = 15
                ws_indiv.column_dimensions['E'].width = 12

                # Configuration Impression Paysage A4 Excel
                for ws in writer.book.worksheets:
                    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
                    ws.page_setup.paperSize = ws.PAPERSIZE_A4
                    ws.sheet_properties.pageSetUpPr.fitToPage = True
                    ws.page_setup.fitToWidth = 1
                    ws.page_setup.fitToHeight = 0

            st.markdown("---")
            st.download_button(
                label="📥 Télécharger le Bilan Officiel FFLDA (Excel)",
                data=output_bilan.getvalue(),
                file_name="Bilan_Officiel_FFLDA.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"Erreur lors de l'analyse du fichier : {e}")

else:
    # --- MODE 1 : GÉNÉRATION DE TOURNOI ---
    col_up1, col_up2 = st.columns([1, 1])
    with col_up1:
        fichier_upload = st.file_uploader("📂 Importez votre liste d'inscrits (.csv ou .xlsx)", type=["xlsx", "csv"])
    with col_up2:
        fichier_arbitres_upload = st.file_uploader("🛡️ (Optionnel) Importez la liste des arbitres (.xlsx ou .csv)", type=["xlsx", "csv"], key="upload_arbitres")

    if fichier_upload is None:
        st.info("👈 Veuillez importer un fichier de participants (format Exalto .csv ou .xlsx) pour lancer l'optimisation des poules et plannings.")
    else:
        try:
            liste_arbitres = charger_liste_arbitres(fichier_arbitres_upload)
            tapis_arbitres = {t: [] for t in range(nb_tapis)}
            if liste_arbitres:
                for idx_arb, arb in enumerate(liste_arbitres):
                    tapis_arbitres[idx_arb % nb_tapis].append(arb)
            if fichier_upload.name.endswith('.csv'):
                df_raw = pd.read_csv(fichier_upload, sep=';', encoding='utf-8')
                if len(df_raw.columns) == 1:
                    fichier_upload.seek(0)
                    df_raw = pd.read_csv(fichier_upload, sep=',', encoding='utf-8')
            else:
                df_temp = pd.read_excel(fichier_upload, nrows=5)
                header_row = 0
                for i, row in df_temp.iterrows():
                    if 'N° Licence' in str(row.values) or 'Nom' in str(row.values) or "Catégorie d'âge" in str(row.values):
                        header_row = i + 1
                        break
                fichier_upload.seek(0)
                df_raw = pd.read_excel(fichier_upload, header=header_row)

            if "Catégorie d'âge" in df_raw.columns: df_raw = df_raw.rename(columns={"Catégorie d'âge": "Age"})
            if "Sigle du Club" in df_raw.columns: df_raw = df_raw.rename(columns={"Sigle du Club": "Club"})
            
            licence_col_found = None
            for col_name in df_raw.columns:
                col_str = str(col_name).strip().lower()
                if any(k in col_str for k in ["licence", "n° licence", "num_licence", "n°licence"]):
                    licence_col_found = col_name
                    break
            if licence_col_found:
                df_raw = df_raw.rename(columns={licence_col_found: "Licence"})
            elif "Licence" not in df_raw.columns:
                df_raw["Licence"] = ""

            comite_col_found = None
            for col_name in df_raw.columns:
                col_str = str(col_name).strip()
                col_lower = col_str.lower()
                if any(k in col_lower for k in ["comité", "comite", "ligue", "région", "region", "c.r."]):
                    comite_col_found = col_name
                    break
            if comite_col_found:
                df_raw = df_raw.rename(columns={comite_col_found: "Comité"})
            if "Comité" not in df_raw.columns: df_raw["Comité"] = "Comité Non Renseigné"

            style_col_found = None
            for col_name in df_raw.columns:
                col_str = str(col_name).strip().lower()
                if any(k in col_str for k in ["style", "discipline", "gréco", "greco", "gr/ll"]):
                    style_col_found = col_name
                    break
            if style_col_found:
                df_raw = df_raw.rename(columns={style_col_found: "Style"})
            if "Style" not in df_raw.columns:
                df_raw["Style"] = ""

            maitrise_col_found = None
            for col_name in df_raw.columns:
                col_str = str(col_name).strip().lower()
                if any(k in col_str for k in ["maîtrise", "maitrise", "niveau", "expéri", "experi", "niv"]):
                    maitrise_col_found = col_name
                    break
            if maitrise_col_found:
                df_raw = df_raw.rename(columns={maitrise_col_found: "Maîtrise"})
            if "Maîtrise" not in df_raw.columns:
                df_raw["Maîtrise"] = ""

            sexe_col_found = None
            for col_name in df_raw.columns:
                col_str = str(col_name).strip().lower()
                if any(k in col_str for k in ["sexe", "genre", "civilité", "civilite", "m/f"]):
                    sexe_col_found = col_name
                    break
            if sexe_col_found:
                df_raw = df_raw.rename(columns={sexe_col_found: "Sexe"})
            if "Sexe" not in df_raw.columns:
                df_raw["Sexe"] = ""

            if "Prénom" in df_raw.columns and "Nom" in df_raw.columns:
                df_raw["Nom"] = df_raw["Nom"].astype(str) + " " + df_raw["Prénom"].astype(str)

            def normaliser_age(val):
                if val is None or pd.isna(val):
                    return ''
                val_str = str(val).strip().upper().replace(' ', '').replace('-', '')
                if 'U7' in val_str or val_str == '7' or val_str.startswith('7'):
                    return 'U7'
                elif 'U9' in val_str or val_str == '9' or val_str.startswith('9'):
                    return 'U9'
                elif 'U11' in val_str or val_str == '11' or val_str.startswith('11'):
                    return 'U11'
                elif 'U13' in val_str or val_str == '13' or val_str.startswith('13'):
                    return 'U13'
                return str(val).strip().upper()

            df_inscr_total = df_raw.copy()
            if "Age" in df_inscr_total.columns:
                df_inscr_total['Age'] = df_inscr_total['Age'].apply(normaliser_age)
                df_inscr_total = df_inscr_total[df_inscr_total['Age'].isin(['U7', 'U9', 'U11', 'U13'])]
            
            total_inscrits_global = len(df_inscr_total)

            def attribuer_niveau(val):
                val_str = str(val).strip().lower()
                if not val_str or val_str in ['nan', 'none', '', '-']:
                    return 'Débutant'
                if any(k in val_str for k in ['confirm', 'expert', 'avanc', 'haut']) or val_str.startswith('c') or val_str in ['2', 'n2', 'niv 2', 'niveau 2']:
                    return 'Confirmé'
                elif any(k in val_str for k in ['début', 'debut', 'novice', 'initia']) or val_str.startswith('d') or val_str in ['1', 'n1', 'niv 1', 'niveau 1']:
                    return 'Débutant'
                else:
                    return 'Débutant'
            
            if poules_par_niveau:
                df_inscr_total['Niveau'] = df_inscr_total['Maîtrise'].apply(attribuer_niveau)
            else:
                df_inscr_total['Niveau'] = ''

            # Nettoyage et conversion du poids
            df_inscr_total['Poids_Clean'] = df_inscr_total['Poids'].astype(str).str.replace(',', '.')
            df_inscr_total['Poids_Num'] = pd.to_numeric(df_inscr_total['Poids_Clean'], errors='coerce').fillna(0)
            df_inscr_total['Poids'] = df_inscr_total['Poids_Num'].apply(formater_poids)

            # --- DÉTECTION ET VALIDATION DES LICENCES MANQUANTES ---
            erreurs_licence = []
            for _, row_test in df_inscr_total.iterrows():
                nom_raw = str(row_test.get('Nom', '')).strip()
                if not nom_raw or nom_raw.lower() in ['nan', 'nan nan', 'none', '']:
                    continue
                
                club_lutteur = str(row_test.get('Club', '')).strip()
                lic_val = str(row_test.get('Licence', '')).strip().lower()
                
                if lic_val in ['', 'nan', 'none', 'null', '-', '0', 'unspecified', 'inconnu'] or pd.isna(row_test.get('Licence')):
                    club_txt = f" ({club_lutteur})" if club_lutteur and club_lutteur.lower() not in ['nan', 'none', '-'] else ""
                    erreurs_licence.append(f"• **{nom_raw}**{club_txt} : Le numéro de licence n'est pas renseigné.")

            if erreurs_licence:
                st.error("❌ **Erreur d'importation dans le fichier de base (Licences manquantes) :**\n\n" + "\n".join(erreurs_licence))
                st.info("💡 *Remarque : Tous les lutteurs enregistrés dans le fichier de base doivent posséder un numéro de licence valide avant de générer la compétition.*")
                st.stop()

            # Fonctions utilitaires d'identification du Sexe et du Style
            def est_sexe_feminin(val):
                v = str(val).strip().upper()
                if not v or v in ['NAN', 'NONE', '']: return False
                return v in ['F', 'FEMME', 'FILLE', 'FEMININ', 'FÉMININ', 'FEM', 'WOMAN', 'GIRL'] or v.startswith('FÉM') or v.startswith('FEM')

            def est_sexe_masculin(val):
                v = str(val).strip().upper()
                if not v or v in ['NAN', 'NONE', '']: return False
                return v in ['M', 'H', 'HOMME', 'GARCON', 'GARÇON', 'MASCULIN', 'MASC', 'MAN', 'BOY'] or v.startswith('MASC') or v.startswith('HOM')

            def est_style_greco(val):
                s = str(val).strip().upper()
                if not s or s in ['NAN', 'NONE', '']: return False
                return any(k in s for k in ['LG', 'GR', 'GRECO', 'GRÉCO', 'ROMAIN']) or s == 'G'

            # --- VALIDATION : INCOMPATIBILITÉ SEXE FÉMININ & LUTTE GRÉCO-ROMAINE (LG/GR) ---
            erreurs_sexe_greco = []
            for _, row_test in df_inscr_total.iterrows():
                nom_lutteur = str(row_test.get('Nom', 'Lutteur Inconnu')).strip()
                if not nom_lutteur or nom_lutteur.lower() in ['nan', 'nan nan', 'none', '']:
                    continue
                
                sexe_val = row_test.get('Sexe', '')
                style_val = row_test.get('Style', '')
                
                if est_sexe_feminin(sexe_val) and est_style_greco(style_val):
                    club_lutteur = str(row_test.get('Club', '')).strip()
                    club_txt = f" ({club_lutteur})" if club_lutteur and club_lutteur.lower() not in ['nan', 'none', '-'] else ""
                    erreurs_sexe_greco.append(f"• **{nom_lutteur}**{club_txt} est de sexe féminin et ne peut pas être référencé en LG.")

            if erreurs_sexe_greco:
                st.error("❌ **Erreur d'importation dans le fichier :**\n\n" + "\n".join(erreurs_sexe_greco))
                st.info("💡 *Remarque : La lutte gréco-romaine (LG / GR) ne s'applique pas aux féminines. Une lutteuse de sexe féminin est automatiquement orientée en LF et ne peut pas être inscrite en LG.*")
                st.stop()

            # --- DÉTECTION ET VALIDATION DU STYLE "JEUNE" SANS SEXE IDENTIFIÉ ---
            erreurs_jeune = []
            for _, row_test in df_inscr_total.iterrows():
                val_style_raw = str(row_test.get('Style', '')).strip().lower()
                poids_val_raw = row_test.get('Poids_Num')
                try:
                    poids_val = float(poids_val_raw) if (poids_val_raw is not None and str(poids_val_raw).strip() != '') else 0.0
                except (ValueError, TypeError):
                    poids_val = 0.0
                nom_lutteur = row_test.get('Nom', 'Lutteur Inconnu')
                club_lutteur = row_test.get('Club', '')
                sexe_val = row_test.get('Sexe', '')
                
                # Si le style est "jeune" mais qu'aucun sexe (F ou M) n'est présent pour déduire LF ou LL
                if val_style_raw == 'jeune' and not (est_sexe_feminin(sexe_val) or est_sexe_masculin(sexe_val)):
                    if poids_val > 0:
                        club_txt = f" ({club_lutteur})" if club_lutteur and str(club_lutteur).lower() not in ['nan', 'none', '-'] else ""
                        erreurs_jeune.append(f"• **{nom_lutteur}**{club_txt} : Poids renseigné (**{poids_val} kg**) avec le style '**jeune**' sans colonne Sexe renseignée. Veuillez renseigner le sexe (M/F) ou le style (LL/LF/LG).")

            if erreurs_jeune:
                st.error("❌ **Erreur d'importation dans le fichier :**\n\n" + "\n".join(erreurs_jeune))
                st.info("💡 *Remarque : Le style 'jeune' nécessite d'indiquer le sexe dans la colonne Sexe/Genre (M ou F) pour orienter automatiquement le lutteur en LL ou LF, ou de corriger directement le style en LL, LF ou LG.*")
                st.stop()

            # RÈGLES DE NORMALISATION DES STYLES :
            # 1. Le sexe féminin est TOUJOURS identifié LF peu importe ce qu'il y a marqué dans la colonne style
            # 2. Le sexe masculin est orienté LL automatiquement peu importe ce qu'il y a marqué dans la colonne style, à part si marqué LG, GR ou gréco
            # 3. Si sexe non renseigné : repli sur la colonne style
            def normaliser_style(row_p):
                sexe_val = row_p.get('Sexe', '')
                style_val = row_p.get('Style', '')
                
                if est_sexe_feminin(sexe_val):
                    return 'LF'
                
                if est_sexe_masculin(sexe_val):
                    if est_style_greco(style_val):
                        return 'LG'
                    return 'LL'
                    
                if est_style_greco(style_val):
                    return 'LG'
                st_upper = str(style_val).strip().upper()
                if any(k in st_upper for k in ['LF', 'FEM', 'FÉM', 'FILLE']) or st_upper == 'F':
                    return 'LF'
                return 'LL'

            df_inscr_total['Style_Norm'] = df_inscr_total.apply(normaliser_style, axis=1)

            # Conserver tous les participants pesés (Poids_Num > 0)
            df_inscr = df_inscr_total[df_inscr_total['Poids_Num'] > 0].copy()

            total_participants_peses = len(df_inscr)
            total_non_peses = total_inscrits_global - total_participants_peses

            if df_inscr.empty:
                st.error("❌ Aucun lutteur U7, U9, U11 ou U13 avec un poids valide et un style de compétition n'a été trouvé dans le fichier.")
                st.stop()
            
            # Définition des groupes de styles selon le réglage de mixité (mixité interdite en U13 selon le règlement officiel FFLDA)
            def attribuer_style_groupe(row_p):
                style_norm = row_p['Style_Norm']
                age = str(row_p.get('Age', '')).strip().upper()
                
                # RÈGLEMENT FFLDA : En U13, il n'existe plus de catégorie mixte !
                # Filles et garçons sont obligatoirement séparés : LF (Féminine), LL (Libre) ou LG (Gréco)
                if age == 'U13':
                    if style_norm == 'LG':
                        return 'LG (Gréco)'
                    elif style_norm == 'LF':
                        return 'LF (Féminine)'
                    else:
                        return 'LL (Libre)'

                # Pour U7, U9, U11 : la mixité dépend du réglage mixte_active
                if mixte_active:
                    if style_norm == 'LG':
                        return 'LG (Gréco)'
                    else:
                        return 'Mixte (LL/LF)'
                else:
                    if style_norm == 'LG':
                        return 'LG (Gréco)'
                    elif style_norm == 'LF':
                        return 'LF (Féminine)'
                    else:
                        return 'LL (Libre)'

            df_inscr['Style_Groupe'] = df_inscr.apply(attribuer_style_groupe, axis=1)
            
            poules_u7, poules_u9, poules_u11, poules_u13 = [], [], [], []
            multiplicateur_poids = 1 + (tolerance_poids / 100.0)
            
            # --- GÉNÉRATION SPÉCIFIQUE U7 : 3 GROUPES HOMOGÈNES EN PLATEAUX D'ACTIVITÉ ---
            df_u7_total = df_inscr[df_inscr['Age'] == 'U7'].sort_values('Poids_Num')
            if not df_u7_total.empty:
                u7_all = df_u7_total.to_dict('records')
                nb_lutteurs_u7 = len(u7_all)
                # Exactement 3 groupes homogènes en effectif (ou moins si < 3 lutteurs au total)
                nb_groupes_u7 = 3 if nb_lutteurs_u7 >= 3 else max(1, nb_lutteurs_u7)
                
                base_len = nb_lutteurs_u7 // nb_groupes_u7
                reste = nb_lutteurs_u7 % nb_groupes_u7
                start_i = 0
                for g_i in range(nb_groupes_u7):
                    taille = base_len + (1 if g_i < reste else 0)
                    parts_g = u7_all[start_i : start_i + taille]
                    start_i += taille
                    
                    p_min = parts_g[0]['Poids_Num']
                    p_max = parts_g[-1]['Poids_Num']
                    p_min_str = formater_poids_court(p_min)
                    p_max_str = formater_poids_court(p_max)
                    poids_range_str = f" ({p_min_str} - {p_max_str})" if p_min_str and p_max_str else ""
                    
                    nom_g = f"U7 | Plateau - Groupe {g_i + 1} ({len(parts_g)} lutteurs){poids_range_str}"
                    poule_obj = {
                        'nom': nom_g,
                        'participants': parts_g,
                        'rondes': [],  # Animation sous forme de plateaux (pas de combats éliminatoires)
                        'type_formule': 'plateau_u7',
                        'num_groupe': g_i + 1
                    }
                    poules_u7.append(poule_obj)

            # --- GÉNÉRATION U9, U11 (Poules morphologiques à tolérance de poids) ---
            for age in ['U9', 'U11']:
                df_age = df_inscr[df_inscr['Age'] == age].sort_values('Poids_Num')
                max_size = 4 if age == 'U9' else 5
                counter_gr = 1  # Numérotation continue des groupes pour la catégorie d'âge (ex: Poule 1 à 14)
                
                for (style_grp, niveau), groupe in df_age.groupby(['Style_Groupe', 'Niveau']):
                    participants = groupe.to_dict('records')
                    poule_courante = []
                    suffixe_niveau = f" | {niveau}" if niveau != "" else ""
                    poules_groupe = []
                    
                    for p in participants:
                        if not poule_courante:
                            poule_courante.append(p)
                        else:
                            poids_min = float(poule_courante[0].get('Poids_Num') or 0.0)
                            p_poids = float(p.get('Poids_Num') or 0.0)
                            if p_poids <= (poids_min * multiplicateur_poids) and len(poule_courante) < max_size:
                                poule_courante.append(p)
                            else:
                                nom_groupe = f"{age} | {style_grp}{suffixe_niveau} ({formater_poids_court(poule_courante[0]['Poids_Num'])} - {formater_poids_court(poule_courante[-1]['Poids_Num'])})"
                                poule_obj = {'nom': nom_groupe, 'participants': list(poule_courante), 'rondes': generer_rondes_fflda(poule_courante), 'type_formule': 'poule'}
                                poules_groupe.append(poule_obj)
                                poule_courante = [p]
                    if poule_courante:
                        nom_groupe = f"{age} | {style_grp}{suffixe_niveau} ({formater_poids_court(poule_courante[0]['Poids_Num'])} - {formater_poids_court(poule_courante[-1]['Poids_Num'])})"
                        poule_obj = {'nom': nom_groupe, 'participants': list(poule_courante), 'rondes': generer_rondes_fflda(poule_courante), 'type_formule': 'poule'}
                        poules_groupe.append(poule_obj)
                    
                    if separer_clubs:
                        poules_groupe = optimiser_poules_clubs(poules_groupe, multiplicateur_poids)

                    # Eviter les poules de 1 en les fusionnant / rééquilibrant tout en respectant la tolérance de poids
                    poules_groupe = fusionner_poules_isolees(poules_groupe, multiplicateur_poids, max_size)
                    
                    if separer_clubs:
                        poules_groupe = optimiser_poules_clubs(poules_groupe, multiplicateur_poids)
                    
                    # Formater et numéroter en CONTINU par catégorie d'âge (ex: Poule 1 à Poule 14)
                    for p_obj in poules_groupe:
                        parts = p_obj['participants']
                        p_min = parts[0]['Poids_Num']
                        p_max = parts[-1]['Poids_Num']
                        p_obj['nom'] = f"{age} | {style_grp}{suffixe_niveau} | Poule {counter_gr} ({formater_poids_court(p_min)} - {formater_poids_court(p_max)})"
                        p_obj['rondes'] = generer_rondes_fflda(parts)
                        p_obj['type_formule'] = 'poule'
                        counter_gr += 1
                    
                    if age == 'U9': poules_u9.extend(poules_groupe)
                    else: poules_u11.extend(poules_groupe)

            # --- GÉNÉRATION U13 (Catégories de poids fixes : 30, 33, 36, 39, 42, 46, 50, 55, 60, +60 kg) ---
            df_u13_total = df_inscr[df_inscr['Age'] == 'U13'].copy()
            if not df_u13_total.empty:
                df_u13_total['Cat_Poids'] = df_u13_total['Poids_Num'].apply(attribuer_categorie_poids_u13)
                for (style_grp, niveau), groupe_style in df_u13_total.groupby(['Style_Groupe', 'Niveau']):
                    suffixe_niveau = f" | {niveau}" if niveau != "" else ""
                    for cat_poids in ORDRE_POIDS_U13:
                        groupe_cat = groupe_style[groupe_style['Cat_Poids'] == cat_poids]
                        if groupe_cat.empty:
                            continue
                        participants_cat = groupe_cat.to_dict('records')
                        comp_obj = generer_competition_u13('U13', style_grp, suffixe_niveau, cat_poids, participants_cat, separer_clubs)
                        if comp_obj:
                            poules_u13.append(comp_obj)

            tous_les_groupes = poules_u7 + poules_u9 + poules_u11 + poules_u13
            participants_par_poule = {p['nom']: p['participants'] for p in tous_les_groupes}
            rondes_par_categorie = {p['nom']: p['rondes'] for p in tous_les_groupes}
            poule_obj_map = {p['nom']: p for p in tous_les_groupes}

            tapis_poules_u7 = {i: [] for i in range(nb_tapis)}
            tapis_poules_u9 = {i: [] for i in range(nb_tapis)}
            tapis_poules_u11 = {i: [] for i in range(nb_tapis)}
            tapis_poules_u13 = {i: [] for i in range(nb_tapis)}
            
            for i, p in enumerate(poules_u7): tapis_poules_u7[i % nb_tapis].append(p)
            for i, p in enumerate(poules_u9): tapis_poules_u9[i % nb_tapis].append(p)
            for i, p in enumerate(poules_u11): tapis_poules_u11[i % nb_tapis].append(p)
            for i, p in enumerate(poules_u13): tapis_poules_u13[i % nb_tapis].append(p)

            dt_pesee_u9 = datetime.combine(datetime.today(), heure_pesee_u9)
            dt_debut_u7 = dt_pesee_u9 + timedelta(minutes=duree_pesee)
            total_matchs_calcules = 0

            ref_usage_count = {}

            def choisir_arbitre_match(c1_club, c2_club, t_idx):
                cand_list = tapis_arbitres.get(t_idx, [])
                if not cand_list and liste_arbitres:
                    cand_list = liste_arbitres
                
                if not cand_list:
                    return "Non attribué"
                
                c1_clean = str(c1_club).strip().lower()
                c2_clean = str(c2_club).strip().lower()
                ignored_clubs = ['', '-', 'indépendant', 'independant', 'none', 'nan']
                
                if eviter_arbitre_meme_club:
                    sans_conflit = []
                    for arb in cand_list:
                        arb_club_clean = str(arb.get('Club', '')).strip().lower()
                        if arb_club_clean in ignored_clubs:
                            sans_conflit.append(arb)
                        elif arb_club_clean != c1_clean and arb_club_clean != c2_clean:
                            sans_conflit.append(arb)
                    pool_choix = sans_conflit if sans_conflit else cand_list
                else:
                    pool_choix = cand_list
                        
                arb_choisi = min(pool_choix, key=lambda a: ref_usage_count.get(a['Nom_Complet'], 0))
                
                ref_usage_count[arb_choisi['Nom_Complet']] = ref_usage_count.get(arb_choisi['Nom_Complet'], 0) + 1
                return arb_choisi['Nom_Complet']

            def nommer_tour_match(poule_obj, r_idx, m_idx, m=None):
                if not poule_obj:
                    return f"Tour {r_idx + 1}"
                type_f = poule_obj.get('type_formule', 'poule')
                rondes = poule_obj.get('rondes', [])
                total_r = len(rondes)
                
                if type_f == 'tableau':
                    if total_r >= 6:
                        if r_idx == 0:
                            return f"1/32 de Finale ({m_idx + 1})"
                        elif r_idx == 1:
                            return f"1/16 de Finale ({m_idx + 1})"
                        elif r_idx == 2:
                            return f"1/8 de Finale ({m_idx + 1})"
                        elif r_idx == 3:
                            return f"1/4 de Finale ({m_idx + 1})"
                        elif r_idx == total_r - 2:
                            if m_idx in (0, 1):
                                return f"Demi-Finale {m_idx + 1}"
                            else:
                                return f"Repêchage 1/4 ({m_idx - 1})"
                        elif r_idx == total_r - 1:
                            if m_idx == 0:
                                return "Grande Finale (Or / Argent)"
                            elif m_idx == 1:
                                return "Finale Bronze 1"
                            elif m_idx == 2:
                                return "Finale Bronze 2"
                            else:
                                return f"Finale {m_idx + 1}"
                        else:
                            return f"Tour {r_idx + 1}"
                    elif total_r == 5:
                        if r_idx == 0:
                            return f"1/16 de Finale ({m_idx + 1})"
                        elif r_idx == 1:
                            return f"1/8 de Finale ({m_idx + 1})"
                        elif r_idx == 2:
                            return f"1/4 de Finale ({m_idx + 1})"
                        elif r_idx == total_r - 2:
                            if m_idx in (0, 1):
                                return f"Demi-Finale {m_idx + 1}"
                            else:
                                return f"Repêchage 1/4 ({m_idx - 1})"
                        elif r_idx == total_r - 1:
                            if m_idx == 0:
                                return "Grande Finale (Or / Argent)"
                            elif m_idx == 1:
                                return "Finale Bronze 1"
                            elif m_idx == 2:
                                return "Finale Bronze 2"
                            else:
                                return f"Finale {m_idx + 1}"
                        else:
                            return f"Tour {r_idx + 1}"
                    elif total_r == 4:
                        if r_idx == 0:
                            return f"Tour Préliminaire 1/8 ({m_idx + 1})"
                        elif r_idx == 1:
                            return f"1/4 de Finale ({m_idx + 1})"
                        elif r_idx == 2:
                            if m_idx in (0, 1):
                                return f"Demi-Finale {m_idx + 1}"
                            else:
                                return f"Repêchage 1/4 ({m_idx - 1})"
                        elif r_idx == 3:
                            if m_idx == 0:
                                return "Grande Finale (Or / Argent)"
                            elif m_idx == 1:
                                return "Finale Bronze 1"
                            elif m_idx == 2:
                                return "Finale Bronze 2"
                            else:
                                return f"Finale {m_idx + 1}"
                        else:
                            return f"Tour {r_idx + 1}"
                    elif total_r == 3:
                        if r_idx == 0:
                            return f"1/4 de Finale ({m_idx + 1})"
                        elif r_idx == 1:
                            if m_idx in (0, 1):
                                return f"Demi-Finale {m_idx + 1}"
                            else:
                                return f"Repêchage 1/4 ({m_idx - 1})"
                        elif r_idx == 2:
                            if m_idx == 0:
                                return "Grande Finale (Or / Argent)"
                            elif m_idx == 1:
                                return "Finale Bronze 1"
                            elif m_idx == 2:
                                return "Finale Bronze 2"
                            else:
                                return f"Finale {m_idx + 1}"
                        else:
                            return f"Tour {r_idx + 1}"
                    else:
                        return f"Tour {r_idx + 1}"
                        
                elif type_f == 'poules_croisees':
                    if r_idx == 4:
                        if m_idx == 0:
                            return "Grande Finale (Or / Argent)"
                        else:
                            return "Finale 3-4 (Bronze unique)"
                    elif r_idx == 3:
                        return f"Demi-Finale Croisée {m_idx + 1}"
                    else:
                        return f"Poule - Tour {r_idx + 1}"
                        
                else:
                    return f"Tour {r_idx + 1}"

            def ordonnancer_phase(poules_phase, heure_debut_phase, duree_combat, meme_tapis=True):
                global total_matchs_calcules
                if not poules_phase:
                    return [heure_debut_phase for _ in range(nb_tapis)], {t: [] for t in range(nb_tapis)}
                
                planning = {t: [] for t in range(nb_tapis)}
                tapis_heure = [heure_debut_phase for _ in range(nb_tapis)]
                last_match_time = {}

                if meme_tapis:
                    # RÈGLE OFFICIELLE : Chaque poule de même style et même catégorie de poids (ex: U13 Gréco 30 kg) est affectée au même tapis fixe.
                    # Une même catégorie de poids d'un autre style (ex: U13 Libre 30 kg ou U13 Féminine 30 kg) peut aller sur un tapis différent,
                    # ce qui permet un équilibrage encore plus optimal du nombre de matchs par tapis (Algorithme LPT).
                    def obtenir_cle_poids_lot(p_obj):
                        nom_p = p_obj.get('nom', '')
                        parts_nom = [x.strip() for x in nom_p.split('|')]
                        age_p = parts_nom[0] if parts_nom else 'U13'
                        style_p = p_obj.get('style_grp') or (parts_nom[1] if len(parts_nom) >= 2 else '')
                        cat_p = p_obj.get('cat_poids')
                        
                        if cat_p:
                            return f"{age_p} | {style_p} | {str(cat_p).strip()}"
                        
                        m_poids = re.search(r'\b(\+?\d+\s*kg)\b', nom_p, re.IGNORECASE)
                        if m_poids:
                            return f"{age_p} | {style_p} | {m_poids.group(1).strip()}"
                        
                        # Pour U9 / U11 : chaque poule morphologique constitue son propre groupe
                        return nom_p

                    # 1. Regroupement par catégorie de poids (tous les combats de la même catégorie de poids restent ensemble)
                    lots_poids = {}
                    for p in poules_phase:
                        cle = obtenir_cle_poids_lot(p)
                        if cle not in lots_poids:
                            lots_poids[cle] = []
                        lots_poids[cle].append(p)

                    # 2. Calcul du nombre de matchs par lot de catégorie de poids
                    lots_avec_poids = []
                    for cle, liste_p in lots_poids.items():
                        nb_m = sum(sum(len(r) for r in p.get('rondes', [])) for p in liste_p)
                        lots_avec_poids.append((cle, liste_p, nb_m))

                    # 3. Tri décroissant par nombre de matchs (LPT - Longest Processing Time First)
                    lots_avec_poids.sort(key=lambda x: x[2], reverse=True)

                    # 4. Affectation équilibrée sur les tapis : chaque lot de poids va sur le tapis actuellement le moins chargé en matchs
                    tapis_poules = {t: [] for t in range(nb_tapis)}
                    charge_matchs_tapis = [0] * nb_tapis

                    for cle, liste_p, nb_m in lots_avec_poids:
                        t_min = min(range(nb_tapis), key=lambda t: (charge_matchs_tapis[t], t))
                        tapis_poules[t_min].extend(liste_p)
                        charge_matchs_tapis[t_min] += nb_m

                    for t in range(nb_tapis):
                        poules_t = tapis_poules[t]
                        if not poules_t:
                            continue
                        
                        matchs_tapis = []
                        max_rondes = max((len(p['rondes']) for p in poules_t), default=0)
                        for r in range(max_rondes):
                            for p in poules_t:
                                if r < len(p['rondes']):
                                    for m_idx, m in enumerate(p['rondes'][r]):
                                        matchs_tapis.append({
                                            'poule': p['nom'],
                                            'poule_obj': p,
                                            'p1': m[0]['Nom'],
                                            'p1_club': m[0].get('Club', ''),
                                            'p1_comite': m[0].get('Comité', m[0].get('Comite', 'Comité Non Renseigné')),
                                            'p2': m[1]['Nom'],
                                            'p2_club': m[1].get('Club', ''),
                                            'p2_comite': m[1].get('Comité', m[1].get('Comite', 'Comité Non Renseigné')),
                                            'tour': r + 1,
                                            'nom_tour': nommer_tour_match(p, r, m_idx, m)
                                        })

                        while matchs_tapis:
                            t_curr_time = tapis_heure[t]
                            
                            match_choisi_idx = None
                            for idx, m in enumerate(matchs_tapis):
                                p1, p2 = m['p1'], m['p2']
                                t_pret_p1 = last_match_time.get(p1, heure_debut_phase)
                                t_pret_p2 = last_match_time.get(p2, heure_debut_phase)
                                if t_pret_p1 <= t_curr_time and t_pret_p2 <= t_curr_time:
                                    match_choisi_idx = idx
                                    break
                            
                            if match_choisi_idx is not None:
                                m = matchs_tapis.pop(match_choisi_idx)
                                heure_debut_match = t_curr_time
                            else:
                                meilleur_idx = 0
                                meilleur_temps = datetime.max
                                for idx, m in enumerate(matchs_tapis):
                                    p1, p2 = m['p1'], m['p2']
                                    t_pret = max(last_match_time.get(p1, heure_debut_phase), last_match_time.get(p2, heure_debut_phase))
                                    if t_pret < meilleur_temps:
                                        meilleur_temps = t_pret
                                        meilleur_idx = idx
                                
                                m = matchs_tapis.pop(meilleur_idx)
                                heure_debut_match = max(t_curr_time, meilleur_temps)
                                
                                if heure_debut_match > t_curr_time:
                                    attente_min = int((heure_debut_match - t_curr_time).total_seconds() // 60)
                                    if attente_min > 0:
                                        planning[t].append({
                                            "Type": "ATTENTE", 
                                            "Heure": t_curr_time.strftime("%H:%M"), 
                                            "Texte": f"⏳ Repos ({attente_min} min)"
                                        })

                            arb_nom = choisir_arbitre_match(m['p1_club'], m['p2_club'], t)

                            planning[t].append({
                                "Type": "MATCH",
                                "Heure": heure_debut_match.strftime("%H:%M"),
                                "Duree": duree_combat,
                                "Cat": m['poule'],
                                "Combattant 1": m['p1'],
                                "Club 1": m['p1_club'],
                                "Comité 1": m['p1_comite'],
                                "Combattant 2": m['p2'],
                                "Club 2": m['p2_club'],
                                "Comité 2": m['p2_comite'],
                                "Tour": m['tour'],
                                "Nom_Tour": m.get('nom_tour', f"Tour {m['tour']}"),
                                "Arbitre": arb_nom
                            })
                            
                            total_matchs_calcules += 1
                            fin_match = heure_debut_match + timedelta(minutes=duree_combat)
                            delai_repos = timedelta(minutes=(repos_matchs * duree_combat))
                            
                            last_match_time[m['p1']] = fin_match + delai_repos
                            last_match_time[m['p2']] = fin_match + delai_repos
                            
                            tapis_heure[t] = fin_match
                else:
                    # Répartition dynamique sur le premier tapis disponible
                    matchs_a_jouer = []
                    max_rondes = max((len(p['rondes']) for p in poules_phase), default=0)
                    for r in range(max_rondes):
                        for p in poules_phase:
                            if r < len(p['rondes']):
                                for m_idx, m in enumerate(p['rondes'][r]):
                                    matchs_a_jouer.append({
                                        'poule': p['nom'],
                                        'poule_obj': p,
                                        'p1': m[0]['Nom'],
                                        'p1_club': m[0].get('Club', ''),
                                        'p1_comite': m[0].get('Comité', m[0].get('Comite', 'Comité Non Renseigné')),
                                        'p2': m[1]['Nom'],
                                        'p2_club': m[1].get('Club', ''),
                                        'p2_comite': m[1].get('Comité', m[1].get('Comite', 'Comité Non Renseigné')),
                                        'tour': r + 1,
                                        'nom_tour': nommer_tour_match(p, r, m_idx, m)
                                    })
                    
                    while matchs_a_jouer:
                        t_min_idx = min(range(nb_tapis), key=lambda t: tapis_heure[t])
                        t_min_time = tapis_heure[t_min_idx]
                        
                        match_choisi_idx = None
                        for idx, m in enumerate(matchs_a_jouer):
                            p1, p2 = m['p1'], m['p2']
                            t_pret_p1 = last_match_time.get(p1, heure_debut_phase)
                            t_pret_p2 = last_match_time.get(p2, heure_debut_phase)
                            if t_pret_p1 <= t_min_time and t_pret_p2 <= t_min_time:
                                match_choisi_idx = idx
                                break
                        
                        if match_choisi_idx is not None:
                            m = matchs_a_jouer.pop(match_choisi_idx)
                            heure_debut_match = t_min_time
                        else:
                            meilleur_idx = 0
                            meilleur_temps = datetime.max
                            for idx, m in enumerate(matchs_a_jouer):
                                p1, p2 = m['p1'], m['p2']
                                t_pret = max(last_match_time.get(p1, heure_debut_phase), last_match_time.get(p2, heure_debut_phase))
                                if t_pret < meilleur_temps:
                                    meilleur_temps = t_pret
                                    meilleur_idx = idx
                            
                            m = matchs_a_jouer.pop(meilleur_idx)
                            heure_debut_match = max(t_min_time, meilleur_temps)
                            
                            if heure_debut_match > t_min_time:
                                attente_min = int((heure_debut_match - t_min_time).total_seconds() // 60)
                                if attente_min > 0:
                                    planning[t_min_idx].append({
                                        "Type": "ATTENTE", 
                                        "Heure": t_min_time.strftime("%H:%M"), 
                                        "Texte": f"⏳ Repos ({attente_min} min)"
                                    })

                        arb_nom = choisir_arbitre_match(m['p1_club'], m['p2_club'], t_min_idx)

                        planning[t_min_idx].append({
                            "Type": "MATCH",
                            "Heure": heure_debut_match.strftime("%H:%M"),
                            "Duree": duree_combat,
                            "Cat": m['poule'],
                            "Combattant 1": m['p1'],
                            "Club 1": m['p1_club'],
                            "Comité 1": m['p1_comite'],
                            "Combattant 2": m['p2'],
                            "Club 2": m['p2_club'],
                            "Comité 2": m['p2_comite'],
                            "Tour": m['tour'],
                            "Nom_Tour": m.get('nom_tour', f"Tour {m['tour']}"),
                            "Arbitre": arb_nom
                        })
                        
                        total_matchs_calcules += 1
                        fin_match = heure_debut_match + timedelta(minutes=duree_combat)
                        delai_repos = timedelta(minutes=(repos_matchs * duree_combat))
                        
                        last_match_time[m['p1']] = fin_match + delai_repos
                        last_match_time[m['p2']] = fin_match + delai_repos
                        
                        tapis_heure[t_min_idx] = fin_match

                return tapis_heure, planning

            # PHASE 1a : Tous les U7 (Format Plateaux d'Activité : 3 rotations de duree_plateau_u7 min)
            planning_u7 = {t: [] for t in range(nb_tapis)}
            tapis_heure_u7 = [dt_debut_u7 for _ in range(nb_tapis)]
            
            if poules_u7:
                nb_rotations = 3
                noms_plateaux = [
                    "Plateau 1 : Motricité & Agilité",
                    "Plateau 2 : Ateliers techniques",
                    "Plateau 3 : Oppositions"
                ]
                
                # Affectation tournante des 3 groupes sur les 3 plateaux d'activité
                rot_affectations = [
                    {0: 0, 1: 1, 2: 2},  # Rot 1: P1->G1, P2->G2, P3->G3
                    {0: 2, 1: 0, 2: 1},  # Rot 2: P1->G3, P2->G1, P3->G2
                    {0: 1, 1: 2, 2: 0}   # Rot 3: P1->G2, P2->G3, P3->G1
                ]
                
                for r in range(nb_rotations):
                    h_rot = dt_debut_u7 + timedelta(minutes=r * duree_plateau_u7)
                    
                    for t in range(nb_tapis):
                        plat_idx = t % 3
                        grp_idx = rot_affectations[r][plat_idx]
                        
                        grp_poule = poules_u7[grp_idx] if grp_idx < len(poules_u7) else poules_u7[0]
                        grp_parts = grp_poule.get('participants', [])
                        
                        arb_nom = choisir_arbitre_match("", "", t)
                        arb_disp = arb_nom if (arb_nom and arb_nom != "Non attribué") else "Animateur FFLDA"
                        
                        item_plateau = {
                            "Type": "MATCH",
                            "Heure": h_rot.strftime("%H:%M"),
                            "Duree": duree_plateau_u7,
                            "Cat": f"U7 | Plateau {plat_idx + 1}",
                            "Combattant 1": f"Groupe {grp_idx + 1} ({len(grp_parts)} lutteurs)",
                            "Club 1": "Tous clubs",
                            "Comité 1": "",
                            "Combattant 2": noms_plateaux[plat_idx],
                            "Club 2": "",
                            "Comité 2": "",
                            "Tour": f"Rotation {r + 1}/3",
                            "Nom_Tour": f"Rotation {r + 1} / 3",
                            "Arbitre": arb_disp
                        }
                        planning_u7[t].append(item_plateau)
                
                fin_u7_globale = dt_debut_u7 + timedelta(minutes=3 * duree_plateau_u7)
                for t in range(nb_tapis):
                    tapis_heure_u7[t] = fin_u7_globale
            else:
                fin_u7_globale = dt_debut_u7

            # PHASE 1b : Tous les U9
            tapis_heure_u9, planning_u9 = ordonnancer_phase(poules_u9, fin_u7_globale, duree_u9, meme_tapis_poule)
            fin_u9_globale = max(tapis_heure_u9) if poules_u9 else fin_u7_globale

            planning_phase1 = {t: planning_u7[t] + planning_u9[t] for t in range(nb_tapis)}

            # PAUSE / PESÉE 2 (U11 / U13)
            dt_pesee_2 = None
            if "2" in type_pesee:
                dt_pesee_2 = fin_u9_globale + timedelta(minutes=duree_pause)
                dt_debut_p2_theorique = dt_pesee_2 + timedelta(minutes=duree_pesee)
            else:
                dt_debut_p2_theorique = fin_u9_globale

            # Égalisation du nombre de lignes en Phase 1 (U7 + U9) pour aligner horizontalement la PAUSE
            max_phase1_lignes = max(len(planning_phase1[t]) for t in range(nb_tapis)) if planning_phase1 else 0
            for t in range(nb_tapis):
                while len(planning_phase1[t]) < max_phase1_lignes:
                    planning_phase1[t].append({"Type": "VIDE"})

            tapis_heure_p1 = [fin_u9_globale for _ in range(nb_tapis)]
            if activer_pause and duree_pause > 0:
                for t in range(nb_tapis):
                    planning_phase1[t].append({"Type": "PAUSE", "Heure": fin_u9_globale.strftime("%H:%M")})
                    tapis_heure_p1[t] = fin_u9_globale + timedelta(minutes=duree_pause)

            debut_p2_reel = max(tapis_heure_p1)
            if dt_debut_p2_theorique and debut_p2_reel < dt_debut_p2_theorique:
                debut_p2_reel = dt_debut_p2_theorique

            for t in range(nb_tapis):
                if (poules_u11 or poules_u13) and tapis_heure_p1[t] < debut_p2_reel:
                    attente = int((debut_p2_reel - tapis_heure_p1[t]).total_seconds() // 60)
                    if attente > 0:
                        planning_phase1[t].append({
                            "Type": "ATTENTE", 
                            "Heure": tapis_heure_p1[t].strftime("%H:%M"), 
                            "Texte": f"Pesée + échauffement U11/U13"
                        })

            # PHASE 2a : Tous les U11
            if poules_u11:
                tapis_heure_u11, planning_u11 = ordonnancer_phase(poules_u11, debut_p2_reel, duree_u11, meme_tapis_poule)
                fin_u11_globale = max(tapis_heure_u11)
            else:
                tapis_heure_u11 = [debut_p2_reel for _ in range(nb_tapis)]
                planning_u11 = {t: [] for t in range(nb_tapis)}
                fin_u11_globale = debut_p2_reel

            # PHASE 2b : Tous les U13
            if poules_u13:
                tapis_heure_u13, planning_u13 = ordonnancer_phase(poules_u13, fin_u11_globale, duree_u13, meme_tapis_poule)
                fin_estimee = max(tapis_heure_u13)
            else:
                planning_u13 = {t: [] for t in range(nb_tapis)}
                fin_estimee = fin_u11_globale

            planning_tapis = {t: planning_phase1[t] + planning_u11[t] + planning_u13[t] for t in range(nb_tapis)}

            texte_pesee_1 = "1ère pesée" if "1" in type_pesee else "Pesée U7/U9"
            valeur_pause = f"{duree_pause} min" if (activer_pause and duree_pause > 0) else "0 min"

            str_comp_u7 = f"{dt_debut_u7.strftime('%H:%M')} - {fin_u7_globale.strftime('%H:%M')}"
            str_comp_u9 = f"{fin_u7_globale.strftime('%H:%M')} - {fin_u9_globale.strftime('%H:%M')}"
            str_comp_u11 = f"{debut_p2_reel.strftime('%H:%M')} - {fin_u11_globale.strftime('%H:%M')}"
            str_comp_u13 = f"{fin_u11_globale.strftime('%H:%M')} - {fin_estimee.strftime('%H:%M')}"

            st.success("✨ Fichier analysé avec succès ! Tournoi généré.")
            
            # --- PRÉPARATION DES DONNÉES DU RÉSUMÉ ---
            lignes_accueil = [
                {"Étape de la journée": texte_pesee_1, "Horaire / Valeur": dt_pesee_u9.strftime('%H:%M')},
            ]
            if poules_u7:
                lignes_accueil.append({"Étape de la journée": "Animation U7 (3 Plateaux)", "Horaire / Valeur": str_comp_u7})
                lignes_accueil.append({"Étape de la journée": "Groupes U7 (Plateaux)", "Horaire / Valeur": f"{len(poules_u7)} groupes"})
            lignes_accueil.extend([
                {"Étape de la journée": "Compétition U9", "Horaire / Valeur": str_comp_u9},
                {"Étape de la journée": "Poules U9 générées", "Horaire / Valeur": f"{len(poules_u9)} poules"},
                {"Étape de la journée": "Pause de la compétition", "Horaire / Valeur": valeur_pause}
            ])
            if "2" in type_pesee and dt_pesee_2:
                lignes_accueil.append({"Étape de la journée": "2ème pesée (U11/U13)", "Horaire / Valeur": dt_pesee_2.strftime('%H:%M')})

            lignes_accueil.extend([
                {"Étape de la journée": "Compétition U11", "Horaire / Valeur": str_comp_u11},
                {"Étape de la journée": "Poules U11 générées", "Horaire / Valeur": f"{len(poules_u11)} poules"},
            ])
            if poules_u13:
                lignes_accueil.extend([
                    {"Étape de la journée": "Compétition U13", "Horaire / Valeur": str_comp_u13},
                    {"Étape de la journée": "Poules & Tableaux U13", "Horaire / Valeur": f"{len(poules_u13)} catégories"},
                ])
            total_poules_calc = len(poules_u7) + len(poules_u9) + len(poules_u11) + len(poules_u13)
            lignes_accueil.extend([
                {"Étape de la journée": "Total Groupes (U7 + U9 + U11 + U13)", "Horaire / Valeur": f"{total_poules_calc} groupes"},
                {"Étape de la journée": "Fin de la compétition estimée", "Horaire / Valeur": fin_estimee.strftime('%H:%M')}
            ])

            # --- GÉNÉRATION DES DOCUMENTS HTML & PDF PAYSAGE DU TOURNOI ---
            sections_tournoi_complet = [
                ("📊 Résumé Prévisionnel de la Journée", pd.DataFrame(lignes_accueil))
            ]
            
            if liste_arbitres:
                lignes_arb_print = []
                for t in range(nb_tapis):
                    noms_arb = ", ".join([a['Nom_Complet'] for a in tapis_arbitres[t]]) if tapis_arbitres[t] else "Aucun arbitre affecté"
                    lignes_arb_print.append({"Tapis": f"Tapis {t + 1}", "Effectif": f"{len(tapis_arbitres[t])} arbitres", "Équipe d'Arbitrage Désignée": noms_arb})
                sections_tournoi_complet.append(("🛡️ Désignation des Équipes d'Arbitrage par Tapis", pd.DataFrame(lignes_arb_print)))

            max_lignes = max(len(liste) for liste in planning_tapis.values()) if planning_tapis else 0
            grille_ui = []
            for row_idx in range(max_lignes):
                ligne = {}
                for t in range(nb_tapis):
                    col = f"Tapis {t + 1}"
                    if row_idx < len(planning_tapis[t]):
                        m = planning_tapis[t][row_idx]
                        if m["Type"] == "PAUSE": ligne[col] = f"[{m['Heure']}] ⏸️ PAUSE"
                        elif m["Type"] == "ATTENTE": ligne[col] = f"[{m['Heure']}] {m['Texte']}"
                        elif m["Type"] == "VIDE": ligne[col] = ""
                        else:
                            arb_str = f" (🛡️ {m['Arbitre']})" if m.get('Arbitre') and m['Arbitre'] != "Non attribué" else ""
                            t_nom = m.get('Nom_Tour') or (f"Tour {m['Tour']}" if m.get('Tour') else "")
                            tour_str = f" [🎯 {t_nom}]" if t_nom else ""
                            ligne[col] = f"[{m['Heure']}] ({m['Duree']}m) [{m['Cat']}]{tour_str} - {m['Combattant 1']} vs {m['Combattant 2']}{arb_str}"
                    else: ligne[col] = ""
                grille_ui.append(ligne)
            
            sections_tournoi_complet.append(("📅 Grille Globale de Passage", pd.DataFrame(grille_ui)))

            # Ajout des sections de Grille par Tapis pour impression/PDF (1 page par onglet)
            for t in range(nb_tapis):
                lignes_tapis_doc = []
                m_count_doc = 0
                for m in planning_tapis[t]:
                    if m["Type"] == "PAUSE":
                        lignes_tapis_doc.append({"N°": "-", "Heure": m["Heure"], "Catégorie": "⏸️ PAUSE", "Tour": "-", "Lutteur Rouge": "-", "Pt Clt (R)": "", "Lutteur Bleu": "-", "Pt Clt (B)": "", "Arbitre": ""})
                    elif m["Type"] == "ATTENTE":
                        lignes_tapis_doc.append({"N°": "-", "Heure": m["Heure"], "Catégorie": f"⏳ {m['Texte']}", "Tour": "-", "Lutteur Rouge": "-", "Pt Clt (R)": "", "Lutteur Bleu": "-", "Pt Clt (B)": "", "Arbitre": ""})
                    elif m["Type"] == "MATCH":
                        m_count_doc += 1
                        c1_t = f"{m['Combattant 1']}"
                        if m.get('Club 1'): c1_t += f" ({m['Club 1']})"
                        c2_t = f"{m['Combattant 2']}"
                        if m.get('Club 2'): c2_t += f" ({m['Club 2']})"
                        lignes_tapis_doc.append({
                            "N°": f"M{m_count_doc}",
                            "Heure": f"{m['Heure']}",
                            "Catégorie": m['Cat'],
                            "Tour": m.get('Nom_Tour') or (f"Tour {m.get('Tour')}" if m.get('Tour') else "-"),
                            "Lutteur Rouge": c1_t,
                            "Pt Clt (R)": "[   ]",
                            "Lutteur Bleu": c2_t,
                            "Pt Clt (B)": "[   ]",
                            "Arbitre": m.get("Arbitre", "")
                        })
                sections_tournoi_complet.append((f"🥋 Grille de Passage & Scores - Tapis {t + 1}", pd.DataFrame(lignes_tapis_doc)))

            for nom_poule, liste_p in participants_par_poule.items():
                p_obj = poule_obj_map.get(nom_poule)
                if p_obj and p_obj.get('type_formule') == 'plateau_u7':
                    df_poule_vue = pd.DataFrame([
                        {
                            "N°": idx,
                            "Nom": p.get('Nom', ''),
                            "Club": p.get('Club', ''),
                            "Poids": formater_poids(p.get('Poids', '')),
                            "Plateau 1 (Motricité & Agilité)": "[ ✓ ]",
                            "Plateau 2 (Ateliers techniques)": "[ ✓ ]",
                            "Plateau 3 (Oppositions)": "[ ✓ ]",
                            "Validation": "🥇 Médaille d'Or"
                        }
                        for idx, p in enumerate(liste_p, 1)
                    ])
                    sections_tournoi_complet.append((f"🏅 Animation U7 (3 Plateaux) : {nom_poule}", df_poule_vue))
                else:
                    df_poule_vue = pd.DataFrame(liste_p)[['Nom', 'Club', 'Poids']]
                    sections_tournoi_complet.append((f"🤼 Feuille de Poule : {nom_poule}", df_poule_vue))
                
            # --- GÉNÉRATION DU CLASSEUR EXCEL OFFICIEL FFLDA ---
            output_excel = io.BytesIO()
            coords_matchs_tapis = {}
            tapis_slots_map = {}
            poule_sheet_names = {}

            with pd.ExcelWriter(output_excel, engine='openpyxl') as writer:
                writer.book.calculation.fullCalcOnLoad = True
                
                resume_data = [
                    {"Étape de la journée": texte_pesee_1, "Horaire / Valeur": dt_pesee_u9.strftime('%H:%M')},
                ]
                if poules_u7:
                    resume_data.append({"Étape de la journée": "Animation U7 (3 Plateaux)", "Horaire / Valeur": str_comp_u7})
                    resume_data.append({"Étape de la journée": "Groupes U7 (Plateaux)", "Horaire / Valeur": f"{len(poules_u7)} groupes"})
                resume_data.extend([
                    {"Étape de la journée": "Compétition U9", "Horaire / Valeur": str_comp_u9},
                    {"Étape de la journée": "Pause de la compétition", "Horaire / Valeur": valeur_pause}
                ])
                if "2" in type_pesee and dt_pesee_2:
                    resume_data.append({"Étape de la journée": "2ème pesée (U11/U13)", "Horaire / Valeur": dt_pesee_2.strftime('%H:%M')})
                
                resume_data.extend([
                    {"Étape de la journée": "Compétition U11", "Horaire / Valeur": str_comp_u11},
                ])
                if poules_u13:
                    resume_data.append({"Étape de la journée": "Compétition U13", "Horaire / Valeur": str_comp_u13})
                resume_data.append({"Étape de la journée": "Fin de la compétition estimée", "Horaire / Valeur": fin_estimee.strftime('%H:%M')})
                pd.DataFrame(resume_data).to_excel(writer, sheet_name="Résumé", index=False)
                
                max_lignes = max(len(liste) for liste in planning_tapis.values()) if planning_tapis else 0
                grille = []
                for row_idx in range(max_lignes):
                    ligne = {}
                    for t in range(nb_tapis):
                        col = f"Tapis {t + 1}"
                        if row_idx < len(planning_tapis[t]):
                            m = planning_tapis[t][row_idx]
                            if m["Type"] == "PAUSE": ligne[col] = f"[{m['Heure']}]\n⏸️ PAUSE DE LA COMPÉTITION"
                            elif m["Type"] == "ATTENTE": ligne[col] = f"[{m['Heure']}]\n{m['Texte']}"
                            elif m["Type"] == "VIDE": ligne[col] = ""
                            else:
                                arb_str = f"\n🛡️ Arbitre : {m['Arbitre']}" if m.get('Arbitre') and m['Arbitre'] != "Non attribué" else ""
                                t_nom = m.get('Nom_Tour') or (f"Tour {m['Tour']}" if m.get('Tour') else "")
                                tour_str = f"\n🎯 Tour : {t_nom}" if t_nom else ""
                                ligne[col] = f"🕘 {m['Heure']} ({m['Duree']} min)\n[{m['Cat']}]{tour_str}\n{m['Combattant 1']} VS {m['Combattant 2']}{arb_str}"
                        else: ligne[col] = ""
                    grille.append(ligne)
                pd.DataFrame(grille).to_excel(writer, sheet_name="Grille de Passage", index=False, startrow=1)
                
                b_style = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                bleu = PatternFill("solid", fgColor="0055A4")
                rouge = PatternFill("solid", fgColor="EF4135")
                bleu_clair = PatternFill("solid", fgColor="DDEBF7") 
                rouge_lutte = PatternFill("solid", fgColor="E53935") 
                bleu_lutte = PatternFill("solid", fgColor="1E88E5")  
                gris_clair = PatternFill("solid", fgColor="F2F2F2")
                entete_noir = PatternFill("solid", fgColor="000000")

                for t in range(nb_tapis):
                    ws_mat = writer.book.create_sheet(f"Grille Tapis {t + 1}")
                    ws_mat.views.sheetView[0].showGridLines = True
                    ws_mat.page_setup.orientation = ws_mat.ORIENTATION_LANDSCAPE
                    ws_mat.page_setup.paperSize = ws_mat.PAPERSIZE_A4
                    ws_mat.sheet_properties.pageSetUpPr.fitToPage = True
                    ws_mat.page_setup.fitToWidth = 1
                    ws_mat.page_setup.fitToHeight = 0
                    
                    ws_mat.row_dimensions[1].height = 35
                    ws_mat.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
                    titre_mat = ws_mat.cell(row=1, column=1, value=f"🏆 {nom_competition.upper()} - GRILLE DE PASSAGE : TAPIS {t + 1} 🏆")
                    titre_mat.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
                    titre_mat.fill = bleu
                    titre_mat.alignment = Alignment(horizontal="center", vertical="center")
                    
                    noms_arb = ", ".join([a['Nom_Complet'] for a in tapis_arbitres[t]]) if (liste_arbitres and tapis_arbitres[t]) else "Aucun arbitre affecté"
                    ws_mat.row_dimensions[2].height = 22
                    ws_mat.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
                    sub_mat = ws_mat.cell(row=2, column=1, value=f"🛡️ Arbitrage : {noms_arb}  |  * Annotations des scores sous chaque match (Pt Clt, Actions, Total Score)")
                    sub_mat.font = Font(name="Arial", size=10, italic=True, bold=True, color="0055A4")
                    sub_mat.alignment = Alignment(horizontal="center", vertical="center")

                    matches_on_tapis = [it for it in planning_tapis[t] if it.get("Type") == "MATCH"]
                    max_nom_t = max([len(str(it.get('Lutteur1', ''))) for it in matches_on_tapis] + [len(str(it.get('Lutteur2', ''))) for it in matches_on_tapis] + [15])
                    max_club_t = max([len(str(it.get('Club1', ''))) for it in matches_on_tapis] + [len(str(it.get('Club2', ''))) for it in matches_on_tapis] + [12])
                    w_nom_t = max(max_nom_t + 4, 25)
                    w_club_t = max(max_club_t + 4, 18)

                    ws_mat.column_dimensions['A'].width = 16
                    ws_mat.column_dimensions['B'].width = 6
                    ws_mat.column_dimensions['C'].width = w_nom_t
                    ws_mat.column_dimensions['D'].width = w_club_t
                    ws_mat.column_dimensions['E'].width = 10
                    ws_mat.column_dimensions['F'].width = 5
                    ws_mat.column_dimensions['G'].width = w_nom_t
                    ws_mat.column_dimensions['H'].width = w_club_t
                    ws_mat.column_dimensions['I'].width = 10

                    r_curr = 4
                    m_count_t = 0
                    for item in planning_tapis[t]:
                        if item["Type"] == "PAUSE":
                            ws_mat.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=9)
                            p_cell = ws_mat.cell(row=r_curr, column=1, value=f"[{item['Heure']}] ⏸️ PAUSE DE LA COMPÉTITION")
                            p_cell.fill, p_cell.font, p_cell.alignment = rouge, Font(bold=True, color="FFFFFF", size=11), Alignment(horizontal="center", vertical="center")
                            ws_mat.row_dimensions[r_curr].height = 24
                            r_curr += 2
                        elif item["Type"] == "ATTENTE":
                            ws_mat.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=9)
                            a_cell = ws_mat.cell(row=r_curr, column=1, value=f"[{item['Heure']}] ⏳ {item['Texte']}")
                            a_cell.fill, a_cell.font, a_cell.alignment = PatternFill("solid", fgColor="EFEFEF"), Font(italic=True, color="666666", size=10), Alignment(horizontal="center", vertical="center")
                            ws_mat.row_dimensions[r_curr].height = 22
                            r_curr += 2
                        elif item["Type"] == "MATCH":
                            m_count_t += 1
                            ws_mat.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=9)
                            arb_txt = f"  |  🛡️ Arbitre : {item['Arbitre']}" if item.get('Arbitre') and item['Arbitre'] != "Non attribué" else ""
                            tour_label = item.get('Nom_Tour') or (f"Tour {item['Tour']}" if item.get('Tour') else "")
                            tour_txt = f"  |  🎯 Tour : {tour_label}" if tour_label else ""
                            hdr_text = f"MATCH N° {m_count_t}  |  🕘 {item['Heure']} ({item['Duree']} min)  |  Catégorie : {item['Cat']}{tour_txt}{arb_txt}"
                            h_cell = ws_mat.cell(row=r_curr, column=1, value=hdr_text)
                            age_m = extraire_age_de_texte(item.get('Cat', ''))
                            cfg_m = COULEURS_AGE_GRILLE.get(age_m, COULEURS_AGE_GRILLE['AUTRE'])
                            h_cell.fill, h_cell.font, h_cell.alignment = cfg_m['bg_excel_header'], Font(name="Arial", bold=True, color="FFFFFF", size=11), Alignment(horizontal="center", vertical="center")
                            ws_mat.row_dimensions[r_curr].height = 24
                            r_curr += 1

                            c_tour_h = ws_mat.cell(row=r_curr, column=1, value="TOUR / PHASE")
                            c_tour_h.fill, c_tour_h.font, c_tour_h.alignment, c_tour_h.border = gris_clair, Font(bold=True, size=8), Alignment(horizontal="center", vertical="center"), b_style

                            c_num_h = ws_mat.cell(row=r_curr, column=2, value="N°")
                            c_num_h.alignment, c_num_h.border, c_num_h.fill = Alignment(horizontal="center", vertical="center"), b_style, gris_clair
                            
                            c_rouge_h = ws_mat.cell(row=r_curr, column=3, value="LUTTEUR ROUGE")
                            c_rouge_h.fill, c_rouge_h.font, c_rouge_h.alignment, c_rouge_h.border = rouge_lutte, Font(bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="center"), b_style
                            ws_mat.merge_cells(start_row=r_curr, start_column=3, end_row=r_curr, end_column=4)
                            ws_mat.cell(row=r_curr, column=4).border = b_style
                            
                            c_ptr = ws_mat.cell(row=r_curr, column=5, value="Pt Clt")
                            c_ptr.font, c_ptr.alignment, c_ptr.border = Font(bold=True), Alignment(horizontal="center", vertical="center"), b_style
                            
                            ws_mat.cell(row=r_curr, column=6, value="VS").alignment = Alignment(horizontal="center", vertical="center")
                            
                            c_bleu_h = ws_mat.cell(row=r_curr, column=7, value="LUTTEUR BLEU")
                            c_bleu_h.fill, c_bleu_h.font, c_bleu_h.alignment, c_bleu_h.border = bleu_lutte, Font(bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="center"), b_style
                            ws_mat.merge_cells(start_row=r_curr, start_column=7, end_row=r_curr, end_column=8)
                            ws_mat.cell(row=r_curr, column=8).border = b_style
                            
                            c_ptb = ws_mat.cell(row=r_curr, column=9, value="Pt Clt")
                            c_ptb.font, c_ptb.alignment, c_ptb.border = Font(bold=True), Alignment(horizontal="center", vertical="center"), b_style
                            
                            r_curr += 1
                            
                            c_tour_v = ws_mat.cell(row=r_curr, column=1, value=tour_label)
                            c_tour_v.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                            c_tour_v.font = Font(bold=True, size=9, color="0055A4")
                            c_tour_v.border = b_style
                            ws_mat.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr+2, end_column=1)
                            ws_mat.cell(row=r_curr+1, column=1).border = b_style
                            ws_mat.cell(row=r_curr+2, column=1).border = b_style

                            ws_mat.cell(row=r_curr, column=2, value=m_count_t).alignment = Alignment(horizontal="center", vertical="center")
                            ws_mat.cell(row=r_curr, column=2).font = Font(bold=True, color="E53935", size=12)
                            ws_mat.cell(row=r_curr, column=2).border = b_style
                            ws_mat.merge_cells(start_row=r_curr, start_column=2, end_row=r_curr+2, end_column=2)
                            ws_mat.cell(row=r_curr+1, column=2).border = b_style
                            ws_mat.cell(row=r_curr+2, column=2).border = b_style
                            
                            c1_str = f"{item['Combattant 1']}"
                            if item.get('Club 1'): c1_str += f" ({item['Club 1']})"
                            if item.get('Comité 1') and item['Comité 1'] != 'Comité Non Renseigné': c1_str += f" - {item['Comité 1']}"
                            
                            c_r_info = ws_mat.cell(row=r_curr, column=3, value=c1_str)
                            c_r_info.border = b_style
                            c_r_info.alignment = Alignment(vertical="center")
                            ws_mat.merge_cells(start_row=r_curr, start_column=3, end_row=r_curr, end_column=4)
                            ws_mat.cell(row=r_curr, column=4).border = b_style
                            
                            box_ptr = ws_mat.cell(row=r_curr, column=5)
                            box_ptr.border, box_ptr.fill = b_style, gris_clair
                            box_ptr.alignment = Alignment(horizontal="center", vertical="center")
                            
                            c_vs_mid = ws_mat.cell(row=r_curr, column=6, value="-")
                            c_vs_mid.alignment = Alignment(horizontal="center", vertical="center")
                            c_vs_mid.border = b_style
                            
                            c2_str = f"{item['Combattant 2']}"
                            if item.get('Club 2'): c2_str += f" ({item['Club 2']})"
                            if item.get('Comité 2') and item['Comité 2'] != 'Comité Non Renseigné': c2_str += f" - {item['Comité 2']}"
                            
                            c_b_info = ws_mat.cell(row=r_curr, column=7, value=c2_str)
                            c_b_info.border = b_style
                            c_b_info.alignment = Alignment(vertical="center")
                            ws_mat.merge_cells(start_row=r_curr, start_column=7, end_row=r_curr, end_column=8)
                            ws_mat.cell(row=r_curr, column=8).border = b_style
                            
                            box_ptb = ws_mat.cell(row=r_curr, column=9)
                            box_ptb.border, box_ptb.fill = b_style, gris_clair
                            box_ptb.alignment = Alignment(horizontal="center", vertical="center")
                            
                            def reg_tapis_slot(k, sl):
                                if k not in tapis_slots_map:
                                    tapis_slots_map[k] = [sl]
                                elif isinstance(tapis_slots_map[k], list):
                                    tapis_slots_map[k].append(sl)
                                else:
                                    tapis_slots_map[k] = [tapis_slots_map[k], sl]

                            cat_m = item.get('Cat', '')
                            c1_m = str(item.get('Combattant 1', ''))
                            c2_m = str(item.get('Combattant 2', ''))
                            s1 = (ws_mat, f"C{r_curr}")
                            s2 = (ws_mat, f"G{r_curr}")
                            
                            reg_tapis_slot((cat_m, c1_m), s1)
                            reg_tapis_slot((cat_m, c2_m), s2)
                            reg_tapis_slot(c1_m, s1)
                            reg_tapis_slot(c2_m, s2)
                            
                            b_c1 = re.sub(r'\s*\[.*?\]', '', c1_m).strip()
                            b_c2 = re.sub(r'\s*\[.*?\]', '', c2_m).strip()
                            if b_c1 and b_c1 != c1_m:
                                reg_tapis_slot((cat_m, b_c1), s1)
                                reg_tapis_slot(b_c1, s1)
                            if b_c2 and b_c2 != c2_m:
                                reg_tapis_slot((cat_m, b_c2), s2)
                                reg_tapis_slot(b_c2, s2)
                            
                            r_curr += 1
                            
                            ws_mat.cell(row=r_curr, column=3, value="Points Techniques (Actions)").font = Font(size=9, italic=True)
                            ws_mat.merge_cells(start_row=r_curr, start_column=3, end_row=r_curr, end_column=4)
                            ws_mat.cell(row=r_curr, column=5, value="Total Score").font = Font(size=9, italic=True)
                            
                            ws_mat.cell(row=r_curr, column=7, value="Points Techniques (Actions)").font = Font(size=9, italic=True)
                            ws_mat.merge_cells(start_row=r_curr, start_column=7, end_row=r_curr, end_column=8)
                            ws_mat.cell(row=r_curr, column=9, value="Total Score").font = Font(size=9, italic=True)
                            
                            r_curr += 1
                            
                            ws_mat.row_dimensions[r_curr].height = 25
                            c_act_r = ws_mat.cell(row=r_curr, column=3)
                            c_act_r.border = b_style
                            ws_mat.cell(row=r_curr, column=4).border = b_style
                            ws_mat.merge_cells(start_row=r_curr, start_column=3, end_row=r_curr, end_column=4)
                            
                            ws_mat.cell(row=r_curr, column=5).border = b_style
                            
                            c_act_b = ws_mat.cell(row=r_curr, column=7)
                            c_act_b.border = b_style
                            ws_mat.cell(row=r_curr, column=8).border = b_style
                            ws_mat.merge_cells(start_row=r_curr, start_column=7, end_row=r_curr, end_column=8)
                            
                            ws_mat.cell(row=r_curr, column=9).border = b_style

                            # Enregistrement des coordonnées des cases Pt Clt, Actions et Total Score sur la Grille Tapis X
                            m_coord_info = {
                                'sheet': f"Grille Tapis {t + 1}",
                                'sheet_name': f"Grille Tapis {t + 1}",
                                'ptr_cell': f"E{r_curr - 2}",
                                'ptb_cell': f"I{r_curr - 2}",
                                'act_r_cell': f"C{r_curr}",
                                'tot_r_cell': f"E{r_curr}",
                                'act_b_cell': f"G{r_curr}",
                                'tot_b_cell': f"I{r_curr}",
                                'p1': item['Combattant 1'],
                                'p2': item['Combattant 2']
                            }
                            coords_matchs_tapis[(cat_m, c1_m, c2_m)] = m_coord_info
                            coords_matchs_tapis[(cat_m, c2_m, c1_m)] = m_coord_info
                            if (b_c1 and b_c1 != c1_m) or (b_c2 and b_c2 != c2_m):
                                coords_matchs_tapis[(cat_m, b_c1 or c1_m, b_c2 or c2_m)] = m_coord_info
                                coords_matchs_tapis[(cat_m, b_c2 or c2_m, b_c1 or c1_m)] = m_coord_info
                            
                            r_curr += 2 
                
                for ws_name in writer.book.sheetnames:
                    ws_sheet = writer.book[ws_name]
                    ws_sheet.page_setup.orientation = ws_sheet.ORIENTATION_LANDSCAPE
                    ws_sheet.page_setup.paperSize = ws_sheet.PAPERSIZE_A4
                    ws_sheet.sheet_properties.pageSetUpPr.fitToPage = True
                    ws_sheet.page_setup.fitToWidth = 1
                    ws_sheet.page_setup.fitToHeight = 0
                
                ws_res = writer.sheets["Résumé"]
                for cell in ws_res[1]: 
                    cell.fill, cell.font, cell.alignment = bleu, Font(bold=True, color="FFFFFF"), Alignment(horizontal="center")
                ws_res.column_dimensions['A'].width = 50
                ws_res.column_dimensions['B'].width = 25
                for row in ws_res.iter_rows(min_row=2, max_row=ws_res.max_row):
                    for cell in row:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.font = Font(size=12)
                
                ws_grille = writer.sheets["Grille de Passage"]
                ws_grille.row_dimensions[1].height = 65
                ws_grille.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nb_tapis)
                titre_cell = ws_grille.cell(row=1, column=1, value=f"🏆 {nom_competition.upper()} — PLANNING OFFICIEL  (🟡 U7 | 🟢 U9 | 🔵 U11 | 🟣 U13) 🏆")
                titre_cell.font = Font(name="Arial", size=20, bold=True, color="FFFFFF")
                titre_cell.fill = bleu
                titre_cell.alignment = Alignment(horizontal="center", vertical="center")
                
                try:
                    from openpyxl.drawing.image import Image as OpenpyxlImage
                    if os.path.exists("logo_fflda.png"):
                        img = OpenpyxlImage("logo_fflda.png")
                    else:
                        url_logo = "https://www.fflutte.com/content/uploads/2021/10/fflutte-bleu-1024x842.png"
                        req = urllib.request.Request(url_logo, headers={'User-Agent': 'Mozilla/5.0'})
                        with urllib.request.urlopen(req) as response: img_data = io.BytesIO(response.read())
                        img = OpenpyxlImage(img_data)
                    img.height, img.width = 30, 36
                    ws_grille.add_image(img, 'A1')
                except Exception: pass 

                ws_grille.freeze_panes = 'A3'
                
                for col in range(1, nb_tapis + 1):
                    c = ws_grille.cell(row=2, column=col)
                    c.fill, c.font, c.border = rouge, Font(bold=True, size=14, color="FFFFFF"), b_style
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    ws_grille.column_dimensions[c.column_letter].width = 45
                    
                for row in ws_grille.iter_rows(min_row=3, max_row=ws_grille.max_row):
                    ws_grille.row_dimensions[row[0].row].height = 90
                    for cell in row:
                        cell.border, cell.alignment = b_style, Alignment(wrap_text=True, horizontal="center", vertical="center")
                        if cell.value:
                            val_s = str(cell.value)
                            if "PAUSE" in val_s:
                                cell.fill, cell.font = rouge, Font(name="Arial", bold=True, color="FFFFFF", size=12)
                            elif any(k in val_s for k in ["Attente", "Pesée", "échauffement", "Repos"]):
                                cell.fill, cell.font = PatternFill("solid", fgColor="EFEFEF"), Font(name="Arial", italic=True, color="666666", size=11)
                            else:
                                age_k = extraire_age_de_texte(val_s)
                                cfg_c = COULEURS_AGE_GRILLE.get(age_k, COULEURS_AGE_GRILLE['AUTRE'])
                                cell.fill = cfg_c['bg_excel_pastel']
                                cell.font = Font(name="Arial", size=10, color="0F172A")

                df_excel_arb = None
                if liste_arbitres:
                    lignes_excel_arb = []
                    for t in range(nb_tapis):
                        for a in tapis_arbitres[t]:
                            lignes_excel_arb.append({
                                "Tapis Affecté": f"Tapis {t + 1}",
                                "Nom": a['Nom'],
                                "Prénom": a['Prenom'],
                                "N° Licence": a['Licence'],
                                "Club": a['Club'],
                                "Comité Régional": a['Comite']
                            })
                    if lignes_excel_arb:
                        df_excel_arb = pd.DataFrame(lignes_excel_arb)
                        df_excel_arb.to_excel(writer, sheet_name="Corps d'Arbitrage", index=False, startrow=4)
                        ws_arb_sheet = writer.sheets["Corps d'Arbitrage"]
                        ws_arb_sheet.views.sheetView[0].showGridLines = True
                        ws_arb_sheet.cell(row=1, column=1, value=f"COMPÉTITION : {nom_competition.upper()}").font = Font(name="Arial", size=15, bold=True, color="0055A4")
                        ws_arb_sheet.cell(row=2, column=1, value="🛡️ CORPS D'ARBITRAGE ET AFFECTATION AUX TAPIS - FFLDA").font = Font(name="Arial", size=12, bold=True, color="666666")
                        ws_arb_sheet.cell(row=3, column=1, value=f"Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}").font = Font(name="Arial", size=9, italic=True, color="888888")
                        
                        for col_idx in range(1, len(df_excel_arb.columns) + 1):
                            cell = ws_arb_sheet.cell(row=5, column=col_idx)
                            cell.fill, cell.font, cell.alignment = bleu, Font(name="Arial", size=10, bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="center")
                        
                        for row_idx in range(6, ws_arb_sheet.max_row + 1):
                            is_even = (row_idx % 2 == 0)
                            for col_idx in range(1, len(df_excel_arb.columns) + 1):
                                cell = ws_arb_sheet.cell(row=row_idx, column=col_idx)
                                cell.border = b_style
                                cell.fill = bleu_clair if is_even else PatternFill(fill_type=None)
                                cell.alignment = Alignment(horizontal="center", vertical="center")

                        ws_arb_sheet.column_dimensions['A'].width = 15
                        ws_arb_sheet.column_dimensions['B'].width = 25
                        ws_arb_sheet.column_dimensions['C'].width = 20
                        ws_arb_sheet.column_dimensions['D'].width = 15
                        ws_arb_sheet.column_dimensions['E'].width = 25
                        ws_arb_sheet.column_dimensions['F'].width = 25

                rouge_lutte = PatternFill("solid", fgColor="E53935") 
                bleu_lutte = PatternFill("solid", fgColor="1E88E5")  
                gris_clair = PatternFill("solid", fgColor="F2F2F2")
                entete_noir = PatternFill("solid", fgColor="000000")

                feuilles_creees = set()
                for nom_poule, liste_p in participants_par_poule.items():
                    nom_base = abreger_nom_onglet(nom_poule)
                    nom_onglet_court = nom_base
                    suffix_i = 1
                    while nom_onglet_court.lower() in feuilles_creees or nom_onglet_court in writer.book.sheetnames:
                        nom_onglet_court = f"{nom_base[:28]}_{suffix_i}"
                        suffix_i += 1
                    feuilles_creees.add(nom_onglet_court.lower())
                    poule_sheet_names[nom_poule] = nom_onglet_court
                    ws_poule = writer.book.create_sheet(nom_onglet_court)
                    
                    p_obj = poule_obj_map.get(nom_poule)
                    if p_obj and p_obj.get('type_formule') == 'tableau':
                        construire_feuille_tableau_excel(ws_poule, p_obj, nom_poule, liste_p, coords_matchs_tapis, nom_competition, tapis_slots=tapis_slots_map)
                    elif p_obj and p_obj.get('type_formule') == 'poules_croisees':
                        construire_feuille_poules_croisees_excel(ws_poule, p_obj, nom_poule, liste_p, coords_matchs_tapis, nom_competition, tapis_slots=tapis_slots_map)
                    elif p_obj and p_obj.get('type_formule') == 'plateau_u7':
                        construire_feuille_plateau_u7_excel(ws_poule, p_obj, nom_poule, liste_p, coords_matchs_tapis, nom_competition)
                    else:
                        construire_feuille_poule_nordique_excel(ws_poule, nom_poule, liste_p, rondes_par_categorie.get(nom_poule, []), coords_matchs_tapis, nom_competition)

            excel_bytes_tournoi_complet = output_excel.getvalue()

            # Source du classeur Excel officiel FFLDA pour la génération PDF
            wb_officiel = getattr(writer, 'book', None)
            if wb_officiel is None:
                try:
                    wb_officiel = openpyxl.load_workbook(io.BytesIO(excel_bytes_tournoi_complet), data_only=False)
                except Exception:
                    wb_officiel = None

            pdf_bytes_tournoi_complet = None
            if wb_officiel is not None:
                try:
                    pdf_bytes_tournoi_complet = generer_pdf_depuis_classeur_excel(wb_officiel, nom_competition, wb=wb_officiel)
                except Exception as e_pdf:
                    st.warning(f"⚠️ Information : génération PDF depuis Excel : {e_pdf}")

            if not pdf_bytes_tournoi_complet:
                try:
                    pdf_bytes_tournoi_complet = generer_pdf_tournoi_complet("Dossier Officiel du Tournoi", nom_competition, sections_tournoi_complet)
                except Exception:
                    pdf_bytes_tournoi_complet = b""

            pdf_grille_bytes = None
            if wb_officiel is not None:
                sheet_gp = None
                if hasattr(wb_officiel, 'sheetnames') and "Grille de Passage" in wb_officiel.sheetnames:
                    sheet_gp = wb_officiel["Grille de Passage"]
                elif hasattr(wb_officiel, 'worksheets'):
                    for s in wb_officiel.worksheets:
                        if s.title == "Grille de Passage":
                            sheet_gp = s
                            break
                if sheet_gp is not None:
                    try:
                        pdf_grille_bytes = generer_pdf_depuis_classeur_excel([sheet_gp], nom_competition, wb=wb_officiel)
                    except Exception:
                        pdf_grille_bytes = None
            html_tournoi_complet = generer_document_html_imprimable("Feuilles Officieuses du Tournoi & Poules FFLDA", nom_competition, sections_tournoi_complet)

            st.markdown("### 📄 Impression & Exportations Officielles (Format Excel FFLDA - A4 Portrait)")
            col_pdf_top, col_xl_top, col_html_top = st.columns([1, 1, 1])
            with col_pdf_top:
                st.download_button(
                    label="📄 Télécharger le Dossier Officiel en PDF (A4 Portrait - Format Excel)",
                    data=pdf_bytes_tournoi_complet,
                    file_name=f"Dossier_Officiel_{nom_competition.replace(' ', '_')}.pdf",
                    mime="application/pdf",
                    key="btn_pdf_top"
                )
            with col_xl_top:
                st.download_button(
                    label="📥 Télécharger le Classeur Officiel Excel (.xlsx)",
                    data=excel_bytes_tournoi_complet,
                    file_name=f"Tournoi_{nom_competition.replace(' ', '_')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="btn_excel_top"
                )
            with col_html_top:
                bouton_imprimer(html_tournoi_complet, filename="Dossier_Tournoi_Impression.html", label="🖨️ Imprimer la Version Web Paysage A4", key="btn_html_top")

            st.markdown("---")

            # --- ONGLETS INTERACTIFS DE L'APPLICATION ---
            noms_onglets = ["📊 Résumé & Stats", "📅 Grille Globale", "🛡️ Équipes d'Arbitrage"] + [f"🥋 Grille Tapis {t + 1}" for t in range(nb_tapis)] + [
                f"🏅 {p[:18]}" if 'plateau' in p.lower() else f"Groupe : {p[:18]}" for p in participants_par_poule.keys()
            ]
            onglets_ui = st.tabs(noms_onglets)
            
            with onglets_ui[0]:
                st.subheader("📊 Résumé prévisionnel de la journée")
                st.table(pd.DataFrame(lignes_accueil))
                
                col_m1, col_m2, col_m3, col_m4, col_m5, col_m6, col_m7, col_m8 = st.columns(8)
                col_m1.metric("Participants (pesés)", total_participants_peses)
                col_m2.metric("Absents / Non pesés", total_non_peses)
                col_m3.metric("Groupes U7 (Plateaux)", len(poules_u7))
                col_m4.metric("Poules U9", len(poules_u9))
                col_m5.metric("Poules U11", len(poules_u11))
                col_m6.metric("Groupes U13", len(poules_u13))
                col_m7.metric("Matchs générés", total_matchs_calcules)
                col_m8.metric("Arbitres engagés", len(liste_arbitres))

                if liste_arbitres:
                    st.markdown("#### 🛡️ Désignation des Équipes d'Arbitrage par Tapis")
                    lignes_arb_sum = []
                    for t in range(nb_tapis):
                        noms_arb = ", ".join([a['Nom_Complet'] for a in tapis_arbitres[t]]) if tapis_arbitres[t] else "Aucun arbitre affecté"
                        lignes_arb_sum.append({"Tapis": f"Tapis {t + 1}", "Effectif": f"{len(tapis_arbitres[t])} arbitres", "Équipe d'Arbitrage Désignée": noms_arb})
                    st.table(pd.DataFrame(lignes_arb_sum))

            with onglets_ui[1]:
                st.subheader("📅 Grille Globale de Passage - Tous les Tapis")
                if liste_arbitres:
                    st.markdown("**🛡️ Équipes d'arbitrage affectées aux tapis :**")
                    cols_arb_disp = st.columns(nb_tapis)
                    for t in range(nb_tapis):
                        with cols_arb_disp[t]:
                            arb_list_txt = "\n".join([f"• **{a['Nom_Complet']}** ({a['Club']})" for a in tapis_arbitres[t]]) if tapis_arbitres[t] else "• Aucun"
                            st.info(f"**Tapis {t+1}** ({len(tapis_arbitres[t])} arbitres) :\n\n{arb_list_txt}")
                
                # Grille interactive avec couleurs par catégorie d'âge
                html_grille_ui = generer_html_grille_coloree(grille_ui, nb_tapis)
                st.markdown(html_grille_ui, unsafe_allow_html=True)
                st.write("")
                
                col_p1, col_h1 = st.columns([1, 1])
                with col_p1:
                    if pdf_grille_bytes:
                        st.download_button(
                            label="📄 Télécharger la Grille Globale en PDF (A4 Portrait - Format Excel)",
                            data=pdf_grille_bytes,
                            file_name=f"Grille_Passage_{nom_competition.replace(' ', '_')}.pdf",
                            mime="application/pdf",
                            key="btn_pdf_grille"
                        )
                with col_h1:
                    bouton_imprimer(html_tournoi_complet, filename="Grille_Tapis_Apercu_Web.html", label="🖨️ Aperçu Web HTML (Optionnel)", key="btn_t1")

            with onglets_ui[2]:
                st.subheader("🛡️ Désignation et Affectation des Arbitres par Tapis")
                if liste_arbitres:
                    if wb_officiel is not None and hasattr(wb_officiel, 'sheetnames') and "Corps d'Arbitrage" in wb_officiel.sheetnames:
                        try:
                            pdf_arb_bytes = generer_pdf_depuis_classeur_excel([wb_officiel["Corps d'Arbitrage"]], nom_competition, wb=wb_officiel)
                            if pdf_arb_bytes:
                                st.download_button(
                                    label="📄 Télécharger le Corps d'Arbitrage en PDF (A4 Portrait - Format Excel)",
                                    data=pdf_arb_bytes,
                                    file_name=f"Corps_Arbitrage_{nom_competition.replace(' ', '_')}.pdf",
                                    mime="application/pdf",
                                    key="btn_pdf_arb"
                                )
                                st.write("")
                        except Exception:
                            pass
                    for t in range(nb_tapis):
                        st.markdown(f"#### 🥋 Tapis {t + 1} ({len(tapis_arbitres[t])} arbitres)")
                        if tapis_arbitres[t]:
                            df_arb_tapis = pd.DataFrame(tapis_arbitres[t])[['Nom', 'Prenom', 'Licence', 'Club', 'Comite']]
                            df_arb_tapis.columns = ['Nom', 'Prénom', 'N° Licence', 'Club', 'Comité Régional']
                            st.table(df_arb_tapis)
                        else:
                            st.info("Aucun arbitre affecté à ce tapis.")
                else:
                    st.info("Aucun fichier d'arbitres n'a été chargé.")

            # Onglets Grille Tapis individuel avec saisie des scores
            for t in range(nb_tapis):
                with onglets_ui[3 + t]:
                    st.subheader(f"🥋 Grille de Passage & Feuille de Marque — Tapis {t + 1}")
                    sheet_mat_name = f"Grille Tapis {t + 1}"
                    pdf_tapis_bytes = None
                    if wb_officiel is not None and hasattr(wb_officiel, 'sheetnames') and sheet_mat_name in wb_officiel.sheetnames:
                        try:
                            pdf_tapis_bytes = generer_pdf_depuis_classeur_excel([wb_officiel[sheet_mat_name]], nom_competition, wb=wb_officiel)
                        except Exception:
                            pdf_tapis_bytes = None
                    
                    if pdf_tapis_bytes:
                        st.download_button(
                            label=f"📄 Télécharger la Grille Tapis {t + 1} en PDF (A4 Portrait - Format Excel)",
                            data=pdf_tapis_bytes,
                            file_name=f"Grille_Tapis_{t + 1}_{nom_competition.replace(' ', '_')}.pdf",
                            mime="application/pdf",
                            key=f"btn_pdf_tapis_{t}"
                        )
                        st.write("")
                    if liste_arbitres and tapis_arbitres[t]:
                        arb_names_st = ", ".join([f"**{a['Nom_Complet']}** ({a['Club']})" for a in tapis_arbitres[t]])
                        st.info(f"🛡️ **Équipe d'arbitrage désignée (Tapis {t + 1})** : {arb_names_st}")
                    else:
                        st.caption("🛡️ Aucun arbitre désigné spécifiquement sur ce tapis.")
                    
                    m_count_st = 0
                    for m in planning_tapis[t]:
                        if m["Type"] == "PAUSE":
                            st.warning(f"[{m['Heure']}] ⏸️ PAUSE DE LA COMPÉTITION")
                        elif m["Type"] == "ATTENTE":
                            st.info(f"[{m['Heure']}] {m['Texte']}")
                        elif m["Type"] == "MATCH":
                            m_count_st += 1
                            age_m = extraire_age_de_texte(m['Cat'])
                            badge_age = COULEURS_AGE_GRILLE.get(age_m, {}).get('badge', '')
                            badge_str = f"{badge_age} " if badge_age else ""
                            arb_info_st = f" | 🛡️ Arbitre : {m['Arbitre']}" if m.get('Arbitre') and m['Arbitre'] != "Non attribué" else ""
                            t_nom = m.get('Nom_Tour') or (f"Tour {m['Tour']}" if m.get('Tour') else "")
                            tour_st = f" | 🎯 Tour : **{t_nom}**" if t_nom else ""
                            st.markdown(f"#### 🤼 MATCH N° {m_count_st} — 🕘 {m['Heure']} ({m['Duree']} min) | Catégorie : {badge_str}`{m['Cat']}`{tour_st}{arb_info_st}")
                            
                            c1_cl = f" ({m.get('Club 1', '')})" if m.get('Club 1') else ""
                            c1_co = f" - {m.get('Comité 1', '')}" if (m.get('Comité 1') and m.get('Comité 1') != 'Comité Non Renseigné') else ""
                            c2_cl = f" ({m.get('Club 2', '')})" if m.get('Club 2') else ""
                            c2_co = f" - {m.get('Comité 2', '')}" if (m.get('Comité 2') and m.get('Comité 2') != 'Comité Non Renseigné') else ""
                            
                            df_m_ui = pd.DataFrame([
                                {
                                    "N°": m_count_st,
                                    "LUTTEUR ROUGE": f"🔴 {m['Combattant 1']}{c1_cl}{c1_co}",
                                    "Pt Clt (Rouge)": "[   ]",
                                    "Points Techniques (Actions Rouge)": "[                                 ]",
                                    "Total Score (Rouge)": "[   ]",
                                    "VS": "VS",
                                    "LUTTEUR BLEU": f"🔵 {m['Combattant 2']}{c2_cl}{c2_co}",
                                    "Pt Clt (Bleu)": "[   ]",
                                    "Points Techniques (Actions Bleu)": "[                                 ]",
                                    "Total Score (Bleu)": "[   ]"
                                }
                            ])
                            st.table(df_m_ui)

            start_idx_poules = 3 + nb_tapis
            for idx, (nom_poule, liste_p) in enumerate(participants_par_poule.items(), start=start_idx_poules):
                with onglets_ui[idx]:
                    nom_safe = nom_poule.replace(' ', '_').replace('|', '_').replace('/', '_')
                    sheet_nom = poule_sheet_names.get(nom_poule)
                    pdf_sheet_bytes = None
                    if wb_officiel is not None and sheet_nom and hasattr(wb_officiel, 'sheetnames') and sheet_nom in wb_officiel.sheetnames:
                        try:
                            pdf_sheet_bytes = generer_pdf_depuis_classeur_excel([wb_officiel[sheet_nom]], nom_competition, wb=wb_officiel)
                        except Exception:
                            pdf_sheet_bytes = None

                    p_obj = poule_obj_map.get(nom_poule)
                    if p_obj and p_obj.get('type_formule') == 'plateau_u7':
                        st.subheader(f"🏅 Animation U7 : {nom_poule}")
                        st.info("Formule officielle FFLDA : Découverte pédagogique sous forme de 3 plateaux d'activités avec rotation (Motricité, Opposition au sol, Lutte debout). Tous les enfants sont récompensés !")
                        df_u7_tab = pd.DataFrame([
                            {
                                "N°": idx_p,
                                "Nom Prénom": p.get('Nom', ''),
                                "Club": p.get('Club', ''),
                                "Poids": formater_poids(p.get('Poids', '')),
                                "Plateau 1 : Motricité & Agilité": "✓ Validé",
                                "Plateau 2 : Ateliers techniques": "✓ Validé",
                                "Plateau 3 : Oppositions": "✓ Validé",
                                "Récompense": "🥇 Médaille d'Or"
                            }
                            for idx_p, p in enumerate(liste_p, 1)
                        ])
                        st.table(df_u7_tab)
                        doc_print_u7 = generer_document_plateau_u7_imprimable(nom_poule, nom_competition, liste_p)
                        col_u7_1, col_u7_2 = st.columns([1, 1])
                        with col_u7_1:
                            if pdf_sheet_bytes:
                                st.download_button(
                                    label="📄 Télécharger cette Feuille en PDF (A4 Portrait - Format Excel)",
                                    data=pdf_sheet_bytes,
                                    file_name=f"Plateau_U7_{nom_safe}.pdf",
                                    mime="application/pdf",
                                    key=f"btn_pdf_u7_{idx}"
                                )
                        with col_u7_2:
                            bouton_imprimer(doc_print_u7, filename=f"Plateau_U7_{nom_safe}.html", label="🖨️ Aperçu Web HTML (Optionnel)", key=f"btn_print_u7_{idx}")

                    else:
                        st.subheader(f"Feuille : {nom_poule}")
                        df_poule_vue = pd.DataFrame(liste_p)[['Nom', 'Club', 'Poids']]
                        st.table(df_poule_vue)

                        if p_obj and p_obj.get('type_formule') == 'poules_croisees':
                            st.markdown("##### 🥋 Répartition en 2 Poules de 3 :")
                            c_pa, c_pb = st.columns(2)
                            with c_pa:
                                st.markdown("**Poule A**")
                                st.table(pd.DataFrame(p_obj['poule_a'])[['Nom', 'Club', 'Poids']])
                            with c_pb:
                                st.markdown("**Poule B**")
                                st.table(pd.DataFrame(p_obj['poule_b'])[['Nom', 'Club', 'Poids']])
                            
                            st.markdown("##### 🤼 Tableau Visuel de la Phase Finale (Gauche ➔ Droite) :")
                            bracket_html = generer_arbre_tableau_html(p_obj)
                            st.markdown(bracket_html, unsafe_allow_html=True)
                            doc_print_bracket = generer_document_bracket_imprimable(nom_poule, nom_competition, bracket_html)
                            col_pc_1, col_pc_2 = st.columns([1, 1])
                            with col_pc_1:
                                if pdf_sheet_bytes:
                                    st.download_button(
                                        label="📄 Télécharger cette Feuille en PDF (A4 Portrait - Format Excel)",
                                        data=pdf_sheet_bytes,
                                        file_name=f"Tableau_Croise_{nom_safe}.pdf",
                                        mime="application/pdf",
                                        key=f"btn_pdf_pc_{idx}"
                                    )
                            with col_pc_2:
                                bouton_imprimer(doc_print_bracket, filename=f"Tableau_{nom_safe}.html", label="🖨️ Aperçu Web HTML (Optionnel)", key=f"btn_print_tab_{idx}")

                        elif p_obj and p_obj.get('type_formule') == 'tableau':
                            st.markdown("##### 🤼 Tableau Visuel Officiel FFLDA (Gauche ➔ Droite) :")
                            bracket_html = generer_arbre_tableau_html(p_obj)
                            st.markdown(bracket_html, unsafe_allow_html=True)
                            doc_print_bracket = generer_document_bracket_imprimable(nom_poule, nom_competition, bracket_html)
                            col_tb_1, col_tb_2 = st.columns([1, 1])
                            with col_tb_1:
                                if pdf_sheet_bytes:
                                    st.download_button(
                                        label="📄 Télécharger cette Feuille en PDF (A4 Portrait - Format Excel)",
                                        data=pdf_sheet_bytes,
                                        file_name=f"Tableau_{nom_safe}.pdf",
                                        mime="application/pdf",
                                        key=f"btn_pdf_tab_{idx}"
                                    )
                            with col_tb_2:
                                bouton_imprimer(doc_print_bracket, filename=f"Tableau_{nom_safe}.html", label="🖨️ Aperçu Web HTML (Optionnel)", key=f"btn_print_tab_{idx}")
                        else:
                            st.markdown("##### 🥋 Combats & Fiche Imprimable :")
                            doc_print_poule = generer_document_poule_imprimable(nom_poule, nom_competition, liste_p, rondes_par_categorie.get(nom_poule, []))
                            col_po_1, col_po_2 = st.columns([1, 1])
                            with col_po_1:
                                if pdf_sheet_bytes:
                                    st.download_button(
                                        label="📄 Télécharger cette Poule en PDF (A4 Portrait - Format Excel)",
                                        data=pdf_sheet_bytes,
                                        file_name=f"Poule_{nom_safe}.pdf",
                                        mime="application/pdf",
                                        key=f"btn_pdf_poule_{idx}"
                                    )
                            with col_po_2:
                                bouton_imprimer(doc_print_poule, filename=f"Poule_{nom_safe}.html", label="🖨️ Aperçu Web HTML (Optionnel)", key=f"btn_print_poule_{idx}")

            st.markdown("---")
            
            # --- TÉLÉCHARGEMENTS COMPLETS DU TOURNOI (FORMATS OFFICIELS FFLDA) ---
            st.markdown("### 📥 Téléchargements Complets du Tournoi (Formats Officiels FFLDA)")
            col_b_pdf, col_b_xl = st.columns(2)
            with col_b_pdf:
                st.download_button(
                    label="📄 Télécharger le Dossier Officiel en PDF (A4 Portrait - Format Excel)",
                    data=pdf_bytes_tournoi_complet,
                    file_name=f"Dossier_Officiel_{nom_competition.replace(' ', '_')}.pdf",
                    mime="application/pdf",
                    key="btn_pdf_bottom"
                )
            with col_b_xl:
                st.download_button(
                    label="📥 Télécharger le Classeur Officiel Excel (.xlsx)",
                    data=excel_bytes_tournoi_complet,
                    file_name=f"Tournoi_{nom_competition.replace(' ', '_')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="btn_excel_bottom"
                )
        except Exception as e:
            import traceback
            st.error(f"Erreur lors de l'analyse du fichier : {e}")
            st.code(traceback.format_exc())
